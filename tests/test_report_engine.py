"""Report Engine P0 (issue #33): XlsxWriter XLSX + WeasyPrint PDF.

Covers: request schema validation, XLSX/PDF structure, Vietnamese text,
filtered-total parity with the existing productivity export and the
Phiên làm việc screen query, formula-injection safety, openpyxl / HTML
fallbacks when libraries are missing, RBAC denial, anonymous 401 and
per-user isolation. No database: ReportRepository is replaced by a fake
that serves fixed fixtures and records the filters it was called with.
"""
from __future__ import annotations

import zipfile
from datetime import date, datetime, timezone
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from mesflow.db.repositories import rbac as rbac_module
from mesflow.reporting import datasets as datasets_module
from mesflow.reporting import pdf as pdf_module
from mesflow.reporting import xlsx as xlsx_module
from mesflow.reporting.datasets import DATASETS, build_document
from mesflow.reporting.engine import ReportForbidden, UnknownDataset, build_report
from mesflow.reporting.schema import ReportValidationError, parse_report_request
from mesflow.web import app as app_module
from mesflow.web import auth as auth_module
from mesflow.web import report_engine as report_engine_module
from mesflow.web.productivity_excel import build_employee_productivity_workbook
from mesflow.web.productivity_export import load_export_data

ROOT = Path(__file__).resolve().parents[1]
UTC = timezone.utc
TZ = 'Asia/Ho_Chi_Minh'
T = lambda d, h, m=0: datetime(2026, 9, d, h, m, tzinfo=UTC)  # 01:00Z = 08:00 ICT
EVIL = '=HYPERLINK("http://evil.example/x","bấm")'


def _emp(emp_id, code, name, dept, sessions, valid, pct, good, defect, seconds):
    return {"employee_id": emp_id, "employee_code": code, "employee_name": name, "department": dept, "team": "",
            "completed_sessions": sessions, "completed_valid_sessions": valid,
            "completed_invalid_sessions": sessions - valid, "repair_sessions": 0, "productivity_percent": pct,
            "good_qty": good, "defect_qty": defect, "worked_seconds": seconds}


def _ps(sid, emp_id, code, name, dept, op, start, end, good, defect, std, pct, **extra):
    row = {"session_id": sid, "employee_id": emp_id, "employee_code": code, "employee_name": name,
           "department": dept, "team": "", "status": "CLOSED", "started_at": start, "ended_at": end,
           "work_date": date(2026, 9, end.day), "actual_seconds": (end - start).total_seconds(),
           "standard_seconds_per_unit": std, "expected_seconds": std * (good + defect), "good_qty": good,
           "defect_qty": defect, "po_code": "PO-1", "part_code": "P-1", "part_name": "Thân áo",
           "operation_code": op, "operation_name": f"Công đoạn {op}", "station_code": "ST1",
           "closed_by_system": False, "quantity_confirmed": True, "close_reason": None, "note": None,
           "is_repair": False, "operation_type": "PRODUCTION", "completion_percent": pct}
    row.update(extra)
    return row


EMPLOYEES = [
    _emp(1, "NV01", "Nguyễn Văn An", "May", 3, 2, 90.0, 60, 0, 9000),
    _emp(2, "NV02", "Trần Thị Bình", "May", 1, 1, 83.33, 45, 5, 3600),
    _emp(3, "NV03", EVIL, "Kho", 1, 1, 100.0, 10, 0, 3600),
]
PSESSIONS = [
    _ps(12, 1, "NV01", "Nguyễn Văn An", "May", "OP10", T(2, 3), T(2, 4), 36, 0, 100, 100.0),
    _ps(11, 1, "NV01", "Nguyễn Văn An", "May", "OP10", T(2, 1), T(2, 2), 24, 0, 120, 80.0),
    _ps(13, 1, "NV01", "Nguyễn Văn An", "May", "OP20", T(3, 1), T(3, 1, 30), 0, 0, 0, None),
    _ps(21, 2, "NV02", "Trần Thị Bình", "May", "OP10", T(2, 5), T(2, 6), 45, 5, 60, 83.33,
        closed_by_system=True, quantity_confirmed=False, close_reason="Hết ca"),
    _ps(31, 3, "NV03", EVIL, "Kho", "OP90", T(2, 1), T(2, 2), 10, 0, 360, 100.0, note='<img src="http://evil.example/p.png">'),
]


def _ws(sid, emp_id, code, name, po_id, po, part_id, op_id, status, good, defect, rework, dur, note=None, **extra):
    row = {"session_id": sid, "employee_id": emp_id, "employee_code": code, "employee_name": name,
           "operation_id": op_id, "station_id": 1, "status": status, "started_at": T(4, 1), "ended_at": T(4, 2) if status == 'CLOSED' else None,
           "good_qty": good, "defect_qty": defect, "rework_qty": rework, "note": note, "close_reason": None,
           "closed_by_system": False, "quantity_confirmed": True, "excluded_from_reports": False,
           "exclusion_reason": None, "po_id": po_id, "po_code": po, "part_id": part_id, "part_code": f"P{part_id}",
           "part_name": "Tay áo", "operation_code": f"OP{op_id}", "operation_name": f"Ủi {op_id}",
           "station_code": "ST1", "station_name": "Trạm 1", "duration_seconds": dur, "work_duration_seconds": dur - 60}
    row.update(extra)
    return row


WSESSIONS = [
    _ws(101, 1, "NV01", "Nguyễn Văn An", 7, "PO-7", 70, 700, "CLOSED", 50, 2, 1, 3600),
    _ws(102, 1, "NV01", "Nguyễn Văn An", 7, "PO-7", 70, 701, "CLOSED", 30, 0, 0, 1800, note="Đổi kim"),
    _ws(103, 2, "NV02", "Trần Thị Bình", 7, "PO-7", 71, 710, "OPEN", 0, 0, 0, 600),
    _ws(104, 3, "NV03", EVIL, 8, "PO-8", 80, 800, "CLOSED", 12, 1, 0, 1200,
        excluded_from_reports=True, exclusion_reason="Phiên thử"),
]


class FakeRepo:
    def __init__(self, employees=None, psessions=None, wsessions=None):
        self.employees = EMPLOYEES if employees is None else employees
        self.psessions = PSESSIONS if psessions is None else psessions
        self.wsessions = WSESSIONS if wsessions is None else wsessions
        self.calls = []

    def employee_productivity(self, date_from, date_to, employee_id, department, team, limit):
        self.calls.append(("summary", date_from, date_to, employee_id, department, team))
        rows = [dict(e) for e in self.employees if (not employee_id or e["employee_id"] == employee_id)
                and (not department or e["department"] == department)]
        return {"summary": {"from": date_from or "2026-09-01", "to": date_to or "2026-09-30"}, "employees": rows}

    def employee_productivity_sessions(self, date_from, date_to, employee_id, department, team):
        self.calls.append(("sessions", date_from, date_to, employee_id, department, team))
        rows = [dict(s) for s in self.psessions if (not employee_id or s["employee_id"] == employee_id)
                and (not department or s["department"] == department)]
        return {"sessions": rows, "truncated": False}

    def session_management(self, po_id=None, part_id=None, operation_id=None, employee_id=None, status=None,
                           date_from=None, date_to=None, limit=3000):
        self.calls.append(("session_management", po_id, part_id, operation_id, employee_id, status, date_from, date_to, limit))
        items = [dict(s) for s in self.wsessions
                 if (not po_id or s["po_id"] == po_id) and (not part_id or s["part_id"] == part_id)
                 and (not operation_id or s["operation_id"] == operation_id)
                 and (not employee_id or s["employee_id"] == employee_id) and (not status or s["status"] == status)]
        return {"items": items, "filters": {}}


ALLOW_ALL = lambda permission: True
NOW = datetime(2026, 10, 10, 9, 30)


def _report(dataset, args, *, role="admin", repo=None, perm=ALLOW_ALL, user="quanly"):
    return build_report(dataset, args, repo=repo or FakeRepo(), role=role, has_permission=perm,
                        generated_by=user, timezone_name=TZ, now=NOW)


def _wb(content):
    return load_workbook(BytesIO(content))


def _data_rows(ws, header_row=5):
    rows = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if all(v is None for v in row):
            continue
        rows.append(row)
    return rows


def _col(ws, label, header_row=5):
    headers = [c.value for c in ws[header_row]]
    return headers.index(label)


# ------------------------------------------------------------------ schema
@pytest.mark.parametrize("args, message", [
    ({"from": "2026-13-01"}, "YYYY-MM-DD"),
    ({"from": "2026-09-10", "to": "2026-09-01"}, "trước hoặc bằng"),
    ({"from": "2025-01-01", "to": "2026-09-01"}, "tối đa 366"),
    ({"po_id": "1"}, "không hỗ trợ lọc theo PO"),
    ({"employee_id": "1 OR 1=1"}, "không hợp lệ"),
    ({"employee_id": "-3"}, "không hợp lệ"),
    ({"format": "csv"}, "xlsx, pdf hoặc html"),
    ({"sql": "select * from users"}, "không được hỗ trợ"),
    ({"sort": "password"}, "sắp xếp"),
    ({"search": "x" * 500}, "quá dài"),
])
def test_schema_rejects_bad_or_unsupported_parameters(args, message):
    spec = DATASETS["employee_productivity"]
    with pytest.raises(ReportValidationError, match=message):
        parse_report_request(spec.key, args, allowed_filters=spec.filters, sort_keys=spec.sort_keys)


def test_schema_accepts_all_work_session_filters():
    spec = DATASETS["work_sessions"]
    req = parse_report_request(spec.key, {"from": "2026-09-01", "to": "2026-09-30", "po_id": "7", "part_id": "70",
                                          "operation_id": "700", "employee_id": "1", "status": "closed", "format": "pdf"},
                               allowed_filters=spec.filters)
    assert req.format == "pdf" and req.status == "CLOSED" and req.po_id == 7 and req.date_from == date(2026, 9, 1)
    assert req.filters() == {"date_from": "2026-09-01", "date_to": "2026-09-30", "po_id": 7, "part_id": 70,
                             "operation_id": 700, "employee_id": 1, "status": "CLOSED"}


def test_unknown_dataset_and_permission_checked_before_any_query():
    repo = FakeRepo()
    with pytest.raises(UnknownDataset):
        _report("users_passwords", {}, repo=repo)
    with pytest.raises(ReportForbidden):
        _report("work_sessions", {"from": "2026-09-01"}, role="viewer", repo=repo)
    with pytest.raises(ReportForbidden):
        _report("employee_productivity", {}, role="operator", repo=repo, perm=lambda p: False)
    assert repo.calls == []


# ------------------------------------------------------------------ XLSX
def test_xlsx_is_written_by_xlsxwriter_with_expected_structure():
    rep = _report("employee_productivity", {"from": "2026-09-01", "to": "2026-09-30", "department": "May"})
    assert rep.engine == "xlsxwriter" and rep.content[:2] == b"PK"
    assert rep.filename == "mesflow_employee_productivity_2026-09-01_2026-09-30.xlsx"
    with zipfile.ZipFile(BytesIO(rep.content)) as z:
        assert z.testzip() is None and "xl/workbook.xml" in z.namelist()
        assert "<dc:creator>MESFlow Report Engine</dc:creator>" in z.read("docProps/core.xml").decode()
    wb = _wb(rep.content)
    assert wb.sheetnames == ["Tổng hợp", "NV01 Nguyễn Văn An", "NV02 Trần Thị Bình"]
    ws = wb["Tổng hợp"]
    assert ws["A1"].value == "BÁO CÁO NĂNG SUẤT NHÂN VIÊN · Tổng hợp năng suất nhân viên"
    assert "Từ ngày: 2026-09-01" in ws["A2"].value and "Múi giờ: Asia/Ho_Chi_Minh" in ws["A2"].value
    assert "Xuất lúc: 10/10/2026 09:30 bởi quanly" in ws["A2"].value
    assert "Bộ phận: May" in ws["A3"].value
    assert [c.value for c in ws[5]][:4] == ["STT", "Mã NV", "Nhân viên", "Bộ phận"]
    assert ws.freeze_panes == "A6" and ws.auto_filter.ref.startswith("A5:")
    assert ws.page_setup.orientation == "landscape" and int(ws.page_setup.paperSize) == 9
    assert ws.print_title_rows == "$5:$5"
    pct = ws.cell(6, _col(ws, "Năng suất TB (%)") + 1)
    assert pct.number_format == "0.0%" and pct.value == pytest.approx(0.90)
    assert ws.cell(6, _col(ws, "Tổng thời gian") + 1).number_format == "[h]:mm:ss"
    detail = wb["NV01 Nguyễn Văn An"]
    assert len(_data_rows(detail)) == 3 + 1  # 3 sessions + totals row
    assert detail.cell(6, _col(detail, "Bắt đầu") + 1).value == datetime(2026, 9, 2, 8, 0)  # local wall clock


def test_xlsx_summary_parity_with_existing_productivity_export():
    """Same filters -> the engine's Tổng hợp rows equal the existing
    /export.xlsx Tổng hợp rows, and the totals row equals their sums."""
    for args in ({}, {"department": "May"}, {"search": "bình"}, {"employee_id": "1"}, {"sort": "good_qty", "dir": "asc"}):
        full = {"from": "2026-09-01", "to": "2026-09-30", **args}
        engine = _wb(_report("employee_productivity", full).content)["Tổng hợp"]
        legacy_data = load_export_data({"from": "2026-09-01", "to": "2026-09-30",
                                        "employee_id": int(args["employee_id"]) if "employee_id" in args else None,
                                        "department": args.get("department"), "team": None,
                                        "search": args.get("search", ""), "sort": args.get("sort", "productivity_percent"),
                                        "dir": args.get("dir", "desc")}, FakeRepo())
        legacy = load_workbook(build_employee_productivity_workbook(legacy_data))["Tổng hợp"]
        e_rows, l_rows = _data_rows(engine), _data_rows(legacy, header_row=4)
        assert len(e_rows) == len(l_rows) + 1, args  # + totals row
        pick = lambda ws, row, labels, hr: tuple(row[_col(ws, lb, hr)] for lb in labels)
        labels = ["Mã NV", "Phiên hoàn tất", "Phiên hợp lệ", "Năng suất TB (%)", "Sản lượng đạt", "Lỗi"]
        for e, l in zip(e_rows, l_rows):
            assert pick(engine, e, labels, 5) == pick(legacy, l, labels, 4), args
            assert e[_col(engine, "Tổng thời gian")] == pytest.approx(l[_col(legacy, "Tổng thời gian", 4)])
        total = e_rows[-1]
        for lb in ("Phiên hoàn tất", "Sản lượng đạt", "Lỗi"):
            assert total[_col(engine, lb)] == sum(r[_col(engine, lb)] for r in e_rows[:-1]), (args, lb)
        assert str(total[_col(engine, "Nhân viên")]).startswith("Tổng cộng")


def test_work_sessions_filters_pass_through_and_totals_match_screen_query():
    repo = FakeRepo()
    args = {"from": "2026-09-01", "to": "2026-09-30", "po_id": "7"}
    rep = _report("work_sessions", args, repo=repo, role="supervisor")
    assert repo.calls == [("session_management", 7, None, None, None, None, "2026-09-01", "2026-09-30", 10000)]
    screen = FakeRepo().session_management(7, None, None, None, None, "2026-09-01", "2026-09-30")["items"]
    wb = _wb(rep.content)
    assert wb.sheetnames == ["Tổng hợp", "Chi tiết phiên"]
    detail = wb["Chi tiết phiên"]
    rows = _data_rows(detail)
    assert len(rows) == len(screen) + 1
    total = rows[-1]
    for label, key in (("Đạt", "good_qty"), ("Lỗi", "defect_qty"), ("Sửa", "rework_qty")):
        assert total[_col(detail, label)] == sum(s[key] for s in screen)
    assert total[_col(detail, "Thời lượng")].total_seconds() == pytest.approx(sum(s["duration_seconds"] for s in screen))
    summary_rows = _data_rows(wb["Tổng hợp"])
    assert summary_rows[-1][_col(wb["Tổng hợp"], "Số phiên")] == len(screen)
    assert "PO: 7" in detail["A3"].value
    statuses = {r[_col(detail, "Trạng thái")] for r in rows[:-1]}
    assert statuses == {"Đã kết thúc", "Đang làm"}


def test_text_is_never_stored_as_a_formula_or_link():
    rep = _report("work_sessions", {"from": "2026-09-01", "to": "2026-09-30"})
    with zipfile.ZipFile(BytesIO(rep.content)) as z:
        for name in z.namelist():
            if name.startswith("xl/worksheets/sheet"):
                xml = z.read(name).decode()
                assert "<f>" not in xml and "<f " not in xml and "hyperlink" not in xml.lower()
    wb = _wb(rep.content)
    found = [c for ws in wb for row in ws.iter_rows() for c in row if c.value == EVIL]
    assert found and all(c.data_type == "s" for c in found)


def test_openpyxl_fallback_when_xlsxwriter_missing(monkeypatch):
    monkeypatch.setattr(xlsx_module, "xlsxwriter_available", lambda: False)
    rep = _report("employee_productivity", {"from": "2026-09-01", "to": "2026-09-30"})
    assert rep.engine == "openpyxl" and rep.content[:2] == b"PK"
    wb = _wb(rep.content)
    assert wb.sheetnames[0] == "Tổng hợp" and len(wb.sheetnames) == 4
    ws = wb["Tổng hợp"]
    rows = _data_rows(ws)
    assert rows[-1][_col(ws, "Sản lượng đạt")] == 115
    evil = [c for row in ws.iter_rows() for c in row if c.value == EVIL]
    assert evil and all(c.data_type == "s" for c in evil)
    with zipfile.ZipFile(BytesIO(rep.content)) as z:
        assert not any("<f>" in z.read(n).decode() for n in z.namelist() if n.startswith("xl/worksheets/"))


def test_empty_result_still_produces_valid_files():
    repo = FakeRepo(employees=[], psessions=[], wsessions=[])
    wb = _wb(_report("employee_productivity", {"from": "2026-09-01", "to": "2026-09-02"}, repo=repo).content)
    assert wb.sheetnames == ["Tổng hợp"]
    rows = _data_rows(wb["Tổng hợp"])
    assert len(rows) == 1 and rows[0][_col(wb["Tổng hợp"], "Sản lượng đạt")] == 0
    html = _report("work_sessions", {"format": "html"}, repo=repo).content.decode()
    assert "Không có dữ liệu trong bộ lọc này." in html


def test_sheet_names_are_valid_and_unique():
    names = xlsx_module.sheet_names(["Tổng hợp", "a/b:c*d?[e]", "tổng HỢP", "x" * 40, "history", ""])
    assert names[0] == "Tổng hợp" and names[2] == "tổng HỢP (2)"
    assert all(len(n) <= 31 and not set(n) & set("[]:*?/\\") for n in names)
    assert names[4] != "history" and len({n.casefold() for n in names}) == len(names)


# ------------------------------------------------------------------ HTML / PDF
def test_html_layout_is_escaped_vietnamese_and_has_page_breaks():
    html = _report("employee_productivity", {"from": "2026-09-01", "to": "2026-09-30", "format": "html"}).content.decode()
    assert '<html lang="vi">' in html and "BÁO CÁO NĂNG SUẤT NHÂN VIÊN" in html and "Trần Thị Bình" in html
    assert '<img src="http' not in html and "&lt;img" in html  # report text cannot inject markup / fetches
    assert html.count('class="part page-break"') == 3  # one per employee, none before the summary
    assert "size: A4 landscape" in html and 'counter(page) "/" counter(pages)' in html
    assert "http://" not in html.split("<body>")[0]  # no external CSS/fonts


def test_pdf_fallback_when_weasyprint_or_native_libs_missing(monkeypatch):
    monkeypatch.setattr(pdf_module, "pdf_status", lambda: (False, "Thiếu thư viện hệ thống cho WeasyPrint (Pango/HarfBuzz): x"))
    with pytest.raises(pdf_module.PdfEngineUnavailable, match="Pango"):
        _report("employee_productivity", {"format": "pdf"})


def _weasy_ok():
    ok, reason = pdf_module.pdf_status()
    return ok, reason


@pytest.mark.skipif(not _weasy_ok()[0], reason=f"WeasyPrint unavailable here: {_weasy_ok()[1]}")
def test_pdf_is_real_vietnamese_a4_landscape_with_page_breaks():
    from pypdf import PdfReader

    rep = _report("employee_productivity", {"from": "2026-09-01", "to": "2026-09-30", "format": "pdf"})
    assert rep.engine == "weasyprint" and rep.content.startswith(b"%PDF-") and rep.content.rstrip().endswith(b"%%EOF")
    reader = PdfReader(BytesIO(rep.content))
    assert len(reader.pages) == 4  # summary + one page per employee
    box = reader.pages[0].mediabox
    assert float(box.width) > float(box.height) and round(float(box.width) / 72 * 25.4) == 297
    text = "\n".join(p.extract_text() for p in reader.pages)
    for needle in ("BÁO CÁO NĂNG SUẤT NHÂN VIÊN", "Nguyễn Văn An", "Trần Thị Bình", "Tổng cộng", "Trang 1/4", "Múi giờ"):
        assert needle in text, needle
    fonts = {str(f.get_object().get("/BaseFont")) for p in reader.pages
             for f in (p["/Resources"].get("/Font") or {}).values()}
    assert any("DejaVu" in f for f in fonts), fonts


@pytest.mark.skipif(not _weasy_ok()[0], reason=f"WeasyPrint unavailable here: {_weasy_ok()[1]}")
def test_pdf_never_fetches_urls_from_report_content(monkeypatch):
    seen = []
    real = pdf_module._deny_all_urls
    monkeypatch.setattr(pdf_module, "_deny_all_urls", lambda url, *a, **k: (seen.append(url), real(url))[1])
    rep = _report("work_sessions", {"format": "pdf"})
    assert rep.content.startswith(b"%PDF-") and seen == []


def test_docker_images_ship_pdf_native_dependencies():
    for name in ("Dockerfile", "Dockerfile.test"):
        text = (ROOT / name).read_text(encoding="utf-8")
        for pkg in ("libpango-1.0-0", "libpangoft2-1.0-0", "libharfbuzz-subset0", "fonts-dejavu-core"):
            assert pkg in text, (name, pkg)
    req = (ROOT / "requirements.txt").read_text(encoding="utf-8")
    assert "XlsxWriter==" in req and "weasyprint==" in req and "openpyxl==3.1.5" in req


# ------------------------------------------------------------------ HTTP / RBAC / isolation
SEED = set(rbac_module.SEED_ROLE_PERMISSIONS)


@pytest.fixture
def make_client(monkeypatch):
    state = {"logged_in": True}
    monkeypatch.setattr(auth_module, "validate_and_touch", lambda: None if state["logged_in"] else "NOT_LOGGED_IN")
    monkeypatch.setattr(rbac_module.RBACRepository, "has_permission",
                        lambda self, role, perm: role == "admin" or (role, perm) in SEED)
    monkeypatch.setattr(report_engine_module, "ReportRepository", FakeRepo)
    app = app_module.create_app()
    app.config["TESTING"] = True

    def make(role="manager", username="quanly", user_id=1):
        c = app.test_client()
        with c.session_transaction() as sess:
            sess["user_id"] = user_id
            sess["username"] = username
            sess["role"] = role
        c.state = state
        return c
    return make


BASE = "/api/reports/engine"
Q = "from=2026-09-01&to=2026-09-30"


def test_http_xlsx_download(make_client):
    r = make_client().get(f"{BASE}/employee_productivity/export?{Q}&format=xlsx&department=May")
    assert r.status_code == 200 and r.data[:2] == b"PK"
    assert r.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert r.headers["Content-Disposition"] == 'attachment; filename="mesflow_employee_productivity_2026-09-01_2026-09-30.xlsx"'
    assert r.headers["Cache-Control"] == "no-store, private" and r.headers["X-Report-Engine"] == "xlsxwriter"
    assert r.headers["X-Report-Rows"] == "2"


def test_http_pdf_503_with_html_fallback_when_engine_missing(make_client, monkeypatch):
    monkeypatch.setattr(pdf_module, "pdf_status", lambda: (False, "Chưa cài thư viện WeasyPrint"))
    c = make_client()
    r = c.get(f"{BASE}/work_sessions/export?{Q}&format=pdf&po_id=7")
    assert r.status_code == 503
    body = r.get_json()
    assert body["error"] == "PDF_ENGINE_UNAVAILABLE" and "format=html" in body["fallback_url"] and "po_id=7" in body["fallback_url"]
    h = c.get(body["fallback_url"])
    assert h.status_code == 200 and h.mimetype == "text/html" and h.headers["Content-Disposition"].startswith("inline")
    assert h.headers["Content-Type"] == "text/html; charset=utf-8"
    assert "BÁO CÁO PHIÊN LÀM VIỆC" in h.get_data(as_text=True)
    caps = c.get(f"{BASE}/datasets").get_json()["capabilities"]
    assert caps["pdf"]["available"] is False and caps["xlsx"]["available"] is True


@pytest.mark.skipif(not _weasy_ok()[0], reason=f"WeasyPrint unavailable here: {_weasy_ok()[1]}")
def test_http_pdf_download(make_client):
    r = make_client().get(f"{BASE}/work_sessions/export?{Q}&format=pdf")
    assert r.status_code == 200 and r.data.startswith(b"%PDF-") and r.mimetype == "application/pdf"
    assert r.headers["Content-Disposition"].startswith("attachment;")


def test_http_validation_and_not_found(make_client):
    c = make_client()
    assert c.get(f"{BASE}/employee_productivity/export?from=bad").status_code == 400
    assert c.get(f"{BASE}/employee_productivity/export?{Q}&po_id=1").status_code == 400
    assert c.get(f"{BASE}/nope/export?{Q}").status_code == 404
    # employee outside the department filter -> 404 (same rule as existing export)
    assert c.get(f"{BASE}/employee_productivity/export?{Q}&department=Kho&employee_id=1").status_code == 404


def test_http_anonymous_is_401(make_client):
    c = make_client()
    c.state["logged_in"] = False
    assert c.get(f"{BASE}/datasets").status_code == 401
    assert c.get(f"{BASE}/employee_productivity/export?{Q}").status_code == 401


@pytest.mark.parametrize("role, dataset, status", [
    ("viewer", "work_sessions", 403),       # has session.view but the screen is admin/manager/supervisor only
    ("operator", "work_sessions", 403),
    ("supervisor", "work_sessions", 200),
    ("manager", "work_sessions", 200),
    ("viewer", "employee_productivity", 200),
    ("guest", "employee_productivity", 403),  # unknown role: no grants
])
def test_http_rbac(make_client, role, dataset, status):
    r = make_client(role=role).get(f"{BASE}/{dataset}/export?{Q}")
    assert r.status_code == status
    if status == 403:
        assert r.get_json()["error"] == "FORBIDDEN"


def test_datasets_listing_only_shows_what_the_role_may_export(make_client):
    keys = lambda role: {d["key"] for d in make_client(role=role).get(f"{BASE}/datasets").get_json()["items"]}
    assert keys("manager") == {"employee_productivity", "work_sessions"}
    assert keys("viewer") == {"employee_productivity"}
    assert keys("guest") == set()


def test_each_user_gets_their_own_uncached_report(make_client):
    a = make_client(role="manager", username="quanly_a", user_id=1)
    b = make_client(role="supervisor", username="todo_b", user_id=2)
    ra = a.get(f"{BASE}/work_sessions/export?{Q}&format=html")
    rb = b.get(f"{BASE}/work_sessions/export?{Q}&format=html")
    assert "bởi quanly_a" in ra.get_data(as_text=True) and "todo_b" not in ra.get_data(as_text=True)
    assert "bởi todo_b" in rb.get_data(as_text=True) and "quanly_a" not in rb.get_data(as_text=True)
    assert ra.headers["Cache-Control"] == rb.headers["Cache-Control"] == "no-store, private"
    wa = _wb(a.get(f"{BASE}/work_sessions/export?{Q}").data)["Tổng hợp"]["A2"].value
    assert "bởi quanly_a" in wa


def test_existing_productivity_export_routes_are_untouched():
    src = (ROOT / "app/mesflow/web/analytics.py").read_text(encoding="utf-8")
    assert "@bp.get('/reports/employee-productivity/export.xlsx')" in src
    assert "build_employee_productivity_workbook(data,timezone_name=settings.timezone_name)" in src
    assert "xlsxwriter" not in (ROOT / "app/mesflow/web/productivity_excel.py").read_text(encoding="utf-8").lower()
    assert "xlsxwriter" not in (ROOT / "app/mesflow/web/router_export.py").read_text(encoding="utf-8").lower()


def test_engine_has_no_sql_of_its_own():
    for path in (ROOT / "app/mesflow/reporting").glob("*.py"):
        text = path.read_text(encoding="utf-8")
        for token in ("fetch_all", "fetch_one", "execute(", "SELECT ", "cursor"):
            assert token not in text, (path.name, token)

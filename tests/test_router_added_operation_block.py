"""Operation người dùng thêm sau khi nhập vẫn phải có tem, trong đúng tờ của Part.

BỐI CẢNH. Người dùng được phép thêm công đoạn vào Part/PO sau khi đã nhập
Template — đó là việc bình thường, không phải hỏng hóc. Nhưng file Excel gốc
chỉ chứa những block có lúc nhập, nên công đoạn mới không khớp block nào.

Bản trước coi MỌI Operation không khớp là lỗi cứng. Sai ở chỗ nó gộp hai
chuyện khác hẳn nhau:

  * Part của Operation CÓ tờ trong template -> người dùng thêm công đoạn, hợp
    lệ. Dựng thêm một block trong đúng tờ đó, sao từ block mẫu của chính tờ ấy.
  * Part KHÔNG có tờ nào -> ánh xạ hỏng, file đang lưu không phải file sinh ra
    PO này. Cái này mới đáng dừng.

Bài ở đây khoá: block cũ không xê dịch một ly, block mới kế thừa đúng hình dạng
của template (style, merge, chiều cao dòng, công thức đã dịch tham chiếu), tem
dán đúng Operation, và trường hợp không có mẫu an toàn thì dừng có lý do.
"""
import pytest
from io import BytesIO

from openpyxl import Workbook, load_workbook

from mesflow.web.router_export import (
    QR_MARKER_TEXT, RouterBlockTemplateError, _stamp_source_workbook)

pytest.importorskip('PIL', reason='Pillow là dependency của tính năng xuất QR')

PO = {'id': 1, 'code': 'PO-6126', 'product': 'GHẾ', 'planned_quantity': 110,
      'source_template_id': 7}
BLOCK_HEIGHT = 12
FIRST_BLOCK_ROW = 8


def _sheet(wb, title, drawing_code, blocks, *, marker_column=None, with_style=True):
    """Sheet Part có style/merge/công thức thật, để bài kiểm có gì mà so."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    ws = wb.create_sheet(title)
    ws['A2'] = 'PO NUMBER:'; ws['C2'] = 6126
    ws['A3'] = 'QTY:'; ws['C3'] = 110
    ws['A5'] = 'MÃ BẢN VẼ:'; ws['C5'] = drawing_code
    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['L'].width = 22
    edge = Side(style='thin')
    row = FIRST_BLOCK_ROW
    for index, (name, minutes, cycle) in enumerate(blocks, start=1):
        title_cell = ws.cell(row=row, column=1, value=f'OPERATION # {index:02d}- {name}')
        if with_style:
            title_cell.font = Font(bold=True, size=12)
            title_cell.alignment = Alignment(horizontal='center')
            title_cell.fill = PatternFill('solid', fgColor='DDEEFF')
            title_cell.border = Border(left=edge, right=edge, top=edge, bottom=edge)
        ws.cell(row=row + 1, column=1, value='Part Number ( Mã bản vẽ )')
        ws.cell(row=row + 1, column=2, value=drawing_code)
        ws.cell(row=row + 2, column=1, value='Ngày/Tháng/Năm')
        ws.cell(row=row + 2, column=12, value='Thời gian Setup ( phút )')
        ws.cell(row=row + 3, column=1, value='SETUP')
        ws.cell(row=row + 3, column=12, value=minutes)
        ws.cell(row=row + 5, column=1, value='Ngày/Tháng/Năm')
        ws.cell(row=row + 5, column=12, value='Thời gian gia công / sản phẩm (s)')
        ws.cell(row=row + 6, column=1, value='Thời gian gia công')
        ws.cell(row=row + 6, column=12, value=cycle)
        # Công thức THẬT: tổng giờ = giây/sp * số lượng / 3600. Khi nhân bản
        # block, tham chiếu L phải đi theo dòng mới.
        ws.cell(row=row + 6, column=13, value=f'=L{row + 6}*$C$3/3600')
        if marker_column:
            ws.cell(row=row + 6, column=marker_column, value=QR_MARKER_TEXT)
        ws.cell(row=row + 7, column=1, value='Hàng đạt:')
        ws.cell(row=row + 8, column=1, value='Nhân viên SX:')
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        ws.row_dimensions[row].height = 22
        row += BLOCK_HEIGHT
    return ws


def _workbook(sheets, **kwargs):
    wb = Workbook(); wb.remove(wb.active)
    for title, code, blocks in sheets:
        _sheet(wb, title, code, blocks, **kwargs)
    buffer = BytesIO(); wb.save(buffer)
    return buffer.getvalue()


def _row(op_id, code, name, *, part_code, part_name, setup_id=None, part_id=1,
         part_sort=0, sort_order=0, cycle=100):
    return {
        'id': op_id, 'code': code, 'name': name, 'sort_order': sort_order,
        'standard_seconds_per_unit': cycle, 'requires_setup': bool(setup_id),
        'expected_setup_minutes': 30 if setup_id else None,
        'part_id': part_id, 'part_code': part_code, 'part_name': part_name,
        'part_sort': part_sort, 'op_qr': f'WF|OPID|{op_id}',
        'setup_id': setup_id, 'setup_code': f'{code}-SU' if setup_id else None,
        'setup_name': f'Setup {name}' if setup_id else None,
        'setup_minutes': 30 if setup_id else None,
        'setup_qr': f'WF|OPID|{setup_id}' if setup_id else None,
    }


# Template có HAI block; PO có BA Operation -- cái thứ ba do người dùng thêm.
TWO_BLOCK_SHEET = [('Chân ghế A', 'KM-001',
                    [('CẮT LASER', 20, 100), ('LÀM NGUỘI', 0, 50)])]


def _three_rows(third_setup_id=None):
    return [
        _row(101, 'PO-6126-KM-001-OP01', 'CẮT LASER', part_code='KM-001',
             part_name='Chân ghế A'),
        _row(102, 'PO-6126-KM-001-OP02', 'LÀM NGUỘI', part_code='KM-001',
             part_name='Chân ghế A', sort_order=1),
        _row(103, 'PO-6126-KM-001-OP03', 'MÀI BAVIA', part_code='KM-001',
             part_name='Chân ghế A', sort_order=2, setup_id=third_setup_id, cycle=75),
    ]


def _export(source, rows):
    wb, placed, matched, unmatched = _stamp_source_workbook(source, PO, rows)
    buffer = BytesIO(); wb.save(buffer); buffer.seek(0)
    return load_workbook(buffer), placed, matched, unmatched


def _block_titles(ws):
    return [(c.row, c.value) for row in ws.iter_rows() for c in row
            if isinstance(c.value, str) and c.value.startswith('OPERATION #')]


# --- OP người dùng thêm được dựng block trong ĐÚNG tờ ---------------------

def test_op_nguoi_dung_them_duoc_dung_block_trong_cung_sheet():
    exported, placed, matched, unmatched = _export(_workbook(TWO_BLOCK_SHEET), _three_rows())
    assert matched == 3 and unmatched == []
    ws = exported['Chân ghế A']
    titles = _block_titles(ws)
    assert len(titles) == 3, f'phải có đúng 3 block, đang là {titles}'
    assert titles[0] == (8, 'OPERATION # 01- CẮT LASER')
    assert titles[1] == (20, 'OPERATION # 02- LÀM NGUỘI')
    # Block mới nối ngay sau block cuối, đánh số tiếp.
    assert titles[2] == (32, 'OPERATION # 03- MÀI BAVIA')
    # KHÔNG sinh sheet phụ nào.
    assert exported.sheetnames == ['Chân ghế A']
    assert {item['operation_id'] for item in placed} == {101, 102, 103}


def test_hai_block_cu_khong_xe_dich_mot_ly():
    source = _workbook(TWO_BLOCK_SHEET)
    exported, _, _, _ = _export(source, _three_rows())
    original = load_workbook(BytesIO(source))
    before, after = original['Chân ghế A'], exported['Chân ghế A']
    # Vùng của hai block gốc (dòng 1..31) phải y hệt.
    for row in range(1, 32):
        for cell in before[row]:
            assert after.cell(row=row, column=cell.column).value == cell.value, (
                f'dòng {row} cột {cell.column} bị đổi')
        assert after.row_dimensions[row].height == before.row_dimensions[row].height
    assert {k: v.width for k, v in before.column_dimensions.items()} == {
        k: v.width for k, v in after.column_dimensions.items()}
    # Merge của hai block gốc còn nguyên (block mới thêm merge của riêng nó).
    original_merges = {str(r) for r in before.merged_cells.ranges}
    assert original_merges <= {str(r) for r in after.merged_cells.ranges}


def test_block_moi_ke_thua_hinh_dang_cua_block_mau():
    """Style, merge, chiều cao dòng của block mới đến TỪ template, không tự chế."""
    source = _workbook(TWO_BLOCK_SHEET)
    exported, _, _, _ = _export(source, _three_rows())
    ws = exported['Chân ghế A']
    template_title, new_title = ws.cell(row=20, column=1), ws.cell(row=32, column=1)
    assert new_title.font.b == template_title.font.b
    assert new_title.font.sz == template_title.font.sz
    assert new_title.alignment.horizontal == template_title.alignment.horizontal
    assert new_title.fill.fgColor.rgb == template_title.fill.fgColor.rgb
    assert new_title.border.left.style == template_title.border.left.style
    assert ws.row_dimensions[32].height == ws.row_dimensions[20].height
    assert 'A32:C32' in {str(r) for r in ws.merged_cells.ranges}
    # Nhãn của block mẫu được mang theo nguyên văn.
    assert ws.cell(row=33, column=1).value == 'Part Number ( Mã bản vẽ )'
    assert ws.cell(row=34, column=12).value == 'Thời gian Setup ( phút )'
    assert ws.cell(row=37, column=12).value == 'Thời gian gia công / sản phẩm (s)'


def test_cong_thuc_trong_block_moi_duoc_dich_tham_chieu():
    """=L26*$C$3/3600 ở block mẫu phải thành =L38*$C$3/3600 ở block mới.

    Không dịch thì block mới hiển thị số giờ của block cũ -- sai lặng lẽ, và
    người đọc tờ giấy không có cách nào biết.
    """
    exported, _, _, _ = _export(_workbook(TWO_BLOCK_SHEET), _three_rows())
    ws = exported['Chân ghế A']
    assert ws.cell(row=26, column=13).value == '=L26*$C$3/3600'
    assert ws.cell(row=38, column=13).value == '=L38*$C$3/3600'


def test_du_lieu_cua_op_moi_duoc_ghi_vao_dung_o():
    exported, _, _, _ = _export(_workbook(TWO_BLOCK_SHEET), _three_rows())
    ws = exported['Chân ghế A']
    assert ws.cell(row=38, column=12).value == 75        # thời gian gia công
    assert ws.cell(row=35, column=12).value == 0         # không setup -> 0
    assert ws.cell(row=33, column=2).value == 'KM-001'   # part number


def test_op_moi_co_setup_thi_co_ca_hai_tem():
    exported, placed, _, _ = _export(_workbook(TWO_BLOCK_SHEET), _three_rows(third_setup_id=203))
    ws = exported['Chân ghế A']
    assert ws.cell(row=35, column=12).value == 30        # phút setup của OP mới
    kinds = {(item['operation_id'], item['kind']) for item in placed}
    assert (103, 'OP') in kinds and (203, 'SETUP') in kinds
    on_new_block = [item for item in placed if item['operation_id'] in (103, 203)]
    assert {item['sheet'] for item in on_new_block} == {'Chân ghế A'}
    assert len(placed) == 4


def test_tem_cua_block_moi_tro_dung_operation_id():
    _, placed, _, _ = _export(_workbook(TWO_BLOCK_SHEET), _three_rows(third_setup_id=203))
    payloads = {item['operation_id']: item['payload'] for item in placed}
    for operation_id in (101, 102, 103, 203):
        assert payloads[operation_id] == f'WF|OPID|{operation_id}'
    assert len(set(payloads.values())) == len(payloads)


def test_marker_trong_block_mau_duoc_nhan_ban_va_dan_dung():
    """Template có marker -> block mới cũng có marker, và tem đi vào đó."""
    source = _workbook(TWO_BLOCK_SHEET, marker_column=14)
    exported, placed, _, _ = _export(source, _three_rows())
    items = {item['operation_id']: item for item in placed}
    assert items[101]['marker_cell'] == 'N14'
    assert items[102]['marker_cell'] == 'N26'
    # Block mới ở dòng 32 -> marker sao chép nằm ở dòng 38.
    assert items[103]['placement'] == 'marker'
    assert items[103]['marker_cell'] == 'N38'
    values = {c.value for row in exported['Chân ghế A'].iter_rows() for c in row}
    assert QR_MARKER_TEXT not in values


def test_nhieu_part_moi_part_them_op_vao_dung_to_cua_minh():
    sheets = [('Chân ghế A', 'KM-001', [('CẮT LASER', 20, 100)]),
              ('Chân ghế B', 'KM-002', [('HÀN', 0, 60)])]
    rows = [
        _row(101, 'PO-6126-KM-001-OP01', 'CẮT LASER', part_code='KM-001',
             part_name='Chân ghế A'),
        _row(102, 'PO-6126-KM-001-OP02', 'MÀI', part_code='KM-001',
             part_name='Chân ghế A', sort_order=1),
        _row(201, 'PO-6126-KM-002-OP01', 'HÀN', part_code='KM-002',
             part_name='Chân ghế B', part_id=2, part_sort=1),
        _row(202, 'PO-6126-KM-002-OP02', 'KIỂM TRA', part_code='KM-002',
             part_name='Chân ghế B', part_id=2, part_sort=1, sort_order=1),
    ]
    exported, placed, matched, unmatched = _export(_workbook(sheets), rows)
    assert matched == 4 and unmatched == []
    assert exported.sheetnames == ['Chân ghế A', 'Chân ghế B']
    assert len(_block_titles(exported['Chân ghế A'])) == 2
    assert len(_block_titles(exported['Chân ghế B'])) == 2
    sheet_of = {item['operation_id']: item['sheet'] for item in placed}
    assert sheet_of[102] == 'Chân ghế A'      # OP mới của Part A ở tờ A
    assert sheet_of[202] == 'Chân ghế B'      # OP mới của Part B ở tờ B


def test_print_area_duoc_noi_de_block_moi_in_duoc():
    wb = Workbook(); wb.remove(wb.active)
    _sheet(wb, 'Chân ghế A', 'KM-001', [('CẮT LASER', 20, 100), ('LÀM NGUỘI', 0, 50)])
    wb['Chân ghế A'].print_area = 'A1:M31'
    buffer = BytesIO(); wb.save(buffer)
    exported, _, _, _ = _export(buffer.getvalue(), _three_rows())
    area = exported['Chân ghế A'].print_area
    area = area[0] if isinstance(area, (list, tuple)) else area
    assert area.endswith('$43') or area.endswith('43'), f'print area chưa nới: {area}'


# --- ánh xạ hỏng thì vẫn phải DỪNG ---------------------------------------

def test_part_khong_co_to_nao_thi_van_la_loi_anh_xa():
    """Không phải mọi 'không khớp' đều là OP người dùng thêm."""
    from mesflow.web.router_export import _stamp_source_workbook as stamp
    rows = [_row(101, 'PO-6126-KM-001-OP01', 'CẮT LASER', part_code='KM-001',
                 part_name='Chân ghế A'),
            _row(999, 'PO-6126-KM-404-OP01', 'CÔNG ĐOẠN LẠ', part_code='KM-404',
                 part_name='Part không có trong file', part_id=9, part_sort=9)]
    _, _, matched, unmatched = stamp(_workbook(TWO_BLOCK_SHEET), PO, rows)
    assert matched == 1
    assert [row['code'] for row in unmatched] == ['PO-6126-KM-404-OP01']


def test_to_khong_co_block_mau_thi_bao_loi_ro_rang():
    """Tờ tồn tại nhưng không có block nào -> không có mẫu để sao, phải dừng."""
    wb = Workbook(); wb.remove(wb.active)
    ws = wb.create_sheet('Chân ghế A')
    ws['A5'] = 'MÃ BẢN VẼ:'; ws['C5'] = 'KM-001'
    ws['A8'] = 'OPERATION # 01- CẮT LASER'
    second = wb.create_sheet('Chân ghế B')
    second['A5'] = 'MÃ BẢN VẼ:'; second['C5'] = 'KM-002'
    buffer = BytesIO(); wb.save(buffer)
    rows = [_row(101, 'PO-6126-KM-001-OP01', 'CẮT LASER', part_code='KM-001',
                 part_name='Chân ghế A')]
    # Part KM-002 không có block nào -> nó cũng không có trong sheet_of_part,
    # nên Operation của nó rơi vào nhánh ánh xạ hỏng chứ không phải dựng block.
    rows.append(_row(201, 'PO-6126-KM-002-OP01', 'HÀN', part_code='KM-002',
                     part_name='Chân ghế B', part_id=2, part_sort=1))
    _, _, matched, unmatched = _stamp_source_workbook(buffer.getvalue(), PO, rows)
    assert matched == 1
    assert [row['code'] for row in unmatched] == ['PO-6126-KM-002-OP01']


def test_loi_block_mau_mang_du_sheet_part_operation():
    """RouterBlockTemplateError phải nói đủ sheet + Part + Operation."""
    error = RouterBlockTemplateError('x', sheet='Chân ghế A', part='KM-001',
                                     operation='OP03', reason='NO_TEMPLATE_BLOCK')
    assert error.sheet == 'Chân ghế A'
    assert error.part == 'KM-001'
    assert error.operation == 'OP03'
    assert error.reason == 'NO_TEMPLATE_BLOCK'


def test_template_khong_co_op_moi_thi_khong_them_block_nao():
    """Hồi quy: PO khớp hết template thì tuyệt đối không sinh thêm block."""
    rows = _three_rows()[:2]
    source = _workbook(TWO_BLOCK_SHEET)
    exported, placed, matched, unmatched = _export(source, rows)
    assert matched == 2 and unmatched == []
    assert len(_block_titles(exported['Chân ghế A'])) == 2
    original = load_workbook(BytesIO(source))
    assert exported['Chân ghế A'].max_row <= original['Chân ghế A'].max_row + 1

"""Tổng thời gian gia công dự kiến: số của khách, số MESFlow tính, và độ lệch.

CÔNG THỨC KHÔNG PHẢI DO TA ĐẶT RA. Nó đọc được nguyên văn từ file thật của
khách (`tests/fixtures/router-newark-arm-chair.xlsx`, mở bằng openpyxl với
``data_only=False``), block đầu của sheet 'Tay ghế - Trái':

    M14 = $I$4*L14/3600 + L11/60          (I4=110 sp, L14=70 giây/sp, L11=15 phút)
    -> 2,3889 giờ = 8.600 giây

nên: ``total_seconds = qty * cycle_giây + setup_phút * 60``. Setup cộng MỘT LẦN
cho cả lô, qty nhân đúng MỘT LẦN, và ba thành phần có ba đơn vị khác nhau ngay
trong một ô (giây / phút / giờ). Cả ba đều là chỗ dễ sai nên có bài riêng.

ĐỘ LỆCH TRÊN FILE THẬT LÀ CÓ THẬT, VÀ CÓ Ý NGHĨA. Đo trên 44 sheet / 112 block:

    Excel   1.111.110 s = 308,64 h
    MESFlow 1.006.610 s = 279,61 h        lệch -104.500 s = -29,03 h (-9,4%)

107/112 block khớp CHÍNH XÁC. 5 block lệch đều vì công thức Excel của riêng
chúng mang thêm một hệ số số lượng chi tiết viết THẲNG TRONG CÔNG THỨC, không
nằm trong ô nào:

    'Lắp ráp sau khi sơn'!M14 = $I$4*L14*10/3600 + L11/60
    'Đóng gói'!M14            = $I$4*L14*2/3600  + L11/60

Trình đọc file chỉ thấy giá trị ô, nên không thể lấy được hệ số đó. Vì vậy bài
này KHÔNG khẳng định hai số bằng nhau -- nó khoá lại đúng con số đo được, để một
thay đổi sau này ở parser hay ở công thức làm xê dịch nó thì phải giải thích.
"""
from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path

import pytest

pytest.importorskip('openpyxl')
from openpyxl import load_workbook  # noqa: E402

from mesflow.domain import expected_time  # noqa: E402
from mesflow.web.excel_io import _parse_go_router_template  # noqa: E402

FIXTURE = Path(__file__).resolve().parent / 'fixtures/router-newark-arm-chair.xlsx'

#: Đo được trên chính file thật. Xem docstring.
NEWARK_SOURCE_SECONDS = 1_111_110.0
NEWARK_CALCULATED_SECONDS = 1_006_610.0
NEWARK_BLOCKS = 112


# --------------------------------------------------------------------------
# 1. Công thức
# --------------------------------------------------------------------------
def test_the_formula_matches_the_customers_own_cell():
    """110 sp x 70 giây + 15 phút -> đúng 2,3889 giờ như ô M14 đã cache."""
    seconds = expected_time.operation_expected_seconds(110, 70, 15)
    assert seconds == pytest.approx(8600.0)
    assert seconds / 3600 == pytest.approx(2.388888888888889)


def test_setup_is_added_once_not_per_unit():
    """Sai kinh điển: nhân setup với sản lượng. 110 x 15 phút = 27,5 giờ setup."""
    assert expected_time.operation_expected_seconds(110, 0, 15) == pytest.approx(900.0)


def test_quantity_multiplies_exactly_once():
    """Sai kinh điển thứ hai: nhân qty hai lần (qty*qty*cycle)."""
    assert expected_time.operation_expected_seconds(10, 5, 0) == pytest.approx(50.0)


def test_a_missing_setup_is_zero_not_an_error():
    """Phần lớn công đoạn không khai setup; None phải là 0, không phải nổ."""
    assert expected_time.operation_expected_seconds(10, 5, None) == pytest.approx(50.0)
    assert expected_time.operation_expected_seconds(None, None, None) == 0.0


# --------------------------------------------------------------------------
# 2. Gộp theo Template
# --------------------------------------------------------------------------
def _summary(parts, operations):
    return expected_time.summarize(parts, operations)


def test_quantity_comes_from_the_part_not_from_one_global_number():
    """Biểu mẫu ghi SỐ LƯỢNG trên TỪNG TỜ; một Part dùng nhiều lần có số riêng."""
    parts = [{'id': 1, 'planned_quantity': 10}, {'id': 2, 'planned_quantity': 100}]
    ops = [{'part_id': 1, 'standard_seconds_per_unit': 60, 'expected_setup_minutes': None},
           {'part_id': 2, 'standard_seconds_per_unit': 60, 'expected_setup_minutes': None}]
    assert _summary(parts, ops)['calculated_seconds'] == pytest.approx(10 * 60 + 100 * 60)


def test_no_source_total_still_reports_the_calculated_one():
    """Template cũ (hoặc dựng tay trên web) không có ô tổng -- vẫn phải có số."""
    parts = [{'id': 1, 'planned_quantity': 10}]
    ops = [{'part_id': 1, 'standard_seconds_per_unit': 60, 'expected_setup_minutes': 5}]
    out = _summary(parts, ops)
    assert out['has_source'] is False
    assert out['source_seconds'] is None and out['delta_seconds'] is None
    assert out['mismatch'] is False
    assert out['calculated_seconds'] == pytest.approx(10 * 60 + 300)


def test_matching_totals_do_not_raise_a_warning():
    parts = [{'id': 1, 'planned_quantity': 10}]
    ops = [{'part_id': 1, 'standard_seconds_per_unit': 60, 'expected_setup_minutes': None,
            'expected_total_seconds': 600.0}]
    out = _summary(parts, ops)
    assert out['has_source'] is True and out['mismatch'] is False
    assert out['delta_seconds'] == pytest.approx(0.0)


def test_a_bom_multiplier_shows_up_as_a_mismatch():
    """Đúng hình dạng của 5 block lệch trên file thật: Excel gấp 10 lần."""
    parts = [{'id': 1, 'planned_quantity': 110}]
    ops = [{'part_id': 1, 'standard_seconds_per_unit': 40, 'expected_setup_minutes': None,
            'expected_total_seconds': 44000.0}]          # = 110*40*10
    out = _summary(parts, ops)
    assert out['mismatch'] is True
    assert out['calculated_seconds'] == pytest.approx(4400.0)
    assert out['delta_seconds'] == pytest.approx(-39600.0)


def test_rounding_noise_is_not_reported_as_a_mismatch():
    """Ô tổng của khách là số thực; vài giây lệch không phải sai dữ liệu."""
    parts = [{'id': 1, 'planned_quantity': 110}]
    ops = [{'part_id': 1, 'standard_seconds_per_unit': 70, 'expected_setup_minutes': 15,
            'expected_total_seconds': 8600.4}]
    assert _summary(parts, ops)['mismatch'] is False


def test_partial_source_coverage_is_stated_not_hidden():
    """Nửa nạc nửa mỡ là ca dễ hiểu nhầm nhất -- phải nói rõ phủ bao nhiêu."""
    parts = [{'id': 1, 'planned_quantity': 10}]
    ops = [{'part_id': 1, 'standard_seconds_per_unit': 60, 'expected_setup_minutes': None,
            'expected_total_seconds': 600.0},
           {'part_id': 1, 'standard_seconds_per_unit': 60, 'expected_setup_minutes': None}]
    out = _summary(parts, ops)
    assert out['source_operation_count'] == 1 and out['operation_count'] == 2
    assert out['source_is_partial'] is True


# --------------------------------------------------------------------------
# 3. File THẬT của khách
# --------------------------------------------------------------------------
@pytest.fixture(scope='module')
def newark():
    if not FIXTURE.is_file():
        pytest.skip(f'thiếu fixture {FIXTURE}')
    return _parse_go_router_template(
        load_workbook(BytesIO(FIXTURE.read_bytes()), data_only=True), FIXTURE.name)


def test_every_block_of_the_real_form_carries_a_source_total(newark):
    ops = newark['operations']
    assert len(ops) == NEWARK_BLOCKS
    assert sum(1 for o in ops if o.get('total_expected_seconds')) == NEWARK_BLOCKS


def test_the_real_form_totals_are_the_measured_numbers(newark):
    """Khoá lại cả hai con số VÀ độ lệch giữa chúng.

    Không khẳng định chúng bằng nhau: chúng không bằng nhau, và lý do đã biết
    (hệ số BOM nằm trong công thức Excel). Bài này bắt mọi thay đổi làm xê dịch
    một trong hai số phải được giải thích.
    """
    parts = [{'id': p['key'], 'planned_quantity': p['planned_quantity']}
             for p in newark['parts']]
    ops = [{'part_id': o['part_key'],
            'standard_seconds_per_unit': o.get('standard_seconds_per_unit'),
            'expected_setup_minutes': o.get('expected_setup_minutes'),
            'expected_total_seconds': o.get('total_expected_seconds')}
           for o in newark['operations']]
    out = expected_time.summarize(parts, ops)

    assert out['source_seconds'] == pytest.approx(NEWARK_SOURCE_SECONDS, abs=1.0)
    assert out['calculated_seconds'] == pytest.approx(NEWARK_CALCULATED_SECONDS, abs=1.0)
    assert out['delta_seconds'] == pytest.approx(
        NEWARK_CALCULATED_SECONDS - NEWARK_SOURCE_SECONDS, abs=1.0)
    assert out['mismatch'] is True
    assert out['source_is_partial'] is False


def test_exactly_five_blocks_disagree_and_all_by_an_integer_factor(newark):
    """Chốt lại chẩn đoán: chỗ lệch là HỆ SỐ NGUYÊN, không phải nhiễu lan man.

    Nếu một ngày con số này đổi, nguyên nhân là parser hoặc công thức đã đổi --
    không phải file khách.
    """
    quantity = {p['key']: p['planned_quantity'] for p in newark['parts']}
    off = []
    for o in newark['operations']:
        source = float(o['total_expected_seconds'])
        calc = expected_time.operation_expected_seconds(
            quantity.get(o['part_key']), o.get('standard_seconds_per_unit'),
            o.get('expected_setup_minutes'))
        if abs(calc - source) > 1.0:
            off.append((o['code'], round(source / calc, 6) if calc else None))
    assert len(off) == 5, off
    assert sorted({factor for _, factor in off}) == [2.0, 4.0, 10.0], off


def test_the_mismatch_names_the_operations_and_their_implied_factor(newark):
    """Chênh 29 giờ là một con số; "5 công đoạn hệ số x2/x4/x10" là một việc làm được."""
    parts = [{'id': p['key'], 'planned_quantity': p['planned_quantity']}
             for p in newark['parts']]
    ops = [{'part_id': o['part_key'], 'code': o['code'], 'name': o['name'],
            'standard_seconds_per_unit': o.get('standard_seconds_per_unit'),
            'expected_setup_minutes': o.get('expected_setup_minutes'),
            'expected_total_seconds': o.get('total_expected_seconds')}
           for o in newark['operations']]
    out = expected_time.summarize(parts, ops)
    assert out['mismatch_operation_count'] == 5
    factors = sorted({o['implied_multiplier'] for o in out['mismatch_operations']})
    assert factors == [2, 4, 10], out['mismatch_operations']
    # ...và mỗi mục phải đủ để đi tìm đúng dòng trên giấy.
    for item in out['mismatch_operations']:
        assert item['code'] and item['source_seconds'] > item['calculated_seconds']


def test_a_non_integer_ratio_reports_no_multiplier():
    """Chỉ suy ra hệ số khi tỉ lệ SẠCH. Lệch lung tung là chuyện khác, đừng đoán bừa."""
    parts = [{'id': 1, 'planned_quantity': 100}]
    ops = [{'part_id': 1, 'code': 'OP01', 'standard_seconds_per_unit': 10,
            'expected_setup_minutes': None, 'expected_total_seconds': 2345.0}]
    out = expected_time.summarize(parts, ops)
    assert out['mismatch_operation_count'] == 1
    assert out['mismatch_operations'][0]['implied_multiplier'] is None


# --------------------------------------------------------------------------
# 4. VỊ TRÍ: đủ để mở file ra là tới thẳng chỗ lệch
# --------------------------------------------------------------------------
def _newark_with_formulas():
    if not FIXTURE.is_file():
        pytest.skip(f'thiếu fixture {FIXTURE}')
    raw = FIXTURE.read_bytes()
    return _parse_go_router_template(
        load_workbook(BytesIO(raw), data_only=True), FIXTURE.name,
        formula_workbook=load_workbook(BytesIO(raw), data_only=False))


def test_every_block_records_where_it_came_from():
    """Sheet + khoảng dòng + ô tổng, cho MỌI block -- không chỉ block lệch."""
    for op in _newark_with_formulas()['operations']:
        assert op['_excel_sheet'], op['code']
        assert op['_excel_row'] and op['_excel_row_end'], op['code']
        assert op['_excel_row_end'] >= op['_excel_row'], op['code']
        assert re.fullmatch(r'[A-Z]+\d+', op['total_expected_cell'] or ''), op['code']


def test_the_first_block_points_at_the_exact_cells_read_by_hand():
    """Chốt bằng toạ độ đã đọc tay từ file: L11 setup, L14 cycle, M14 tổng,
    block chiếm dòng 8-19 của sheet 'Tay ghế - Trái'."""
    ops = _newark_with_formulas()['operations']
    first = next(o for o in ops if o['_excel_sheet'] == 'Tay ghế - Trái'
                 and o['_excel_row'] == 8)
    assert (first['setup_cell'], first['cycle_cell'], first['total_expected_cell']) == (
        'L11', 'L14', 'M14')
    assert first['_excel_row_end'] == 19


def test_the_original_excel_formula_is_captured():
    """Công thức là thứ DUY NHẤT giải thích được hệ số ẩn -- nó không nằm trong
    ô nào khác, nên không đọc được nó thì chỉ còn cách đoán."""
    ops = _newark_with_formulas()['operations']
    first = next(o for o in ops if o['_excel_sheet'] == 'Tay ghế - Trái'
                 and o['_excel_row'] == 8)
    assert first['total_expected_formula'] == '=$I$4*L14/3600+L11/60'


def test_the_five_mismatches_are_fully_locatable():
    """Yêu cầu của người dùng, từng mục một: mỗi dòng lệch phải có sheet, dòng,
    ô, mã, tên, công thức gốc, hệ số ẩn, hai con số và độ chênh."""
    parsed = _newark_with_formulas()
    quantity = {p['key']: p['planned_quantity'] for p in parsed['parts']}
    parts = [{'id': p['key'], 'planned_quantity': p['planned_quantity']}
             for p in parsed['parts']]
    ops = [{'part_id': o['part_key'], 'code': o['code'], 'name': o['name'],
            'standard_seconds_per_unit': o.get('standard_seconds_per_unit'),
            'expected_setup_minutes': o.get('expected_setup_minutes'),
            'expected_total_seconds': o.get('total_expected_seconds'),
            'source_sheet': o.get('_excel_sheet'),
            'source_row_start': o.get('_excel_row'),
            'source_row_end': o.get('_excel_row_end'),
            'expected_total_cell': o.get('total_expected_cell'),
            'expected_total_formula': o.get('total_expected_formula')}
           for o in parsed['operations']]

    out = expected_time.summarize(parts, ops)
    assert out['mismatch_operation_count'] == 5
    for item in out['mismatch_operations']:
        assert item['sheet'], item
        assert item['row_start'] and item['row_end'], item
        assert re.fullmatch(r'[A-Z]+\d+', item['cell'] or ''), item
        assert item['code'] and item['name'], item
        assert item['formula'] and item['formula'].startswith('='), item
        assert item['implied_multiplier'] in (2, 4, 10), item
        assert item['source_seconds'] > item['calculated_seconds']
        assert item['delta_seconds'] == pytest.approx(
            item['calculated_seconds'] - item['source_seconds'])
        # Hệ số ẩn phải xuất hiện NGUYÊN VĂN trong công thức -- đó là bằng chứng
        # cho chẩn đoán, không phải một con số ta bịa ra từ tỉ lệ.
        assert f"*{item['implied_multiplier']}/" in item['formula'].replace(' ', ''), item

    # Và đúng hai tờ đó, không phải rải rác khắp nơi.
    assert {i['sheet'] for i in out['mismatch_operations']} == {
        'Lắp ráp sau khi sơn', 'Đóng gói'}


def test_a_template_imported_before_the_location_columns_still_works():
    """Tương thích ngược: Template cũ không có vị trí -- vẫn phải ra đủ hai con
    số và độ lệch, chỉ thiếu phần chỉ đường."""
    parts = [{'id': 1, 'planned_quantity': 110}]
    ops = [{'part_id': 1, 'code': 'OP01', 'name': 'CŨ',
            'standard_seconds_per_unit': 40, 'expected_setup_minutes': None,
            'expected_total_seconds': 44000.0}]
    out = expected_time.summarize(parts, ops)
    assert out['mismatch'] is True
    item = out['mismatch_operations'][0]
    assert item['sheet'] is None and item['cell'] is None and item['formula'] is None
    assert item['implied_multiplier'] == 10

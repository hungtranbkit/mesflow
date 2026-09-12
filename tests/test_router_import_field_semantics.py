"""Parser Router phải đọc HẾT ý nghĩa của file, không chỉ tên Part/OP.

Mọi field trong tờ router là dữ liệu nghiệp vụ. Bài này khoá ba nhóm:

  * ĐƠN VỊ THỜI GIAN đọc từ chính nhãn (giây/phút/giờ + biến thể), không đóng
    cứng — đóng cứng là cách chắc chắn nhất để một ngày nào đó nhập sai gấp 60
    lần mà không ai thấy;
  * SỐ LIỆU NGUỒN giữ nguyên: số OP gốc, tiêu đề gốc, `-` khác `0`, tổng thời
    gian dự kiến là field độc lập chứ không phải thứ suy ra rồi vứt;
  * SỐ LƯỢNG của sheet độc lập với QTY của PO (bội số BOM 1×/2×/4×).

Phần chạy trên chính file 44 sheet của xưởng nằm ở
`tests/test_router_real_workbook_semantics.py`.
"""
import pytest
from io import BytesIO

from openpyxl import Workbook, load_workbook

from mesflow.web.excel_io import _parse_go_router_template, _time_unit_from_label

BLOCK_HEIGHT = 12
FIRST_BLOCK_ROW = 8


def _sheet(wb, title, drawing_code, blocks, *, sheet_qty=110,
           setup_label='Thời gian Setup ( phút )',
           cycle_label='Thời gian gia công / sản phẩm (s)',
           total_label='Tổng thời gian gia công dự kiến ( giờ )',
           drawing_name='__same_as_title__', order_type='SẢN XUẤT HÀNG LOẠT'):
    """Sheet Part đúng bố cục thật: nhãn trái->phải, nhãn phải->xuống dưới."""
    ws = wb.create_sheet(title)
    ws['A2'] = 'PO NUMBER:'; ws['C2'] = 6126
    ws['G2'] = 'LOẠI ĐƠN HÀNG'; ws['I2'] = 'SỐ LƯỢNG'; ws['J2'] = 'HÌNH ẢNH'
    ws['A3'] = 'QTY:'; ws['C3'] = 110
    ws['A4'] = 'TÊN BẢN VẼ:'
    # Tờ bản vẽ thật khai CẢ tên lẫn mã; bài nào muốn tờ mức quy trình thì
    # truyền drawing_name=None cùng drawing_code=None.
    if drawing_name == '__same_as_title__':
        drawing_name = title
    if drawing_name is not None:
        ws['C4'] = drawing_name
    ws['G4'] = order_type
    ws['I4'] = sheet_qty
    ws['A5'] = 'MÃ BẢN VẼ:'
    if drawing_code is not None:
        ws['C5'] = drawing_code
    ws['J1'] = 'Mã số: SX-1/BM1\nLần ban hành: 1\nNgày BH: 7/9/2026'
    row = FIRST_BLOCK_ROW
    for index, spec in enumerate(blocks, start=1):
        name, setup_value, cycle_value, total_value = spec[:4]
        sequence = spec[4] if len(spec) > 4 else index
        part_number = spec[5] if len(spec) > 5 else drawing_code
        ws.cell(row=row, column=1, value=f'OPERATION # {sequence:02d}- {name}')
        ws.cell(row=row + 1, column=1, value='Part Number ( Mã bản vẽ )')
        if part_number is not None:
            ws.cell(row=row + 1, column=2, value=part_number)
        ws.cell(row=row + 2, column=1, value='Ngày/Tháng/Năm')
        ws.cell(row=row + 2, column=12, value=setup_label)
        ws.cell(row=row + 3, column=1, value='SETUP')
        ws.cell(row=row + 3, column=12, value=setup_value)
        ws.cell(row=row + 5, column=12, value=cycle_label)
        ws.cell(row=row + 5, column=13, value=total_label)
        ws.cell(row=row + 6, column=1, value='Thời gian gia công')
        ws.cell(row=row + 6, column=12, value=cycle_value)
        ws.cell(row=row + 6, column=13, value=total_value)
        row += BLOCK_HEIGHT
    return ws


def _parse(sheets, **kwargs):
    wb = Workbook(); wb.remove(wb.active)
    for title, code, blocks in sheets:
        _sheet(wb, title, code, blocks, **kwargs)
    buffer = BytesIO(); wb.save(buffer); buffer.seek(0)
    return _parse_go_router_template(load_workbook(buffer, data_only=True), 'router.xlsx')


def _ops(parsed):
    return {op['source_title']: op for op in parsed['operations']}


# --- đơn vị đọc từ nhãn ---------------------------------------------------

@pytest.mark.parametrize('label,unit', [
    ('Thời gian gia công / sản phẩm (s)', 'second'),
    ('Thời gian gia công / sản phẩm (sec)', 'second'),
    ('Thời gian gia công ( giây )', 'second'),
    ('Cycle time (seconds)', 'second'),
    ('Thời gian Setup ( phút )', 'minute'),
    ('Thời gian Setup (phút)', 'minute'),
    ('Setup (min)', 'minute'),
    ('Setup time (minutes)', 'minute'),
    ('Tổng thời gian gia công dự kiến ( giờ )', 'hour'),
    ('Tổng thời gian ( h )', 'hour'),
    ('Total (hr)', 'hour'),
    ('Total time (hours)', 'hour'),
])
def test_don_vi_doc_duoc_tu_nhan(label, unit):
    assert _time_unit_from_label(label) == unit


@pytest.mark.parametrize('label', [
    'Thời gian gia công', 'Tổng cộng', 'Thời gian Setup', 'Cycle time'])
def test_nhan_khong_noi_don_vi_thi_khong_doan(label):
    assert _time_unit_from_label(label) is None


def test_nhan_khong_ro_don_vi_bi_chan_kem_vi_tri():
    """Không đoán: thà dừng còn hơn nhập sai gấp 60 lần mà không ai thấy."""
    with pytest.raises(ValueError) as excinfo:
        _parse([('Chân ghế A', 'KM-001', [('CẮT LASER', 20, 100, 3.39)])],
               cycle_label='Thời gian gia công / sản phẩm')
    message = str(excinfo.value)
    assert 'không nói rõ đơn vị' in message
    assert 'Chân ghế A' in message


@pytest.mark.parametrize('setup_label,setup_value,expected_seconds', [
    ('Thời gian Setup ( phút )', 60, 3600.0),
    ('Thời gian Setup ( giây )', 60, 60.0),
    ('Thời gian Setup ( giờ )', 1, 3600.0),
    ('Setup (min)', 30, 1800.0),
])
def test_setup_quy_doi_theo_dung_don_vi_cua_nhan(setup_label, setup_value, expected_seconds):
    parsed = _parse([('Chân ghế A', 'KM-001', [('CẮT LASER', setup_value, 100, 3.39)])],
                    setup_label=setup_label)
    op = parsed['operations'][0]
    assert op['setup_seconds'] == expected_seconds
    assert op['setup_raw'] == setup_value
    # Cột schema là PHÚT (migration 0047), nên quy về phút khi lưu.
    assert op['expected_setup_minutes'] == int(round(expected_seconds / 60))


# --- ví dụ bắt buộc trong hợp đồng ---------------------------------------

def test_vi_du_chot_60_phut_180_giay_6_5_gio():
    """setup 60 phút + 110 × 180 giây = 6,5 giờ — đúng phép tính của tờ giấy."""
    parsed = _parse([('Chân ghế A', 'KM-001', [('CẮT LASER', 60, 180, 6.5)])])
    op = parsed['operations'][0]
    assert op['setup_raw'] == 60 and op['setup_unit'] == 'minute'
    assert op['setup_seconds'] == 3600.0
    assert op['expected_setup_minutes'] == 60
    assert op['requires_setup'] is True
    assert op['cycle_raw'] == 180 and op['cycle_unit'] == 'second'
    assert op['standard_seconds_per_unit'] == 180.0
    assert op['total_expected_raw'] == 6.5 and op['total_expected_unit'] == 'hour'
    assert op['total_expected_seconds'] == 23400.0
    # Và phép đối chiếu khớp: 3600 + 110*180 = 23400.
    assert 3600 + 110 * 180 == op['total_expected_seconds']


# --- 0 / '-' / trống là ba trạng thái KHÁC NHAU --------------------------

def test_setup_0_va_gach_ngang_va_trong_khac_nhau():
    parsed = _parse([('Chân ghế A', 'KM-001', [
        ('CÓ SETUP', 20, 100, 3.39),
        ('SETUP BẰNG 0', 0, 50, 1.53),
        ('SETUP GẠCH NGANG', '-', 50, 1.53),
        ('SETUP TRỐNG', None, 50, 1.53),
    ])])
    ops = _ops(parsed)
    has = ops['OPERATION # 01- CÓ SETUP']
    zero = ops['OPERATION # 02- SETUP BẰNG 0']
    dash = ops['OPERATION # 03- SETUP GẠCH NGANG']
    blank = ops['OPERATION # 04- SETUP TRỐNG']
    # Chỉ >0 mới sinh OP SETUP.
    assert has['requires_setup'] is True
    assert (zero['requires_setup'], dash['requires_setup'], blank['requires_setup']) == (
        False, False, False)
    # Nhưng trạng thái THÔ phải phân biệt được: '-' là "không áp dụng", 0 là
    # "có khai báo và bằng không", trống là "không điền".
    assert zero['setup_raw'] == 0
    assert str(dash['setup_raw']).strip() == '-'
    assert blank['setup_raw'] is None
    assert all(op['_setup_declared'] for op in (zero, dash, blank))


@pytest.mark.parametrize('bad', [-5, -0.5])
def test_gia_tri_am_bi_chan_khong_coerce_ve_0(bad):
    with pytest.raises(ValueError) as excinfo:
        _parse([('Chân ghế A', 'KM-001', [('CẮT LASER', bad, 100, 3.39)])])
    assert 'không được âm' in str(excinfo.value)


@pytest.mark.parametrize('bad', ['hai mươi', '??', '20 phút'])
def test_gia_tri_khong_doc_duoc_bi_chan_kem_vi_tri(bad):
    with pytest.raises(ValueError) as excinfo:
        _parse([('Chân ghế A', 'KM-001', [('CẮT LASER', bad, 100, 3.39)])])
    message = str(excinfo.value)
    assert 'phải là số' in message and 'Chân ghế A' in message


# --- tổng thời gian là field NGUỒN, không phải thứ suy ra rồi vứt --------

def test_tong_thoi_gian_duoc_giu_nguyen_ke_ca_khi_lech_voi_phep_tinh():
    """OP có số lần thao tác nhiều hơn số thành phẩm -- không được tự sửa.

    File thật có 5 ca như vậy (ĐÓNG ECU VÀO NAN GỖ: 40s/sp nhưng tổng 12,2222h
    => 1100 lần, gấp 10 lần QTY 110). Nếu bỏ tổng rồi tự tính lại theo số
    lượng, định mức của 5 công đoạn đó biến mất.
    """
    parsed = _parse([('Lắp ráp sau khi sơn', 'KM-900', [
        ('ĐÓNG ECU VÀO NAN GỖ', 0, 40, 12.222222222222221)])])
    op = parsed['operations'][0]
    assert op['total_expected_seconds'] == pytest.approx(44000.0, abs=1)
    # Suy ra số lần thực hiện: (tổng - setup) / cycle = 1100, gấp 10 lần 110.
    work_units = (op['total_expected_seconds'] - op['setup_seconds']) / op['standard_seconds_per_unit']
    assert round(work_units) == 1100
    # Và giá trị nguồn vẫn nguyên văn, không bị làm tròn/ghi đè.
    # openpyxl round-trip đổi bit cuối của float; giá trị nguồn không bị LÀM TRÒN.
    assert op['total_expected_raw'] == pytest.approx(12.222222222222221, rel=1e-12)


# --- số lượng: Part độc lập với PO ---------------------------------------

def test_so_luong_sheet_doc_lap_voi_qty_cua_po():
    parsed = _parse([('Chân ghế A', 'KM-001', [('CẮT LASER', 20, 100, 3.39)])],
                    sheet_qty=440)
    assert parsed['qty'] == 110, 'QTY của PO'
    assert parsed['parts'][0]['planned_quantity'] == 440, 'SỐ LƯỢNG của Part'
    assert parsed['parts'][0]['source_quantity_raw'] == 440


def test_so_luong_khong_nhat_nham_o_ben_canh():
    """'SỐ LƯỢNG' là TIÊU ĐỀ CỘT, giá trị ở dòng dưới -- không phải ô kế bên.

    Đọc theo kiểu nhãn->phải sẽ nhặt phải 'HÌNH ẢNH' ở ô kế tiếp. Đây là lỗi
    thật đã xảy ra khi chạy trên file gốc.
    """
    parsed = _parse([('Chân ghế A', 'KM-001', [('CẮT LASER', 20, 100, 3.39)])],
                    sheet_qty=220)
    assert parsed['parts'][0]['planned_quantity'] == 220


# --- danh tính: giữ nguyên số OP gốc -------------------------------------

def test_so_op_trung_trong_cung_part_duoc_ghi_nhan_day_du_de_chan():
    """Trùng số trong CÙNG một Part: ghi nhận đủ tờ/Part/số/hai dòng, không tự sửa.

    Parser không quyết định chặn hay không -- nó chỉ phải ghi đủ thông tin để
    tầng nhập dựng câu lỗi chỉ đúng chỗ phải sửa trên file.
    """
    parsed = _parse([('Thanh la khung ngồi', 'KM-317', [
        ('CẮT LASER', 20, 100, 3.39, 1),
        ('CHAMFER LỖ', 0, 40, 1.22, 2),
        ('LÀM NGUỘI', 0, 50, 1.53, 2),
    ])])
    assert len(parsed['operations']) == 3
    assert [op['source_op_no'] for op in parsed['operations']] == [1, 2, 2]
    # Mã giữ nguyên, KHÔNG hậu tố -- hợp thức hoá lỗi là điều bị cấm.
    assert [op['code'] for op in parsed['operations']] == [
        'KM-317-OP01', 'KM-317-OP02', 'KM-317-OP02']
    note = next(n for n in parsed['notes'] if n['kind'] == 'DUPLICATE_OP_IN_PART')
    assert note['part'] == 'KM-317' and note['op_code'] == 'OP02'
    assert note['first_row'] != note['row']


def test_so_op_trung_giua_hai_part_khong_phai_loi():
    """Mỗi tờ bản vẽ đánh số lại từ OP01 -- không cảnh báo, không lỗi."""
    parsed = _parse([
        ('Chân ghế A', 'KM-001', [('CẮT LASER', 20, 100, 3.39, 1)]),
        ('Chân ghế B', 'KM-002', [('CẮT LASER', 20, 100, 3.39, 1)]),
    ])
    assert [op['source_op_no'] for op in parsed['operations']] == [1, 1]
    assert [op['code'] for op in parsed['operations']] == ['KM-001-OP01', 'KM-002-OP01']
    assert not [n for n in parsed['notes'] if n['kind'] == 'DUPLICATE_OP_IN_PART']

def test_part_number_trong_block_trong_thi_ke_thua_tu_sheet():
    parsed = _parse([('Chân ghế A', 'KM-001', [
        ('CẮT LASER', 20, 100, 3.39, 1, 'KM-001'),
        ('LÀM NGUỘI', 0, 50, 1.53, 2, None),
    ])])
    ops = parsed['operations']
    assert ops[0]['source_part_number'] == 'KM-001'
    assert ops[1]['source_part_number'] == ''
    # Cả hai vẫn thuộc Part của sheet -- trống không phải lỗi.
    assert {op['part_key'] for op in ops} == {'KM-001'}


# --- sheet đặc biệt phải được NÓI RA -------------------------------------

def test_to_muc_quy_trinh_khong_bi_coi_la_thieu_du_lieu():
    """Tờ không khai TÊN lẫn MÃ BẢN VẼ là tờ MỨC QUY TRÌNH, không phải tờ lỗi.

    Công đoạn sơn/kiểm tra/đóng gói làm trên cả cụm, không gắn vào một bản vẽ
    nào. Bắt chúng khai mã bản vẽ là bắt xưởng bịa ra một con số.
    """
    parsed = _parse([('SƠN TĨNH ĐIỆN', None, [('SƠN', 0, 30, 0.92)])],
                    drawing_name=None)
    assert parsed['parts'][0]['sheet_kind'] == 'PROCESS_LEVEL_SHEET'
    assert not [n for n in parsed['notes']
                if n['kind'] == 'DRAWING_SHEET_MISSING_IDENTITY']
    # Và công đoạn của tờ đó vẫn được nhập bình thường.
    assert len(parsed['operations']) == 1


def test_to_khai_ten_ma_thieu_ma_ban_ve_thi_moi_canh_bao():
    """Có TÊN mà thiếu MÃ -- tự nhận là tờ bản vẽ nhưng không có khoá để khớp."""
    parsed = _parse([('Chân ghế A', None, [('CẮT LASER', 20, 100, 3.39)])],
                    drawing_name='Chân ghế A')
    note = next(n for n in parsed['notes']
                if n['kind'] == 'DRAWING_SHEET_MISSING_CODE')
    assert 'MÃ BẢN VẼ' in note['message']


def test_to_co_ma_ma_thieu_ten_thi_khong_canh_bao():
    """Thiếu TÊN mà có MÃ thì không mơ hồ: mã chính là danh tính."""
    parsed = _parse([('Chân ghế A', 'KM-001', [('CẮT LASER', 20, 100, 3.39)])],
                    drawing_name=None)
    assert parsed['parts'][0]['sheet_kind'] == 'DRAWING_PART_SHEET'
    assert not [n for n in parsed['notes'] if n['kind'] == 'DRAWING_SHEET_MISSING_CODE']


def test_metadata_po_duoc_giu_lai():
    parsed = _parse([('Chân ghế A', 'KM-001', [('CẮT LASER', 20, 100, 3.39)])])
    assert parsed['po'] == '6126'
    assert parsed['qty'] == 110
    assert parsed['order_type'] == 'SẢN XUẤT HÀNG LOẠT'
    assert 'SX-1/BM1' in parsed['source_document_meta']
    assert 'SX-1/BM1' in parsed['parts'][0]['source_document_meta']


def test_nhap_lai_cung_file_cho_ket_qua_y_het():
    sheets = [('Chân ghế A', 'KM-001', [
        ('CẮT LASER', 60, 180, 6.5, 1), ('LÀM NGUỘI', '-', 50, 1.53, 2)])]
    first, second = _parse(sheets), _parse(sheets)
    fields = ('code', 'source_op_no', 'source_title', 'setup_seconds',
              'standard_seconds_per_unit', 'total_expected_seconds',
              'requires_setup', 'expected_setup_minutes')
    assert ([{k: op[k] for k in fields} for op in first['operations']]
            == [{k: op[k] for k in fields} for op in second['operations']])
    assert first['notes'] == second['notes']

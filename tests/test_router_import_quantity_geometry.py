"""Sản lượng: đọc được ở cả hai kiểu đặt nhãn, và DỪNG khi hai kiểu đá nhau.

LỖI THẬT (XL-03, audit 2026-09-14). Đầu tờ router dùng HAI kiểu nhãn cùng lúc:

  * nhãn bên trái, giá trị BÊN PHẢI cùng dòng   (``A3='QTY:'`` -> ``C3=110``)
  * nhãn là TIÊU ĐỀ CỘT, giá trị Ở DƯỚI         (``I2='SỐ LƯỢNG'`` -> ``I4=110``)

`SỐ LƯỢNG` đã được sửa sang kiểu thứ hai, kèm hẳn một docstring giải thích --
nhưng `QTY` bị bỏ lại ở kiểu thứ nhất. Trên biến thể đặt `QTY` làm tiêu đề cột,
nó nhặt ô bên cạnh: sản lượng thật 110, nhập vào **2026** (một cái năm).

`planned_quantity` là mẫu số của mọi phép tính tiến độ/năng suất và chỉ được đặt
MỘT LẦN lúc tạo PO. Một PO mở cho 2026 thay vì 110 sẽ hiện ~5% hoàn thành mãi
mãi và không bao giờ chuyển COMPLETED -- không ai thấy vì không có gì báo.

VÌ SAO KHÔNG ĐƠN GIẢN LÀ "ĐỔI SANG ĐỌC Ở DƯỚI". Vì biểu mẫu chuẩn đặt giá trị
BÊN PHẢI (``A3='QTY:'`` -> ``C3``), nên đổi cứng sang kiểu kia sẽ làm hỏng file
thật. Cách đúng là đọc cả hai và chỉ từ chối khi chúng THỰC SỰ mâu thuẫn: ứng
viên phải là số nguyên mới được tính, nên ô ``A4='TÊN BẢN VẼ:'`` nằm dưới
``A3`` tự loại và biểu mẫu chuẩn vẫn chạy.
"""
from __future__ import annotations

import io

import pytest

pytest.importorskip('openpyxl')
from openpyxl import Workbook, load_workbook  # noqa: E402

from mesflow.web.excel_io import _parse_go_router_template  # noqa: E402


def _block(ws, row=8, code='KM-001'):
    ws.cell(row=row, column=1, value='OPERATION # 01- CẮT LASER')
    ws.cell(row=row + 1, column=1, value='Part Number ( Mã bản vẽ )')
    ws.cell(row=row + 1, column=2, value=code)
    ws.cell(row=row + 2, column=12, value='Thời gian Setup ( phút )')
    ws.cell(row=row + 3, column=1, value='SETUP')
    ws.cell(row=row + 3, column=12, value=0)
    ws.cell(row=row + 5, column=12, value='Thời gian gia công / sản phẩm (s)')
    ws.cell(row=row + 6, column=1, value='Thời gian gia công')
    ws.cell(row=row + 6, column=12, value=60)


def _standard_layout(qty=110):
    """Biểu mẫu CHUẨN: nhãn trái, giá trị bên phải. Phải tiếp tục chạy."""
    wb = Workbook(); wb.remove(wb.active)
    ws = wb.create_sheet('Chân ghế A')
    ws['A2'] = 'PO NUMBER:'; ws['C2'] = 'PO-GEO'
    ws['A3'] = 'QTY:'; ws['C3'] = qty
    ws['A4'] = 'TÊN BẢN VẼ:'; ws['C4'] = 'Chân ghế A'
    ws['A5'] = 'MÃ BẢN VẼ:'; ws['C5'] = 'KM-001'
    ws['I2'] = 'SỐ LƯỢNG'; ws['I4'] = qty
    _block(ws)
    buffer = io.BytesIO(); wb.save(buffer); return buffer.getvalue()


def _column_header_layout(qty=110, neighbour=2026):
    """Biến thể: QTY là TIÊU ĐỀ CỘT, và ô bên cạnh có một con số khác."""
    wb = Workbook(); wb.remove(wb.active)
    ws = wb.create_sheet('Chân ghế A')
    ws['A2'] = 'PO NUMBER:'; ws['C2'] = 'PO-GEO'
    ws['H2'] = 'QTY'; ws['I2'] = neighbour       # <- cái bẫy
    ws['H4'] = qty                               # <- sản lượng thật
    ws['A4'] = 'TÊN BẢN VẼ:'; ws['C4'] = 'Chân ghế A'
    ws['A5'] = 'MÃ BẢN VẼ:'; ws['C5'] = 'KM-001'
    _block(ws)
    buffer = io.BytesIO(); wb.save(buffer); return buffer.getvalue()


def _parse(data):
    return _parse_go_router_template(
        load_workbook(io.BytesIO(data), data_only=True), 'geo.xlsx')


# --------------------------------------------------------------------------
def test_the_standard_layout_still_reads_the_cell_to_the_right():
    """Bản vá không được làm hỏng biểu mẫu thật -- đây là ca thường gặp nhất."""
    parsed = _parse(_standard_layout(qty=110))
    assert int(parsed['qty']) == 110
    assert parsed['parts'][0]['planned_quantity'] == 110


def test_a_column_header_layout_never_silently_imports_the_neighbour():
    """Ca đo được: sản lượng thật 110, ô bên cạnh 2026.

    Điều BẮT BUỘC là 2026 không được lọt vào. Dừng kèm lời giải thích là kết
    quả chấp nhận được; nhập âm thầm số sai thì không.
    """
    try:
        parsed = _parse(_column_header_layout(qty=110, neighbour=2026))
    except ValueError as exc:
        assert '2026' in str(exc) and '110' in str(exc), str(exc)
        assert 'QTY' in str(exc)
        return
    assert int(parsed['qty']) != 2026, 'đã nhập nhầm ô bên cạnh làm sản lượng'
    assert int(parsed['qty']) == 110


def test_the_error_names_both_numbers_so_the_user_can_fix_the_sheet():
    """Một lời từ chối không nói rõ hai số nào thì người dùng không sửa được."""
    with pytest.raises(ValueError) as excinfo:
        _parse(_column_header_layout(qty=110, neighbour=2026))
    message = str(excinfo.value)
    assert '2026' in message and '110' in message
    assert 'bên phải' in message and 'bên dưới' in message


def test_agreeing_values_are_fine_in_both_geometries():
    """Hai kiểu cùng chỉ về một số thì không có gì mâu thuẫn để từ chối."""
    wb = Workbook(); wb.remove(wb.active)
    ws = wb.create_sheet('Chân ghế A')
    ws['A2'] = 'PO NUMBER:'; ws['C2'] = 'PO-GEO'
    ws['H2'] = 'QTY'; ws['I2'] = 110      # bên phải
    ws['H4'] = 110                        # bên dưới, cùng số
    ws['A4'] = 'TÊN BẢN VẼ:'; ws['C4'] = 'Chân ghế A'
    ws['A5'] = 'MÃ BẢN VẼ:'; ws['C5'] = 'KM-001'
    _block(ws)
    buffer = io.BytesIO(); wb.save(buffer)
    assert int(_parse(buffer.getvalue())['qty']) == 110


def test_a_text_cell_under_the_label_is_not_a_rival_candidate():
    """Chính điều làm biểu mẫu chuẩn sống sót: ô dưới 'QTY:' là 'TÊN BẢN VẼ:',
    một chuỗi, nên nó không được coi là một ứng viên đối chọi."""
    from mesflow.web.excel_io import _maybe_int
    assert _maybe_int('TÊN BẢN VẼ:') is None
    assert _maybe_int(None) is None and _maybe_int('') is None
    assert _maybe_int(True) is None, 'bool không phải sản lượng'
    assert _maybe_int(110) == 110 and _maybe_int('110') == 110
    assert _maybe_int(110.0) == 110
    assert _maybe_int(110.5) is None


def test_the_per_sheet_quantity_uses_the_same_rule():
    """SỐ LƯỢNG của từng tờ là mẫu số riêng của Part (bội số BOM), nên nó phải
    được canh y hệt -- không để một nửa hệ thống cẩn thận, nửa kia đoán bừa."""
    wb = Workbook(); wb.remove(wb.active)
    ws = wb.create_sheet('Chân ghế A')
    ws['A2'] = 'PO NUMBER:'; ws['C2'] = 'PO-GEO'
    ws['A3'] = 'QTY:'; ws['C3'] = 110
    ws['A4'] = 'TÊN BẢN VẼ:'; ws['C4'] = 'Chân ghế A'
    ws['A5'] = 'MÃ BẢN VẼ:'; ws['C5'] = 'KM-001'
    ws['I2'] = 'SỐ LƯỢNG'; ws['J2'] = 999      # bên phải: mâu thuẫn
    ws['I4'] = 220                             # bên dưới: số thật
    _block(ws)
    buffer = io.BytesIO(); wb.save(buffer)
    with pytest.raises(ValueError) as excinfo:
        _parse(buffer.getvalue())
    assert '999' in str(excinfo.value) and '220' in str(excinfo.value)
    assert 'SỐ LƯỢNG' in str(excinfo.value)

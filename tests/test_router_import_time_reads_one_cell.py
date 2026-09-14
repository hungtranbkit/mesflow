"""Thời gian đọc từ ĐÚNG ô ngay dưới nhãn, không dò xuống dòng kế.

LỖI THẬT (XL-02, audit 2026-09-14). `_block_labeled_time` quét
``block_rows[idx+1:idx+3]`` và lấy ô KHÔNG RỖNG ĐẦU TIÊN. Khi ô ngay dưới nhãn
trống, nó nhảy xuống dòng kế và đọc một con số hoàn toàn khác làm thời gian --
im lặng, không cảnh báo, và con số sai đó trở thành master data.

Hai đường vào, cả hai đều tầm thường ở xưởng:

  * ô để trống (chưa điền, hoặc điền ở chỗ khác);
  * mở file bằng LibreOffice rồi lưu lại -- công thức mất giá trị cache, và với
    ``data_only=True`` openpyxl trả về None cho ô đó, đúng như ô trống.

Đo được trước bản vá: ô setup ghi ``=5*2`` (10 phút), số nhập vào là **777** --
số thợ nằm ở dòng dưới. Và với ô trống hẳn: **999**.

`standard_seconds_per_unit` là mẫu số của mọi phép tính takt, tiến độ và năng
suất; `expected_setup_minutes` sinh ra hẳn một OP SETUP. Đọc nhầm ô nghĩa là
sai toàn bộ những thứ đó mà không có dấu hiệu nào.

Biểu mẫu thật LUÔN đặt giá trị ngay dưới nhãn (L10/L11, L13/L14, M13/M14), nên
phép quét hai dòng chưa bao giờ là dung sai cần thiết -- nó chỉ là một đường
đoán mò, và đoán sai thì ra dữ liệu sai không ai biết.
"""
from __future__ import annotations

import io

import pytest

pytest.importorskip('openpyxl')
from openpyxl import Workbook, load_workbook  # noqa: E402

from mesflow.web.excel_io import _parse_go_router_template  # noqa: E402

BLOCK_HEIGHT = 12


def _router(*, setup_value, cycle_value=100, row_below_setup=None,
            row_below_cycle=None):
    """Một tờ, một block. `row_below_*` đặt số vào dòng KẾ ô giá trị -- đúng
    chỗ mà phép quét hai dòng cũ sẽ vớ phải."""
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet('Chân ghế A')
    ws['A2'] = 'PO NUMBER:'
    ws['C2'] = 'PO-ONECELL'
    ws['A3'] = 'QTY:'
    ws['C3'] = 110
    ws['A4'] = 'TÊN BẢN VẼ:'
    ws['C4'] = 'Chân ghế A'
    ws['A5'] = 'MÃ BẢN VẼ:'
    ws['C5'] = 'KM-001'
    row = 8
    ws.cell(row=row, column=1, value='OPERATION # 01- CẮT LASER')
    ws.cell(row=row + 1, column=1, value='Part Number ( Mã bản vẽ )')
    ws.cell(row=row + 1, column=2, value='KM-001')
    ws.cell(row=row + 2, column=12, value='Thời gian Setup ( phút )')
    ws.cell(row=row + 3, column=1, value='SETUP')
    if setup_value is not None:
        ws.cell(row=row + 3, column=12, value=setup_value)
    if row_below_setup is not None:
        ws.cell(row=row + 4, column=12, value=row_below_setup)
    ws.cell(row=row + 5, column=12, value='Thời gian gia công / sản phẩm (s)')
    ws.cell(row=row + 5, column=13, value='Tổng thời gian gia công dự kiến ( giờ )')
    ws.cell(row=row + 6, column=1, value='Thời gian gia công')
    if cycle_value is not None:
        ws.cell(row=row + 6, column=12, value=cycle_value)
    if row_below_cycle is not None:
        ws.cell(row=row + 7, column=12, value=row_below_cycle)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _op(data):
    parsed = _parse_go_router_template(
        load_workbook(io.BytesIO(data), data_only=True), 'onecell.xlsx')
    return parsed['operations'][0]


# --------------------------------------------------------------------------
def test_a_blank_setup_cell_does_not_borrow_the_row_below():
    """Ca đo được: ô setup trống, dòng dưới có 999 (số thợ) -> KHÔNG được lấy."""
    op = _op(_router(setup_value=None, row_below_setup=999))
    assert op['setup_raw'] is None, f"đã vớ phải {op['setup_raw']!r} ở dòng dưới"
    assert op['setup_seconds'] in (None, 0.0)
    assert op['requires_setup'] is False
    assert op['expected_setup_minutes'] is None


def test_a_blank_cycle_cell_does_not_borrow_the_row_below():
    """Nguy hiểm hơn setup: cycle là mẫu số của mọi phép tính năng suất."""
    op = _op(_router(setup_value=15, cycle_value=None, row_below_cycle=777))
    assert op['cycle_raw'] is None, f"đã vớ phải {op['cycle_raw']!r} ở dòng dưới"
    assert op['standard_seconds_per_unit'] == 0.0


def test_the_normal_case_still_reads_the_cell_right_under_the_label():
    """Bản vá không được làm hỏng đường đi bình thường."""
    op = _op(_router(setup_value=15, cycle_value=70, row_below_setup=999,
                     row_below_cycle=777))
    assert op['setup_raw'] == 15 and op['setup_seconds'] == 900.0
    assert op['expected_setup_minutes'] == 15
    assert op['cycle_raw'] == 70 and op['standard_seconds_per_unit'] == 70.0


def test_zero_and_dash_still_mean_what_they_meant():
    """Hợp đồng ba trạng thái 0 / '-' / trống không bị bản vá này đụng vào."""
    zero = _op(_router(setup_value=0, row_below_setup=999))
    assert zero['setup_raw'] == 0 and zero['setup_seconds'] == 0.0
    assert zero['_setup_declared'] is True
    dash = _op(_router(setup_value='-', row_below_setup=999))
    assert str(dash['setup_raw']).strip() == '-' and dash['setup_seconds'] == 0.0
    assert dash['_setup_declared'] is True
    blank = _op(_router(setup_value=None, row_below_setup=999))
    assert blank['_setup_declared'] is True, 'nhãn vẫn có trên tờ giấy'


def test_an_uncached_formula_behaves_like_blank_not_like_the_row_below():
    """ĐÂY LÀ CA ĐO ĐƯỢC TRONG AUDIT, và nó là lý do bản vá tồn tại.

    Mở file bằng LibreOffice rồi lưu lại thì công thức mất giá trị cache; với
    ``data_only=True`` openpyxl trả về None -- đúng như ô trống (kiểm ngay dưới
    đây, để giả định này không âm thầm sai đi). Bản cũ vì thế nhảy xuống dòng kế
    và đọc 777. Bản mới dừng lại.
    """
    data = _router(setup_value='=5*2', row_below_setup=777)
    assert load_workbook(io.BytesIO(data), data_only=True)['Chân ghế A']['L11'].value is None
    assert load_workbook(io.BytesIO(data), data_only=False)['Chân ghế A']['L11'].value == '=5*2'

    op = _op(data)
    assert op['setup_raw'] is None, f"đã vớ phải {op['setup_raw']!r} (777 = số thợ)"
    assert op['expected_setup_minutes'] is None
    assert op['requires_setup'] is False


def test_a_literal_equals_string_is_named_not_called_not_a_number():
    """Ô chứa CHỮ '=5*2' (gõ thành văn bản) thì openpyxl trả về chuỗi thật.
    Nói thẳng phải mở bằng Excel và lưu lại, thay vì một câu "phải là số"."""
    from mesflow.web.excel_io import SETUP_TIME_LABELS, _block_labeled_time
    with pytest.raises(ValueError) as excinfo:
        _block_labeled_time([('Thời gian Setup ( phút )',), ('=5*2',)],
                            SETUP_TIME_LABELS, where='test', label_vi='Thời gian Setup')
    message = str(excinfo.value)
    assert 'công thức chưa có giá trị' in message
    assert 'Excel' in message and 'lưu lại' in message


def test_the_value_found_flag_distinguishes_label_only_from_a_real_number():
    """Cờ mới, để chỗ nào cần thì phân biệt được mà không đổi nghĩa `declared`."""
    from mesflow.web.excel_io import SETUP_TIME_LABELS, _block_labeled_time
    rows = [('SETUP', None), ('Thời gian Setup ( phút )', None), (None, None)]
    out = _block_labeled_time([(r[0],) for r in rows], SETUP_TIME_LABELS,
                              where='test', label_vi='Thời gian Setup')
    assert out['declared'] is True and out['value_found'] is False
    rows2 = [('Thời gian Setup ( phút )',), (15,)]
    out2 = _block_labeled_time(rows2, SETUP_TIME_LABELS,
                               where='test', label_vi='Thời gian Setup')
    assert out2['declared'] is True and out2['value_found'] is True
    assert out2['seconds'] == 900.0

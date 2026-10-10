"""Authenticated reporting boundaries and export integrity, with isolated MES data."""
import csv
import hashlib
import importlib.util
import io
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
BASE={'report_type':'productivity','from':'2026-10-01','to':'2026-10-10','po_id':None,'operation_id':None,'employee_id':None}


@pytest.fixture
def service(monkeypatch,tmp_path):
    monkeypatch.syspath_prepend(str(ROOT/'services/report-assistant'))
    spec=importlib.util.spec_from_file_location('report_test_server',ROOT/'services/report-assistant/server.py')
    module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
    module.app.config.update(TESTING=True)
    monkeypatch.setattr(module,'AUDIT',tmp_path/'audit.jsonl')
    monkeypatch.setattr(module,'GATEWAY_KEY','')
    module.calls=[]
    def backend(path,params=None):
        module.calls.append((path,params))
        if path=='/api/auth/me':
            cookie=module.session_cookie()
            if 'alice' in cookie:return {'ok':True,'user':{'id':1,'role':'manager','permissions':['dashboard.view','employees.view','po.view','session.view','exceptions.view']}}
            if 'bob' in cookie:return {'ok':True,'user':{'id':2,'role':'manager','permissions':['dashboard.view','employees.view','po.view','session.view','exceptions.view']}}
            if 'denied' in cookie:return {'ok':True,'user':{'id':3,'role':'operator','permissions':['overview.view']}}
            raise module.ReportError('unauthorized',401)
        if path=='/api/reports/employee-productivity':
            return {'ok':True,'employees':[{'employee_id':7,'employee_code':'NV-7','employee_name':'=HYPERLINK("https://invalid")','completed_sessions':2,'good_qty':12,'defect_qty':1,'worked_seconds':3600,'productivity_percent':88.5}]}
        raise AssertionError('Unexpected source '+path)
    monkeypatch.setattr(module,'backend',backend)
    return module


def authenticated(service,who='alice'):
    client=service.app.test_client();client.set_cookie('session',who);return client


def preview(service,client,mode='live',intent=None):
    response=client.post('/reports-api/preview',json={'mode':mode,'intent':intent or BASE})
    assert response.status_code==200,response.json
    return response.json


def test_anonymous_demo_never_queries_mes_and_is_marked_in_files(service):
    client=service.app.test_client();record=preview(service,client,'demo')
    assert service.calls==[]
    assert record['mode']=='demo' and record['rows'][0]['sample'].startswith('DEMO')
    exported=client.post('/reports-api/export',json={'report_id':record['id'],'format':'xlsx'})
    assert exported.status_code==200
    book=load_workbook(io.BytesIO(exported.data))
    assert book.sheetnames[0]=='DEMO'
    assert 'DEMO' in exported.headers['Content-Disposition']
    assert service.calls==[]


def test_anonymous_live_denied_and_role_denied_before_data(service):
    assert service.app.test_client().post('/reports-api/preview',json={'mode':'live','intent':BASE}).status_code==401
    denied=authenticated(service,'denied')
    assert denied.post('/reports-api/preview',json={'mode':'live','intent':BASE}).status_code==403
    assert all(path=='/api/auth/me' for path,_ in service.calls)


@pytest.mark.parametrize('extra',[{'tenant_id':'other'},{'database':'other'},{'sql':'SELECT * FROM users'},{'source_url':'https://evil'},{'user_id':1}])
def test_no_cross_tenant_database_identity_or_query_override(service,extra):
    response=authenticated(service).post('/reports-api/preview',json={'mode':'live','intent':{**BASE,**extra}})
    assert response.status_code==400
    assert service.calls==[]


def test_snapshot_owner_and_current_auth_rechecked_at_download(service):
    alice=authenticated(service);record=preview(service,alice)
    for who,status in [('bob',403),('expired',401),('denied',403)]:
        response=authenticated(service,who).post('/reports-api/export',json={'report_id':record['id'],'format':'csv'})
        assert response.status_code==status
    assert record['id'] not in json.dumps(authenticated(service,'bob').get('/reports-api/status').json)


def test_revoked_permission_for_same_user_blocks_download(service,monkeypatch):
    alice=authenticated(service);record=preview(service,alice)
    monkeypatch.setattr(service,'current_user',lambda required=True:{'id':1,'role':'viewer','permissions':[]})
    assert alice.post('/reports-api/export',json={'report_id':record['id'],'format':'csv'}).status_code==403


def test_csv_xlsx_bytes_match_preview_filters_hash_audit_and_formula_safety(service):
    client=authenticated(service);record=preview(service,client,intent={**BASE,'employee_id':7})
    query=next(params for path,params in service.calls if path.endswith('employee-productivity'))
    assert query=={'from':BASE['from'],'to':BASE['to'],'limit':501,'employee_id':7}
    for format in ('csv','xlsx'):
        response=client.post('/reports-api/export',json={'report_id':record['id'],'format':format})
        assert response.status_code==200
        assert hashlib.sha256(response.data).hexdigest()==response.headers['X-Report-SHA256']
        if format=='csv':
            rows=list(csv.DictReader(io.StringIO(response.data.decode('utf-8-sig'))))
            assert rows[0]['employee_name'].startswith("'=HYPERLINK")
            assert rows[0]['good_qty']=='12'
            assert json.loads(rows[0]['applied_filters'])==record['intent']
        else:
            book=load_workbook(io.BytesIO(response.data),data_only=False)
            assert book['Data']['B2'].value.startswith("'=HYPERLINK")
            assert book['Data']['B2'].data_type=='s'
            assert book['Data']['D2'].value==12
    events=[json.loads(line) for line in service.AUDIT.read_text().splitlines()]
    assert [e['action'] for e in events]==['REPORT_PREVIEW','REPORT_DOWNLOAD_GRANTED','REPORT_DOWNLOAD_GRANTED']
    assert all(e['user_id']==1 and e['filters']['employee_id']==7 for e in events)
    assert 'HYPERLINK' not in service.AUDIT.read_text()


def test_expiry_count_and_audit_failure_fail_closed(service,monkeypatch):
    client=authenticated(service);record=preview(service,client)
    for _ in range(3): assert client.post('/reports-api/export',json={'report_id':record['id'],'format':'csv'}).status_code==200
    assert client.post('/reports-api/export',json={'report_id':record['id'],'format':'csv'}).status_code==429
    service.snapshots[record['id']]['created']-=301
    assert client.post('/reports-api/export',json={'report_id':record['id'],'format':'csv'}).status_code==404
    monkeypatch.setattr(service,'AUDIT',service.AUDIT/'missing/file')
    assert client.post('/reports-api/preview',json={'mode':'live','intent':BASE}).status_code==503


@pytest.mark.parametrize('change',[{'from':None},{'to':'2026-02-30'},{'from':'2026-10-11'},{'to':'2026-12-30'},{'employee_id':True},{'employee_id':0},{'po_id':-1},{'operation_id':'1 OR 1=1'}])
def test_invalid_or_missing_filters_never_query_data(service,change):
    response=authenticated(service).post('/reports-api/preview',json={'mode':'live','intent':{**BASE,**change}})
    assert response.status_code==400
    assert service.calls==[]


def test_intent_gateway_receives_only_public_enums_not_raw_text_ids_or_data(service,monkeypatch):
    captured=[]
    def select(candidates,period):
        captured.append((candidates,period));return {'report_type':'productivity','period':period},'gateway'
    monkeypatch.setattr(service,'ai_select',select)
    response=authenticated(service).post('/reports-api/intent',json={'mode':'live','request':'Năng suất nhân viên #7 hôm qua secret-customer-code'})
    assert response.json['intent']['employee_id']==7
    assert captured==[(['productivity'],'yesterday')]
    assert all(path=='/api/auth/me' for path,_ in service.calls)
    assert response.json['source']=='gateway'


def test_missing_report_type_and_dates_are_explicit_clarifications(service):
    r=service.app.test_client().post('/reports-api/intent',json={'mode':'demo','request':'Tạo báo cáo giúp tôi'})
    assert set(r.json['missing'])=={'report_type','from','to'}
    assert r.json['source']=='rules'


def test_body_origin_export_schema_limits(service):
    client=service.app.test_client()
    assert client.post('/reports-api/intent',json={'request':'x'*3000,'mode':'demo'}).status_code==413
    assert client.post('/reports-api/intent',json={'request':'năng suất','mode':'demo'},headers={'Origin':'https://evil'}).status_code==403
    assert client.post('/reports-api/export',json={'report_id':'anything','format':'pdf'}).status_code==400
    assert client.post('/reports-api/preview',json={'mode':'demo','intent':BASE,'rows':[{'leaked':1}]}).status_code==400


def test_readonly_api_allowlist_excludes_mutating_exception_read(service):
    assert not service.ALLOWED.fullmatch('/api/session-exceptions')
    assert not service.ALLOWED.fullmatch('/api/users')
    assert service.ALLOWED.fullmatch('/api/session-management/123')


def test_scoped_rows_and_no_silent_truncation(service):
    import data
    intent={**BASE,'report_type':'operation_output','employee_id':7}
    with pytest.raises(service.ReportError,match='ngoài phạm vi'):
        data.project(intent,lambda path,params:{'items':[{'session_id':1,'employee_id':8}]})
    with pytest.raises(service.ReportError,match='quá rộng'):
        data.project(intent,lambda path,params:{'items':[{}]*501})


def test_po_operation_and_employee_filters_preserve_source_semantics(service):
    import data
    intent={**BASE,'report_type':'po_progress','po_id':12,'operation_id':34,'employee_id':7}
    calls=[]
    def get(path,params):
        calls.append((path,params))
        if path.endswith('production-orders/12'):
            return {'report':{'production_order':{'id':12,'code':'PO12'},'operations':[{'id':34,'code':'OP34','planned_qty':100,'done_qty':50,'status':'IN_PROGRESS'},{'id':35,'code':'OP35','planned_qty':200}]}}
        if path=='/api/operations/34': return {'item':{'id':34,'code':'OP34','production_order_id':12}}
        return {'items':[{'session_id':1,'po_id':12,'operation_id':34,'employee_id':7,'good_qty':8,'excluded_from_reports':False},{'session_id':2,'po_id':12,'operation_id':34,'employee_id':7,'good_qty':99,'excluded_from_reports':True}]}
    result=data.project(intent,get)
    assert len(result['rows'])==1
    assert result['rows'][0]['done_qty_now']==50
    assert result['rows'][0]['good_qty_in_period']==8
    assert any('ẢNH CHỤP HIỆN TẠI' in note for note in result['notes'])
    assert calls[-1][1]=={'from':BASE['from'],'to':BASE['to'],'limit':501,'po_id':12,'operation_id':34,'employee_id':7}


def test_exceptions_readonly_detail_and_limit(service):
    import data
    calls=[]
    def get(path,params):
        calls.append(path)
        if path=='/api/session-management': return {'items':[{'session_id':9,'employee_code':'NV7','po_code':'PO12','operation_code':'OP34'}]}
        assert path=='/api/session-management/9'
        return {'session':{'session_id':9},'exceptions':[{'exception_code':'MISSING_END','severity':'WARN','workflow_status':'NEW','exception_message':'sample'}]}
    result=data.project({**BASE,'report_type':'exceptions'},get)
    assert result['rows'][0]['exception_code']=='MISSING_END'
    assert '/api/session-exceptions' not in calls


def test_planned_actual_uses_existing_session_scores_and_excludes_invalidated_rows(service):
    import data
    def get(path,params):
        if path.endswith('employee-productivity'): return {'employees':[{'employee_id':7,'employee_code':'NV7','employee_name':'A'}]}
        return {'employee':{'employee_id':7},'sessions':[
            {'session_id':1,'po_code':'PO12','operation_code':'OP34','actual_seconds':100,'expected_seconds':120,'completion_percent':120,'good_qty':10},
            {'session_id':2,'excluded_from_reports':True,'actual_seconds':99,'expected_seconds':99,'completion_percent':100}]}
    result=data.project({**BASE,'report_type':'planned_actual'},get)
    assert len(result['rows'])==1
    assert result['rows'][0]['actual_seconds']==100
    assert result['rows'][0]['expected_seconds']==120
    assert result['rows'][0]['completion_percent']==120


@pytest.mark.parametrize('bad_content',['{"report_type":"DROP TABLE","period":null}', '{"report_type":"productivity","period":null,"sql":"SELECT 1"}', '{"report_type":"productivity","period":"all_time"}'])
def test_hallucinated_gateway_intent_is_rejected(service,monkeypatch,bad_content):
    monkeypatch.setattr(service,'GATEWAY_KEY','test-key-never-a-real-secret')
    sent=[]
    class Reply:
        status=200
        def __init__(self,value): self.value=value
        def __enter__(self): return self
        def __exit__(self,*args): pass
        def read(self,_): return json.dumps(self.value).encode()
    def call(req,timeout):
        sent.append(req.data)
        if req.full_url.endswith('/models'):
            return Reply({'data':[{'id':'facebook-chat','candidates':['gemini-web'],'capabilities':{'queued':False}}]})
        return Reply({'gateway':{'provider':'gemini-web'},'choices':[{'message':{'content':bad_content}}]})
    monkeypatch.setattr(service,'urlopen',call)
    selected,mode=service.ai_select(['productivity'],None)
    assert selected is None and mode=='rules'
    assert len(sent)==2
    assert 'employee_id' not in sent[-1].decode()


def test_non_claude_provider_policy_is_checked_before_intent_completion(service,monkeypatch):
    monkeypatch.setattr(service,'GATEWAY_KEY','test-key-never-a-real-secret')
    class Reply:
        status=200
        def __enter__(self):return self
        def __exit__(self,*args):pass
        def read(self,_):return json.dumps({'data':[{'id':'facebook-chat','candidates':['claude-review'],'capabilities':{'queued':False}}]}).encode()
    def call(req,timeout):
        assert req.full_url.endswith('/models'), 'must not call a costly provider'
        return Reply()
    monkeypatch.setattr(service,'urlopen',call)
    assert service.ai_select(['productivity'],None)==(None,'rules')


def test_operation_only_filter_does_not_mix_reused_codes_across_pos(service):
    import data
    def get(path,params):
        if path=='/api/operations/34':return {'item':{'id':34,'code':'CUT','production_order_id':12}}
        if path.endswith('production-orders/12'):return {'report':{'production_order':{'id':12,'code':'PO12'}}}
        if path.endswith('employee-productivity'):return {'employees':[{'employee_id':7,'employee_code':'NV7','employee_name':'A'}]}
        return {'employee':{'employee_id':7},'sessions':[
            {'session_id':1,'po_code':'PO12','operation_code':'CUT','good_qty':10},
            {'session_id':2,'po_code':'PO99','operation_code':'CUT','good_qty':99}]}
    result=data.project({**BASE,'report_type':'productivity','operation_id':34},get)
    assert result['rows'][0]['good_qty']==10
    assert result['rows'][0]['completed_sessions']==1


def test_nginx_report_patch_keeps_root_auth_and_business_proxy_unchanged():
    path=ROOT/'scripts/prepare-report-assistant-nginx.py'
    spec=importlib.util.spec_from_file_location('report_nginx',path)
    module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
    original='http { server { listen 80 default_server;\n'+module.MARKER+'location / { proxy_pass http://mesflow_backend; } }\nserver { listen 443 ssl default_server;\n'+module.MARKER+'location / { proxy_pass http://mesflow_backend; } } }'
    patched=module.prepare(original)
    assert patched.replace(module.HTTP,'').replace(module.HTTPS,'')==original
    assert 'proxy_set_header Authorization "";' in patched
    assert 'proxy_set_header X-Report-Client-IP $remote_addr;' in patched
    with pytest.raises(ValueError):module.prepare(patched)


def test_audit_failure_blocks_preview_and_download(service,monkeypatch,tmp_path):
    client=authenticated(service);record=preview(service,client)
    monkeypatch.setattr(service,'AUDIT',tmp_path/'nonexistent'/'audit.jsonl')
    assert client.post('/reports-api/preview',json={'mode':'live','intent':BASE}).status_code==503
    assert client.post('/reports-api/export',json={'report_id':record['id'],'format':'csv'}).status_code==503


def test_postgres_decimal_strings_are_real_excel_numbers(service):
    import data
    def get(path,params):
        if path.endswith('employee-productivity'): return {'employees':[{'employee_id':7,'employee_code':'NV7','employee_name':'A'}]}
        return {'employee':{'employee_id':7},'sessions':[{'session_id':1,'po_code':'PO12','operation_code':'CUT','actual_seconds':'100.50','expected_seconds':'120.00','completion_percent':'119.40','good_qty':'10.00'}]}
    result=data.project({**BASE,'report_type':'planned_actual'},get)
    assert result['rows'][0]['actual_seconds']==100.5
    assert result['rows'][0]['good_qty']==10
    for bad in ['NaN','Infinity','not a number',True]:
        with pytest.raises(service.ReportError):data.number(bad)

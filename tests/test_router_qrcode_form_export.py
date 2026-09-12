"""Biểu mẫu THẬT có ô ``QRCODE``: tem phải nằm trong ô đó, trên file đã xuất.

VÌ SAO CÓ BÀI NÀY. Đường marker trước giờ chỉ được kiểm trên workbook dựng tay
bằng openpyxl. Workbook dựng tay không có merge của biểu mẫu, không có bề rộng
cột 49/39/35 ký tự, không có dòng cao 30-50 point, và -- quan trọng nhất -- nó
ghi nhãn bằng CHỮ ở mọi block. Biểu mẫu thật thì không: chỉ block đầu của tờ
đầu có chữ, mọi block sau nhắc lại nhãn bằng công thức (``=$L$10``), các tờ
khác nhắc lại bằng công thức liên sheet. Vì thế bài dựng tay không chạm nổi
chỗ hỏng thật.

File thật `router-newark-arm-chair.xlsx` KHÔNG có ô QRCODE nào (0 lần trong
sharedStrings và trong toàn bộ sheet XML) -- marker là cơ chế dành cho biểu mẫu
thế hệ sau. Nên biểu mẫu-có-marker ở đây được DỰNG RA TỪ CHÍNH FILE THẬT: mở
gói .xlsx, ghi chữ QRCODE vào hai ô trống có sẵn trong mỗi block, giữ nguyên
mọi thứ còn lại -- kể cả giá trị Excel đã cache, thứ mà openpyxl lưu lại sẽ làm
mất. Nhờ vậy fixture không thể trôi khỏi biểu mẫu thật, và không phải nhân đôi
1,9 MB nhị phân trong repo.

Bài này mở FILE ĐÃ GHI RA rồi mới kiểm, không kiểm workbook còn trong bộ nhớ:
lỗi P0 người dùng gặp chỉ nhìn thấy khi mở file.
"""
import re
import zipfile
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from mesflow.web.excel_io import _parse_go_router_template
from mesflow.web.router_export import (
    QR_MARKER_TEXT, QR_MIN_READABLE_PX, _region_size_px, _stamp_source_workbook)
from mesflow.db.repositories.master_data import operation_code_suffix

pytest.importorskip('PIL', reason='Pillow là dependency của tính năng xuất QR')

REAL_FIXTURE = Path(__file__).resolve().parent / 'fixtures/router-newark-arm-chair.xlsx'
#: Đánh dấu MỌI tờ có block, không phải vài tờ. Luật của hợp đồng là "biểu mẫu
#: đã có marker thì phải có đủ", nên một biểu mẫu đánh dấu nửa vời chính là ca
#: bị từ chối -- nó được kiểm riêng ở test_router_qr_marker.py.
FIRST_SHEET = 'Chân ghế A  - Trái'
SECOND_SHEET = 'Chân ghế A  - Phải'

PO = {'id': 1, 'code': 'PO-NEWARK', 'product': 'NEWARK ARM CHAIR',
      'planned_quantity': 110, 'source_template_id': 7}

#: Chỗ đặt marker, tính từ dòng tiêu đề block. +3 nằm giữa nhãn 'Thời gian
#: Setup' (+2) và nhãn thời gian gia công (+5) nên thuộc vùng SETUP; +7 nằm
#: ngoài vùng đó nên thuộc tem OP, và ô J:K dòng ấy là merged range CÓ SẴN của
#: biểu mẫu -- đường "căn giữa trong merge" được kiểm trên hình học thật.
SETUP_MARKER_OFFSET = 3
OP_MARKER_OFFSET = 7
MARKER_COLUMN = 10  # J
MARKER_COLUMN_LETTER = 'J'


def _sheet_parts(archive):
    """Tên sheet -> đường dẫn XML của nó trong gói .xlsx."""
    workbook = archive.read('xl/workbook.xml').decode('utf-8')
    rels = archive.read('xl/_rels/workbook.xml.rels').decode('utf-8')
    target_by_rel = dict(re.findall(r'Id="([^"]+)"[^>]*Target="([^"]+)"', rels))
    parts = {}
    for entry in re.findall(r'<sheet [^>]*/>', workbook):
        name = re.search(r'name="([^"]+)"', entry).group(1)
        rel = re.search(r'r:id="([^"]+)"', entry).group(1)
        target = target_by_rel[rel].lstrip('/')
        parts[name] = target if target.startswith('xl/') else f'xl/{target}'
    return parts


def _write_marker(sheet_xml, coordinate):
    """Ghi chữ QRCODE vào một ô ĐÃ CÓ SẴN trong XML, giữ nguyên style của nó.

    Chỉ đụng vào ô rỗng: nếu ô đang có nội dung thì fixture đang đặt marker đè
    lên dữ liệu của biểu mẫu, và bài kiểm sẽ nói dối.
    """
    pattern = re.compile(r'<c r="%s"([^>]*?)/>' % coordinate)
    match = pattern.search(sheet_xml)
    assert match, f'ô {coordinate} không có trong XML hoặc không rỗng'
    marker = (f'<c r="{coordinate}"{match.group(1)} t="inlineStr">'
              f'<is><t>{QR_MARKER_TEXT}</t></is></c>')
    return sheet_xml[:match.start()] + marker + sheet_xml[match.end():]


def _marker_cell(ws, row, column):
    """Ô thật để ghi marker: nếu (row, column) nằm trong merge thì là góc trên-trái.

    Mười tờ lắp ráp của biểu mẫu có thêm merge B11:K11, nên ô J11 ở đó KHÔNG
    tồn tại với Excel -- ghi vào đấy thì chữ không hiện ra. Đây đúng là hình
    dạng thật mà workbook dựng tay không có.
    """
    for merged in ws.merged_cells.ranges:
        if (merged.min_row <= row <= merged.max_row
                and merged.min_col <= column <= merged.max_col):
            row, column = merged.min_row, merged.min_col
            break
    cell = ws.cell(row=row, column=column)
    assert cell.value in (None, ''), \
        f'{ws.title}!{cell.coordinate} đang có nội dung, không được đè marker lên'
    return cell.coordinate


def _marked_form(real_bytes, sheets=None):
    """File thật + chữ QRCODE trong mỗi block. Mọi thứ khác y nguyên.

    Trả ``(bytes, {tên tờ: [(dòng tiêu đề, ô marker SETUP, ô marker OP)]})``.
    """
    source = zipfile.ZipFile(BytesIO(real_bytes))
    parts = _sheet_parts(source)
    blocks = {}
    edited = {}
    workbook = load_workbook(BytesIO(real_bytes))
    for name in (sheets or workbook.sheetnames):
        ws = workbook[name]
        starts = sorted(cell.row for row in ws.iter_rows() for cell in row
                        if isinstance(cell.value, str)
                        and cell.value.strip().upper().startswith('OPERATION #'))
        if not starts:
            # Tờ không có block thì không có Operation nào để dán tem.
            continue
        xml = source.read(parts[name]).decode('utf-8')
        placed = []
        for start in starts:
            setup_cell = _marker_cell(ws, start + SETUP_MARKER_OFFSET, MARKER_COLUMN)
            op_cell = _marker_cell(ws, start + OP_MARKER_OFFSET, MARKER_COLUMN)
            xml = _write_marker(xml, setup_cell)
            xml = _write_marker(xml, op_cell)
            placed.append((start, setup_cell, op_cell))
        blocks[name] = placed
        edited[parts[name]] = xml.encode('utf-8')

    out = BytesIO()
    with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            target.writestr(item, edited.get(item.filename, source.read(item.filename)))
    return out.getvalue(), blocks


@pytest.fixture(scope='module')
def real_bytes():
    assert REAL_FIXTURE.is_file(), f'thiếu fixture {REAL_FIXTURE}'
    return REAL_FIXTURE.read_bytes()


@pytest.fixture(scope='module')
def marked(real_bytes):
    return _marked_form(real_bytes)


@pytest.fixture(scope='module')
def source_bytes(marked):
    return marked[0]


@pytest.fixture(scope='module')
def parsed(source_bytes):
    return _parse_go_router_template(load_workbook(BytesIO(source_bytes), data_only=True),
                                     REAL_FIXTURE.name)


@pytest.fixture(scope='module')
def rows(parsed):
    part_code_by_key = {p['key']: p['code'] for p in parsed['parts']}
    part_index = {p['key']: i for i, p in enumerate(parsed['parts'])}
    built = []
    next_id = 1000
    for op in parsed['operations']:
        part_code = part_code_by_key[op['part_key']]
        suffix = operation_code_suffix(part_code, op['code'])
        next_id += 1
        op_id = next_id
        setup_id = None
        if op['requires_setup']:
            next_id += 1
            setup_id = next_id
        built.append({
            'id': op_id, 'code': f"{PO['code']}-{suffix}", 'name': op['name'],
            'sort_order': op['sort_order'],
            'standard_seconds_per_unit': op['standard_seconds_per_unit'],
            'requires_setup': op['requires_setup'],
            'expected_setup_minutes': op['expected_setup_minutes'],
            'part_id': part_index[op['part_key']], 'part_code': part_code,
            'part_name': op['_excel_sheet'], 'part_sort': part_index[op['part_key']],
            'op_qr': f'WF|OPID|{op_id}',
            'setup_id': setup_id,
            'setup_code': f"{PO['code']}-{suffix}-SU" if setup_id else None,
            'setup_name': f"Setup {op['name']}" if setup_id else None,
            'setup_minutes': op['expected_setup_minutes'],
            'setup_qr': f'WF|OPID|{setup_id}' if setup_id else None,
            '_sheet': op['_excel_sheet'], '_block_row': op['_excel_row'],
        })
    return built


@pytest.fixture(scope='module')
def stamped(source_bytes, rows):
    wb, placed, matched, unmatched = _stamp_source_workbook(source_bytes, PO, rows)
    assert wb is not None
    assert unmatched == [], 'mọi Operation của fixture phải khớp một block'
    assert matched == len(rows)
    return wb, placed


@pytest.fixture(scope='module')
def exported(stamped):
    """FILE đã ghi ra đĩa rồi mở lại — đúng thứ người dùng mở."""
    wb, _ = stamped
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return load_workbook(buffer)


# --- fixture đúng là "biểu mẫu có marker" --------------------------------

def test_fixture_thuc_su_co_o_qrcode(source_bytes, marked):
    """Nếu fixture mất chữ QRCODE thì cả bài này thành vô nghĩa mà vẫn xanh."""
    _, blocks = marked
    source = load_workbook(BytesIO(source_bytes))
    markers = [(name, cell.coordinate)
               for name in source.sheetnames
               for row in source[name].iter_rows() for cell in row
               if isinstance(cell.value, str) and cell.value.strip().upper() == QR_MARKER_TEXT]
    assert len(markers) == 2 * sum(len(v) for v in blocks.values()), markers


def test_file_that_shop_uses_has_no_marker_yet(real_bytes):
    """Đối chứng: biểu mẫu ĐANG dùng chưa có marker — đó là lý do fixture này tồn tại."""
    wb = load_workbook(BytesIO(real_bytes))
    found = [cell.coordinate for name in wb.sheetnames
             for row in wb[name].iter_rows() for cell in row
             if isinstance(cell.value, str) and cell.value.strip().upper() == QR_MARKER_TEXT]
    assert found == []


def test_gia_tri_cache_cua_excel_van_con_sau_khi_them_marker(source_bytes):
    """Nhãn của block 2 trở đi là công thức; chữ thật nằm ở giá trị đã cache.

    Nếu bước dựng fixture làm mất cache (openpyxl lưu lại là mất), thì phân loại
    marker SETUP/OP mất chỗ dựa và bài kiểm không còn phản ánh file thật.
    """
    values = load_workbook(BytesIO(source_bytes), data_only=True)
    ws = values[FIRST_SHEET]
    assert 'Setup' in str(ws['L22'].value), ws['L22'].value
    assert values[SECOND_SHEET]['L10'].value, 'nhãn liên sheet phải có giá trị cache'


def test_marker_trong_vung_setup_ra_tem_setup_du_nhan_la_cong_thuc(stamped, rows):
    """Chỗ hỏng thật: block 2+ ghi nhãn bằng ``=$L$10``, không có chữ 'Thời gian Setup'.

    Đọc nhãn từ workbook giữ công thức thì mọi marker đều bị xếp là OP, và block
    có hai marker sẽ bị từ chối là 'trùng marker' -- file thật không xuất được.
    """
    _, placed = stamped
    setup_stamps = [item for item in placed if item['kind'] == 'SETUP']
    assert setup_stamps, 'phải có tem SETUP'
    later_blocks = [row for row in rows if row['_block_row'] > 8 and row['setup_id']]
    assert later_blocks, 'fixture phải có block sau block đầu tiên có setup'
    by_operation = {item['operation_id']: item for item in placed}
    for row in later_blocks:
        assert by_operation[row['setup_id']]['kind'] == 'SETUP'


# --- tem đi đúng đường marker --------------------------------------------

def test_moi_tem_deu_theo_marker_khong_cai_nao_roi_vao_lan_doan(stamped, rows):
    _, placed = stamped
    assert placed, 'phải có tem'
    assert {item['placement'] for item in placed} == {'marker'}
    expected = {row['id'] for row in rows} | {row['setup_id'] for row in rows if row['setup_id']}
    assert {item['operation_id'] for item in placed} == expected


def test_tem_dan_vao_dung_o_qrcode_cua_block_minh(stamped, rows, marked):
    """Marker nào thuộc block nào — dán nhầm block là quét nhầm việc."""
    _, blocks = marked
    cells = {(sheet, start): (setup_cell, op_cell)
             for sheet, entries in blocks.items() for start, setup_cell, op_cell in entries}
    _, placed = stamped
    by_operation = {item['operation_id']: item for item in placed}
    for row in rows:
        setup_cell, op_cell = cells[(row['_sheet'], row['_block_row'])]
        op_item = by_operation[row['id']]
        assert op_item['sheet'] == row['_sheet']
        assert op_item['marker_cell'] == op_cell
        if row['setup_id']:
            setup_item = by_operation[row['setup_id']]
            assert setup_item['sheet'] == row['_sheet']
            assert setup_item['marker_cell'] == setup_cell


def test_payload_la_danh_tinh_canonical_cua_dung_operation(stamped, rows):
    _, placed = stamped
    by_operation = {item['operation_id']: item for item in placed}
    for row in rows:
        assert by_operation[row['id']]['payload'] == f"WF|OPID|{row['id']}"
        if row['setup_id']:
            assert by_operation[row['setup_id']]['payload'] == f"WF|OPID|{row['setup_id']}"


# --- trên FILE đã xuất ----------------------------------------------------

def test_file_xuat_ra_khong_con_chu_qrcode_nao(exported):
    """Để nguyên chữ thì tờ giấy in ra có chữ QRCODE nằm dưới tem."""
    left = [(name, cell.coordinate) for name in exported.sheetnames
            for row in exported[name].iter_rows() for cell in row
            if isinstance(cell.value, str) and cell.value.strip().upper() == QR_MARKER_TEXT]
    assert left == []


def test_anh_trong_file_xuat_ra_nam_gon_trong_vung_marker(source_bytes, exported, stamped):
    """Đúng lỗi P0: neo đúng ô nhưng ảnh giữ cỡ gốc nên tràn ra ngoài ô.

    Kiểm trên file đã ghi: mỗi tem phải neo vào góc trên-trái của vùng marker,
    và toàn bộ bề ngang/bề dọc của nó phải nằm lọt trong vùng ấy.
    """
    from openpyxl.utils.units import EMU_to_pixels

    source = load_workbook(BytesIO(source_bytes))
    _, placed = stamped
    regions = {}
    for item in placed:
        min_cell, max_cell = item['region'].split(':')
        regions.setdefault(item['sheet'], []).append((min_cell, max_cell, item))

    checked = 0
    for sheet_name, entries in regions.items():
        ws_source = source[sheet_name]
        ws_out = exported[sheet_name]
        added = [image for image in ws_out._images
                 if type(image.anchor).__name__ == 'OneCellAnchor']
        anchored = {(image.anchor._from.col, image.anchor._from.row): image
                    for image in added}
        assert len(anchored) == len(added), 'hai tem chồng lên cùng một ô'
        for min_cell, _max_cell, item in entries:
            from openpyxl.utils.cell import coordinate_to_tuple
            row0, col0 = coordinate_to_tuple(min_cell)
            image = anchored.get((col0 - 1, row0 - 1))
            assert image is not None, f'{sheet_name}: không có tem neo ở {min_cell}'
            region = _region_from(item['region'])
            width, height = _region_size_px(ws_source, region)
            side_x = EMU_to_pixels(image.anchor.ext.cx)
            side_y = EMU_to_pixels(image.anchor.ext.cy)
            offset_x = EMU_to_pixels(image.anchor._from.colOff)
            offset_y = EMU_to_pixels(image.anchor._from.rowOff)
            assert side_x == side_y, 'QR méo thì quét không ra'
            assert offset_x + side_x <= width, (
                f'{sheet_name} {min_cell}: tem tràn ngang ({offset_x}+{side_x} > {width})')
            assert offset_y + side_y <= height, (
                f'{sheet_name} {min_cell}: tem tràn dọc ({offset_y}+{side_y} > {height})')
            # Căn giữa: lề hai bên chênh nhau nhiều nhất 1 px do chia nguyên.
            assert abs((width - side_x) - 2 * offset_x) <= 1
            assert abs((height - side_y) - 2 * offset_y) <= 1
            checked += 1
    assert checked == len(placed), 'phải kiểm hết mọi tem'


def test_tem_trong_merged_range_cua_bieu_mau_that(stamped):
    """Marker OP rơi vào vùng J:K đã merge SẴN của biểu mẫu, không phải merge do test tạo."""
    _, placed = stamped
    op_items = [item for item in placed if item['kind'] == 'OP']
    assert op_items
    merged = [item for item in op_items if item['region'].split(':')[0] != item['region'].split(':')[1]]
    assert merged, 'không có tem nào rơi vào merged range -> bài này không kiểm được gì'
    for item in merged:
        min_cell, max_cell = item['region'].split(':')
        assert min_cell.startswith(MARKER_COLUMN_LETTER) and max_cell.startswith('K'), item


def test_o_marker_cao_mot_dong_bi_bao_la_qua_nho_chu_khong_im_lang(stamped):
    """Phát hiện thật về biểu mẫu này: ô cao một dòng cho ra tem quá nhỏ để quét.

    Dòng của tờ router cao 30 point (~40 px), nên một ô QRCODE cao đúng một
    dòng chỉ chứa nổi tem ~34 px -- dưới ngưỡng đọc ổn định. Hệ thống KHÔNG
    được im lặng thu nhỏ rồi coi như xong: nó vẫn dán, nhưng phải nói ra để
    người vẽ biểu mẫu gộp ô cao hơn.
    """
    _, placed = stamped
    one_row = [item for item in placed
               if item['region'].split(':')[0][1:] == item['region'].split(':')[1][1:]]
    assert one_row, 'biểu mẫu này toàn ô marker cao một dòng'
    assert all(item['too_small'] for item in one_row)
    assert all(item['size_px'] < QR_MIN_READABLE_PX for item in one_row)
    # Và tem vẫn được dán, không phải bỏ qua.
    assert all(item['placement'] == 'marker' for item in one_row)


# --- phần còn lại của biểu mẫu không đổi ----------------------------------

def test_ngoai_tem_va_o_marker_bieu_mau_khong_doi_gi(source_bytes, exported):
    """Merge, kích thước ô, khổ in, công thức và mọi ô khác giữ nguyên."""
    source = load_workbook(BytesIO(source_bytes))
    # Khác biệt DUY NHẤT được phép: những ô đặt chỗ QRCODE bị xoá chữ.
    cleared = {(name, cell.coordinate) for name in source.sheetnames
               for row in source[name].iter_rows() for cell in row
               if isinstance(cell.value, str)
               and cell.value.strip().upper() == QR_MARKER_TEXT}

    assert exported.sheetnames == source.sheetnames
    for name in source.sheetnames:
        before, after = source[name], exported[name]
        assert {str(r) for r in after.merged_cells.ranges} == \
            {str(r) for r in before.merged_cells.ranges}, f'{name}: merge đổi'
        assert {k: v.width for k, v in after.column_dimensions.items()} == \
            {k: v.width for k, v in before.column_dimensions.items()}, f'{name}: bề rộng cột đổi'
        assert {k: v.height for k, v in after.row_dimensions.items()} == \
            {k: v.height for k, v in before.row_dimensions.items()}, f'{name}: chiều cao dòng đổi'
        assert after.print_area == before.print_area
        assert after.page_setup.orientation == before.page_setup.orientation
        for row in before.iter_rows():
            for cell in row:
                if (name, cell.coordinate) in cleared:
                    assert after[cell.coordinate].value is None
                    continue
                assert after[cell.coordinate].value == cell.value, \
                    f'{name}!{cell.coordinate} đổi giá trị'


def test_anh_co_san_cua_bieu_mau_van_con(source_bytes, exported):
    """Logo / hình sản phẩm của khách không được biến mất vì ta thêm tem."""
    source = load_workbook(BytesIO(source_bytes))
    for name in source.sheetnames:
        before = len(source[name]._images)
        after = len([image for image in exported[name]._images
                     if type(image.anchor).__name__ != 'OneCellAnchor'])
        assert after >= before, f'{name}: mất ảnh gốc ({before} -> {after})'


def _region_from(text):
    from openpyxl.utils.cell import coordinate_to_tuple
    min_cell, max_cell = text.split(':')
    min_row, min_col = coordinate_to_tuple(min_cell)
    max_row, max_col = coordinate_to_tuple(max_cell)
    return min_row, min_col, max_row, max_col


def test_marker_thua_bi_xoa_chu_nhung_khong_dan_tem(stamped, source_bytes, exported):
    """Block không có setup vẫn có ô chừa cho tem SETUP — chữ phải biến mất.

    'QRCODE' là chữ đặt chỗ, không phải nội dung biểu mẫu. Để nguyên thì tờ giấy
    in ra có chữ QRCODE nằm giữa block mà chẳng có tem nào. Nhưng cũng KHÔNG
    được bịa ra một tem SETUP cho công đoạn không hề set máy.
    """
    _, placed = stamped
    source = load_workbook(BytesIO(source_bytes))
    markers = {(name, cell.coordinate) for name in source.sheetnames
               for row in source[name].iter_rows() for cell in row
               if isinstance(cell.value, str)
               and cell.value.strip().upper() == QR_MARKER_TEXT}
    stamped_cells = {(item['sheet'], item['marker_cell']) for item in placed}
    unused = markers - stamped_cells
    assert unused, 'biểu mẫu phải có ô chừa cho setup ở block không set máy'
    for name, coordinate in unused:
        assert exported[name][coordinate].value is None, f'{name}!{coordinate} còn chữ'
    setup_stamps = {item['operation_id'] for item in placed if item['kind'] == 'SETUP'}
    assert len(setup_stamps) == len([item for item in placed if item['kind'] == 'SETUP'])

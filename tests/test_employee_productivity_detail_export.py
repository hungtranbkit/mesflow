"""Excel + print năng suất nhân viên (hotfix 2026-09-29, one file + print).

- "Xuất Excel" = ONE .xlsx: sheet "Tổng hợp" (unchanged summary) + one
  detail sheet per employee in the current filters, one row per session.
- "In" = printable HTML of the same data (all employees, or one employee
  that must belong to the current filters).
"""
from datetime import date, datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from mesflow.db.repositories.base import NotFoundError
from mesflow.web import analytics as analytics_module
from mesflow.web import app as app_module
from mesflow.web import auth as auth_module
from mesflow.web.productivity_excel import (
    EMPLOYEE_HEADERS,
    EXPORT_HEADERS,
    SUMMARY_SHEET,
    build_employee_productivity_workbook,
    employee_sheet_names,
    session_status_note,
)
from mesflow.web.productivity_export import group_sessions, load_export_data, parse_filters

ROOT = Path(__file__).resolve().parents[1]
UTC = timezone.utc


def _employee(emp_id, code, name, dept, sessions, valid, pct, good, defect, seconds):
    return {"employee_id": emp_id, "employee_code": code, "employee_name": name, "department": dept, "team": "",
            "completed_sessions": sessions, "completed_valid_sessions": valid, "completed_invalid_sessions": sessions - valid,
            "repair_sessions": 0, "productivity_percent": pct, "good_qty": good, "defect_qty": defect,
            "worked_seconds": seconds}


def _session(sid, emp_id, code, name, dept, op, start, end, good, defect, std, pct, **extra):
    row = {"session_id": sid, "employee_id": emp_id, "employee_code": code, "employee_name": name, "department": dept,
           "team": "", "status": "CLOSED", "started_at": start, "ended_at": end, "work_date": date(2026, 9, end.day),
           "actual_seconds": (end - start).total_seconds(), "standard_seconds_per_unit": std,
           "expected_seconds": std * (good + defect), "good_qty": good, "defect_qty": defect, "po_code": "PO-1",
           "part_code": "P-1", "part_name": "Thân", "operation_code": op, "operation_name": f"Công đoạn {op}",
           "station_code": "ST1", "closed_by_system": False, "quantity_confirmed": True, "close_reason": None,
           "note": None, "is_repair": False, "operation_type": "PRODUCTION", "completion_percent": pct}
    row.update(extra)
    return row


T = lambda d, h, m=0: datetime(2026, 9, d, h, m, tzinfo=UTC)  # 01:00Z = 08:00 ICT
SESSIONS = [
    # An (NV01): 3 sessions, TWO on the same OP10 -> two separate rows.
    _session(12, 1, "NV01", "An", "May", "OP10", T(2, 3), T(2, 4), 36, 0, 100, 100.0),
    _session(11, 1, "NV01", "An", "May", "OP10", T(2, 1), T(2, 2), 24, 0, 120, 80.0),
    _session(13, 1, "NV01", "An", "May", "OP20", T(3, 1), T(3, 1, 30), 0, 0, 0, None),
    # Bình (NV02): auto-closed, quantity never confirmed.
    _session(21, 2, "NV02", "Bình", "May", "OP10", T(2, 5), T(2, 6), 45, 5, 60, 83.33,
             closed_by_system=True, quantity_confirmed=False, close_reason="Hết ca"),
    # Cường (NV03): other department.
    _session(31, 3, "NV03", "Cường", "Kho", "OP90", T(2, 1), T(2, 2), 10, 0, 360, 100.0),
]
EMPLOYEES = [
    _employee(1, "NV01", "An", "May", 3, 2, 90.0, 60, 0, 9000),
    _employee(2, "NV02", "Bình", "May", 1, 1, 83.33, 45, 5, 3600),
    _employee(3, "NV03", "Cường", "Kho", 1, 1, 100.0, 10, 0, 3600),
]


class FakeRepo:
    """Stands in for ReportRepository: records calls so tests can prove the
    export uses exactly two queries and passes the filters through."""

    def __init__(self):
        self.calls = []

    def employee_productivity(self, date_from, date_to, employee_id, department, team, limit):
        self.calls.append(("summary", date_from, date_to, employee_id, department, team))
        rows = [dict(e) for e in EMPLOYEES if (not employee_id or e["employee_id"] == employee_id)
                and (not department or e["department"] == department)]
        return {"summary": {"from": date_from or "2026-09-01", "to": date_to or "2026-09-29"}, "employees": rows}

    def employee_productivity_sessions(self, date_from, date_to, employee_id, department, team):
        self.calls.append(("sessions", date_from, date_to, employee_id, department, team))
        rows = [dict(s) for s in SESSIONS if (not employee_id or s["employee_id"] == employee_id)
                and (not department or s["department"] == department)]
        return {"sessions": rows, "truncated": False}


def _data(**args):
    return load_export_data(parse_filters({"from": "2026-09-01", "to": "2026-09-29", **args}), FakeRepo())


def _wb(**args):
    return load_workbook(BytesIO(build_employee_productivity_workbook(_data(**args)).getvalue()))


def _rows(ws, first=5):
    return [[ws.cell(r, c).value for c in range(1, ws.max_column + 1)] for r in range(first, ws.max_row + 1)
            if ws.cell(r, 1).value is not None]


# ------------------------------------------------------------ workbook shape
def test_one_workbook_summary_first_then_one_sheet_per_employee():
    wb = _wb()
    assert wb.sheetnames[0] == SUMMARY_SHEET and len(wb.sheetnames) == 1 + 3
    # Summary order == employee sheet order (productivity desc: Cường 100, An 90, Bình 83.33).
    summary_codes = [r[1] for r in _rows(wb[SUMMARY_SHEET])]
    assert summary_codes == ["NV03", "NV01", "NV02"]
    assert [ws.title for ws in wb.worksheets[1:]] == ["NV03 Cường", "NV01 An", "NV02 Bình"]


def test_filters_decide_which_employee_sheets_exist():
    assert _wb(department="May").sheetnames == [SUMMARY_SHEET, "NV01 An", "NV02 Bình"]
    assert _wb(search="bình").sheetnames == [SUMMARY_SHEET, "NV02 Bình"]
    one = _wb(employee_id="1")
    assert one.sheetnames == [SUMMARY_SHEET, "NV01 An"]
    assert [r[1] for r in _rows(one[SUMMARY_SHEET])] == ["NV01"]


def test_employee_sheet_one_row_per_session_grouped_and_mapped():
    ws = _wb()["NV01 An"]
    assert [ws.cell(4, c).value for c in range(1, len(EMPLOYEE_HEADERS) + 1)] == EMPLOYEE_HEADERS
    for required in ("Mã NV", "Tên NV", "Ngày", "PO", "Part", "Mã OP", "Operation", "Bắt đầu", "Kết thúc",
                     "Thời gian thực tế", "Sản lượng đạt", "Thời gian định mức", "% năng suất",
                     "Trạng thái / ghi chú", "NV xác nhận"):
        assert required in EMPLOYEE_HEADERS
    col = {h: i for i, h in enumerate(EMPLOYEE_HEADERS)}
    rows = _rows(ws)
    assert [r[col["Mã phiên"]] for r in rows] == [11, 12, 13]  # by start time, not input order
    assert [r[col["Mã OP"]] for r in rows] == ["OP10", "OP10", "OP20"]  # same OP twice = 2 rows
    first = rows[0]
    assert first[col["Mã NV"]] == "NV01" and first[col["Tên NV"]] == "An"
    assert first[col["Part"]] == "P-1 · Thân"
    assert first[col["Bắt đầu"]] == datetime(2026, 9, 2, 8, 0)  # local wall-clock
    assert first[col["Thời gian thực tế"]] == timedelta(hours=1)
    assert first[col["Thời gian định mức"]] == timedelta(seconds=2880)
    assert first[col["% năng suất"]] == 0.8
    assert rows[2][col["% năng suất"]] is None
    assert first[col["NV xác nhận"]] is None
    assert "Năng suất TB: 90.0%" in ws["A2"].value and "Bộ phận: May" in ws["A2"].value
    assert ws.freeze_panes == "A5" and ws.auto_filter.ref == "A4:S7"
    assert ws["I5"].number_format == "dd/mm/yyyy hh:mm" and ws["K5"].number_format == "[h]:mm:ss"
    assert ws["P5"].number_format == "0.0%" and ws["D5"].number_format == "dd/mm/yyyy"
    assert ws.column_dimensions["Q"].width >= 30


def test_employee_percent_is_the_average_of_their_session_percents():
    for group in _data()["groups"]:
        emp, sessions = group["employee"], group["sessions"]
        scored = [s["completion_percent"] for s in sessions if s["completion_percent"] is not None]
        assert len(sessions) == emp["completed_sessions"] and len(scored) == emp["completed_valid_sessions"]
        assert round(sum(scored) / len(scored), 2) == emp["productivity_percent"]


def test_summary_sheet_semantics_unchanged():
    ws = _wb(search="an", department="May", sort="good_qty", dir="desc")[SUMMARY_SHEET]
    assert [ws.cell(4, c).value for c in range(1, len(EXPORT_HEADERS) + 1)] == EXPORT_HEADERS
    assert ws["A1"].value == "BÁO CÁO NĂNG SUẤT NHÂN VIÊN"
    assert "Bộ phận: May" in ws["A2"].value and "Tìm nhân viên: an" in ws["A2"].value
    assert "Năng suất TB theo bộ lọc: 90.0%" in ws["A2"].value
    assert ws["B5"].value == "NV01" and ws["H5"].value == 0.9 and ws["L5"].value is None
    assert ws.freeze_panes == "A5" and ws["K5"].number_format == "[h]:mm"


def test_exactly_two_queries_whatever_the_employee_count():
    repo = FakeRepo()
    data = load_export_data(parse_filters({"from": "2026-09-01", "department": "May"}), repo)
    assert [c[0] for c in repo.calls] == ["summary", "sessions"]
    assert repo.calls[1][4] == "May"  # filters reach the bulk session query
    assert len(data["groups"]) == 2


# ------------------------------------------------------------ sheet names
def test_sheet_names_are_valid_unique_and_deterministic():
    emps = [
        {"employee_id": 1, "employee_code": "NV01", "employee_name": "Nguyễn Thị Minh Khai Phương Hoàng Anh"},
        {"employee_id": 2, "employee_code": "NV01", "employee_name": "Nguyễn Thị Minh Khai Phương Hoàng Anh"},
        {"employee_id": 3, "employee_code": "A/B", "employee_name": "x[y]:z*?\\"},
        {"employee_id": 4, "employee_code": "", "employee_name": "Tổng hợp"},
        {"employee_id": 5, "employee_code": "", "employee_name": "TỔNG HỢP"},
        {"employee_id": 6, "employee_code": "'Q'", "employee_name": ""},
        {"employee_id": 7, "employee_code": "", "employee_name": ""},
    ]
    names = employee_sheet_names(emps)
    assert names == employee_sheet_names(emps)  # deterministic
    assert all(0 < len(n) <= 31 for n in names)
    assert all(not set(n) & set('[]:*?/\\') for n in names)
    assert all(not n.startswith("'") and not n.endswith("'") for n in names)
    assert len({n.casefold() for n in names}) == len(names)
    assert SUMMARY_SHEET.casefold() not in {n.casefold() for n in names}
    assert names[0] == "NV01 Nguyễn Thị Minh Khai Phương"[:31]
    assert names[1].endswith(" (2)") and len(names[1]) <= 31
    assert names[2] == "A B x y z"
    assert names[3] == "Tổng hợp (2)" and names[4] == "TỔNG HỢP (3)"
    assert names[5] == "Q" and names[6] == "NV 7"
    # openpyxl accepts every generated name in one workbook.
    wb = build_employee_productivity_workbook({"report": {"summary": {}}, "employees": [],
        "groups": [{"employee": e, "sessions": []} for e in emps], "filters": {}})
    assert load_workbook(BytesIO(wb.getvalue())).sheetnames == [SUMMARY_SHEET, *names]


def test_status_note_uses_existing_session_flags_only():
    by_id = {s["session_id"]: s for s in SESSIONS}
    assert session_status_note(by_id[11]) == "Tính năng suất"
    assert session_status_note(by_id[13]) == "Thiếu định mức · không tính năng suất"
    assert session_status_note(by_id[21]) == "Tính năng suất · Hệ thống tự đóng · Chưa xác nhận sản lượng · Hết ca"
    assert session_status_note({**by_id[11], "completion_percent": None, "good_qty": 0}) == "Không có sản lượng · không tính năng suất"


def test_group_sessions_drops_employees_outside_the_list():
    groups = group_sessions([EMPLOYEES[1]], SESSIONS)
    assert [g["employee"]["employee_code"] for g in groups] == ["NV02"]
    assert [s["session_id"] for s in groups[0]["sessions"]] == [21]


def test_selected_employee_must_be_inside_the_filters():
    with pytest.raises(NotFoundError):
        _data(employee_id="3", department="May")
    with pytest.raises(ValueError):
        parse_filters({"employee_id": "1; drop"})


# ------------------------------------------------------------ HTTP routes
@pytest.fixture
def client(monkeypatch):
    state = {"logged_in": True}
    monkeypatch.setattr(auth_module, "validate_and_touch", lambda: None if state["logged_in"] else "NOT_LOGGED_IN")
    monkeypatch.setattr(analytics_module, "ReportRepository", FakeRepo)
    app = app_module.create_app()
    app.config["TESTING"] = True
    c = app.test_client()
    with c.session_transaction() as sess:
        sess["user_id"] = 1
        sess["username"] = "quanly"
        sess["role"] = "manager"
    c.state = state
    return c


BASE = "/api/reports/employee-productivity"
Q = "from=2026-09-01&to=2026-09-29"


def test_export_route_returns_one_workbook_with_employee_sheets(client):
    r = client.get(f"{BASE}/export.xlsx?{Q}&department=May")
    assert r.status_code == 200
    assert r.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "nang-suat-nhan-vien_2026-09-01_2026-09-29.xlsx" in r.headers["Content-Disposition"]
    assert load_workbook(BytesIO(r.data)).sheetnames == [SUMMARY_SHEET, "NV01 An", "NV02 Bình"]


def test_print_all_has_summary_and_a_page_per_employee(client):
    html = client.get(f"{BASE}/print?{Q}&department=May&autoprint=1").get_data(as_text=True)
    assert "BÁO CÁO NĂNG SUẤT NHÂN VIÊN" in html and "Tổng hợp" in html
    assert html.count('class="employee page-break"') == 2
    assert 'data-employee-id="1"' in html and 'data-employee-id="2"' in html and 'data-employee-id="3"' not in html
    assert "Bộ phận:</b> May" in html and "size: A4 landscape" in html
    assert "window.print()" in html and "thead{display:table-header-group}" in html
    assert "Hệ thống tự đóng · Chưa xác nhận sản lượng" in html  # session detail rendered
    assert "80,0%" in html and "90,0%" in html  # session % and employee % (vi format)


def test_print_one_employee_only(client):
    r = client.get(f"{BASE}/print?{Q}&employee_id=2")
    html = r.get_data(as_text=True)
    assert r.status_code == 200
    assert "BÁO CÁO NĂNG SUẤT · NV02 · Bình" in html
    assert 'data-employee-id="2"' in html and 'data-employee-id="1"' not in html
    assert "page-break" not in html.split("<main", 1)[1]  # no blank first page for one employee
    assert "setTimeout(() => window.print()" not in html  # autoprint only when asked


def test_print_rejects_employee_outside_filters_and_bad_ids(client):
    assert client.get(f"{BASE}/print?{Q}&department=Kho&employee_id=1").status_code == 404
    assert client.get(f"{BASE}/print?{Q}&employee_id=abc").status_code == 400
    assert client.get(f"{BASE}/export.xlsx?{Q}&department=Kho&employee_id=1").status_code == 404


def test_print_and_export_require_login(client):
    client.state["logged_in"] = False
    assert client.get(f"{BASE}/print?{Q}").status_code == 401
    assert client.get(f"{BASE}/export.xlsx?{Q}").status_code == 401


# ------------------------------------------------------------ source contracts
def test_repository_detail_uses_the_same_scope_and_score_as_the_summary():
    src = (ROOT / "app/mesflow/db/repositories/analytics.py").read_text(encoding="utf-8")
    agg = src.split("    def employee_productivity(self", 1)[1].split("\n    def ", 1)[0]
    det = src.split("    def employee_productivity_sessions(self", 1)[1].split("\n    def ", 1)[0]
    for body in (agg, det):
        assert "self._productivity_scope(date_from,date_to,employee_id,department,team)" in body
        assert "{self._SESSION_COMPLETION_PERCENT_SQL} completion_percent" in body
    assert det.count("fetch_all(") == 1


def test_frontend_single_excel_button_and_print_dialog():
    js = (ROOT / "app/mesflow/web/static/pages/employee-productivity.js").read_text(encoding="utf-8")
    # Legacy layout: plain .btn; P3 layout: "Xuất Excel" is the one .primary.
    assert 'id="epExport" type="button">Xuất Excel</button>' in js
    assert js.count('id="epExport"') == 1
    assert '<button class="btn" id="epPrint" type="button" aria-haspopup="dialog">In</button>' in js
    assert "rowMenu" not in js and "mode=detail" not in js and "'detail'" not in js
    assert "document.getElementById('epExport').onclick = exportExcel;" in js
    assert "In toàn bộ" in js and "In theo nhân viên" in js
    assert "window.open(`/api/reports/employee-productivity/print?${q.toString()}`, '_blank')" in js
    assert "rows.map(x => `<option value=\"${x.employee_id}\">" in js  # choices = current filtered rows

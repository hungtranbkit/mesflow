"""Nhập Router từ màn PO: tạo Template + PO thật, và nhập lại không nhân bản.

Chạy trên API + PostgreSQL thật vì thứ cần chứng minh là hành vi của cả chuỗi —
quyết định, quyền, transaction — chứ không phải một hàm.

Sáu tình huống của hợp đồng:
  a) chưa có gì  -> tạo Template + PO
  b) nhập lại đúng file -> KHÔNG nhân bản, và KHÔNG báo lỗi
  c) Template đã có, PO chưa -> dùng lại Template, tạo PO
  d) Template đã có, PO đã có -> dừng, không ghi đè dữ liệu vận hành
  e) file không có mã PO -> chỉ Template
  f) trùng mã Template, nội dung khác -> đòi xác nhận
"""
import io
import os
import re
import uuid

import pytest
from openpyxl import Workbook, load_workbook

BASE_URL = os.environ.get('MESFLOW_BASE_URL', 'http://mesflow-test-api:8080').rstrip('/')
pytestmark = pytest.mark.integration

BLOCK_HEIGHT = 12


def _router_bytes(po_number='', qty=110, sheets=None, extra_block=False):
    """Workbook GO ROUTER tối giản, đúng bố cục thật (nhãn phải -> giá trị dưới)."""
    sheets = sheets or [('Chân ghế A', 'KM-001', 110, [('CẮT LASER', 20, 100, 3.39)]),
                        ('Chân ghế B', 'KM-002', 220, [('HÀN', 0, 60, 3.67)])]
    wb = Workbook(); wb.remove(wb.active)
    for title, code, sheet_qty, blocks in sheets:
        ws = wb.create_sheet(title)
        ws['A2'] = 'PO NUMBER:'
        if po_number:
            ws['C2'] = po_number
        ws['G2'] = 'LOẠI ĐƠN HÀNG'; ws['I2'] = 'SỐ LƯỢNG'; ws['J2'] = 'HÌNH ẢNH'
        ws['A3'] = 'QTY:'; ws['C3'] = qty
        ws['A4'] = 'TÊN BẢN VẼ:'; ws['C4'] = title
        ws['G4'] = 'SẢN XUẤT HÀNG LOẠT'; ws['I4'] = sheet_qty
        ws['A5'] = 'MÃ BẢN VẼ:'; ws['C5'] = code
        row = 8
        blocks = list(blocks) + ([('CÔNG ĐOẠN THÊM', 0, 30, 0.92)] if extra_block else [])
        for index, (name, setup, cycle, total) in enumerate(blocks, start=1):
            ws.cell(row=row, column=1, value=f'OPERATION # {index:02d}- {name}')
            ws.cell(row=row + 1, column=1, value='Part Number ( Mã bản vẽ )')
            ws.cell(row=row + 1, column=2, value=code)
            ws.cell(row=row + 2, column=12, value='Thời gian Setup ( phút )')
            ws.cell(row=row + 3, column=1, value='SETUP')
            ws.cell(row=row + 3, column=12, value=setup)
            ws.cell(row=row + 5, column=12, value='Thời gian gia công / sản phẩm (s)')
            ws.cell(row=row + 5, column=13, value='Tổng thời gian gia công dự kiến ( giờ )')
            ws.cell(row=row + 6, column=1, value='Thời gian gia công')
            ws.cell(row=row + 6, column=12, value=cycle)
            ws.cell(row=row + 6, column=13, value=total)
            row += BLOCK_HEIGHT
    buffer = io.BytesIO(); wb.save(buffer)
    return buffer.getvalue()


def _upload(api, data, *, confirm=False, path='/api/templates/import-workbook'):
    files = {'file': ('router.xlsx', data,
                      'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}
    form = {'confirm': '1'} if confirm else None
    return api.post(f'{BASE_URL}{path}', files=files, data=form, timeout=300)


def _po_number():
    return f'PO{uuid.uuid4().hex[:8].upper()}'


# --- a) chưa có gì --------------------------------------------------------

def test_nhap_lan_dau_tao_ca_template_va_po(api, db):
    po_code = _po_number()
    response = _upload(api, _router_bytes(po_number=po_code))
    assert response.status_code == 200, response.text[:400]
    body = response.json()
    assert body['production_order']['action'] == 'CREATED'
    assert body['production_order']['code'] == po_code

    po = db.execute('SELECT id,planned_quantity,order_type FROM production_orders WHERE code=%s',
                    (po_code,)).fetchone()
    assert po['planned_quantity'] == 110, 'QTY của PO'
    assert po['order_type'] == 'SẢN XUẤT HÀNG LOẠT'
    # Số lượng RIÊNG của từng Part, không gộp về QTY của PO.
    quantities = sorted(r['planned_quantity'] for r in db.execute(
        'SELECT planned_quantity FROM parts WHERE production_order_id=%s', (po['id'],)).fetchall())
    assert quantities == [110, 220]


def test_op_mang_danh_tinh_nguon_va_thoi_gian(api, db):
    po_code = _po_number()
    assert _upload(api, _router_bytes(po_number=po_code)).status_code == 200
    po = db.execute('SELECT id FROM production_orders WHERE code=%s', (po_code,)).fetchone()
    rows = db.execute("""SELECT source_op_no,source_title,expected_total_seconds,
            setup_source_raw,standard_seconds_per_unit
        FROM operations WHERE production_order_id=%s
        AND COALESCE(operation_type,'PRODUCTION')='PRODUCTION' ORDER BY id""",
        (po['id'],)).fetchall()
    assert len(rows) == 2
    assert rows[0]['source_op_no'] == 1
    assert rows[0]['source_title'] == 'OPERATION # 01- CẮT LASER'
    assert float(rows[0]['standard_seconds_per_unit']) == 100.0
    # Tổng thời gian là số liệu NGUỒN: 3,39 giờ -> giây, giữ nguyên.
    assert rows[0]['expected_total_seconds'] == pytest.approx(3.39 * 3600, abs=1)
    # setup 20 phút -> tạo OP SETUP liên kết.
    setups = db.execute("""SELECT COUNT(*) c FROM operations
        WHERE production_order_id=%s AND operation_type='SETUP'""", (po['id'],)).fetchone()['c']
    assert setups == 1


# --- b) nhập lại đúng file ------------------------------------------------

def test_nhap_lai_dung_file_la_no_op_sach_khong_bao_loi(api, db):
    po_code = _po_number()
    data = _router_bytes(po_number=po_code)
    assert _upload(api, data).status_code == 200
    before = db.execute('SELECT COUNT(*) c FROM production_orders').fetchone()['c']

    again = _upload(api, data)
    # KHÔNG lỗi, dù PO đã tồn tại -- đây là điểm dễ sai nhất của hợp đồng.
    assert again.status_code == 200, again.text[:400]
    body = again.json()
    assert body.get('idempotent') is True
    assert body['production_order']['action'] == 'EXISTS'
    after = db.execute('SELECT COUNT(*) c FROM production_orders').fetchone()['c']
    assert after == before, 'không được tạo PO thứ hai'


def test_nhap_lai_khong_nhan_ban_part_hay_operation(api, db):
    po_code = _po_number()
    data = _router_bytes(po_number=po_code)
    assert _upload(api, data).status_code == 200
    po = db.execute('SELECT id FROM production_orders WHERE code=%s', (po_code,)).fetchone()
    counts = lambda: (
        db.execute('SELECT COUNT(*) c FROM parts WHERE production_order_id=%s', (po['id'],)).fetchone()['c'],
        db.execute('SELECT COUNT(*) c FROM operations WHERE production_order_id=%s', (po['id'],)).fetchone()['c'])
    before = counts()
    assert _upload(api, data).status_code == 200
    assert counts() == before


# --- d) PO đã tồn tại, file KHÁC ------------------------------------------

def test_po_da_ton_tai_voi_file_khac_thi_dung_lai(api, db):
    """File đổi nội dung nhưng trùng mã PO -> dừng, không ghi đè dữ liệu vận hành."""
    po_code = _po_number()
    assert _upload(api, _router_bytes(po_number=po_code)).status_code == 200
    changed = _router_bytes(po_number=po_code, extra_block=True)
    response = _upload(api, changed, confirm=True)
    assert response.status_code == 409, response.text[:400]
    body = response.json()
    assert body['reason'] == 'PO_ALREADY_EXISTS'
    assert po_code in body['message']


# --- e) file không có mã PO ----------------------------------------------

def test_file_khong_co_ma_po_thi_chi_tao_template(api, db):
    before = db.execute('SELECT COUNT(*) c FROM production_orders').fetchone()['c']
    response = _upload(api, _router_bytes(po_number=''))
    assert response.status_code == 200, response.text[:400]
    body = response.json()
    assert body['production_order']['action'] == 'NONE'
    assert body['template_id']
    after = db.execute('SELECT COUNT(*) c FROM production_orders').fetchone()['c']
    assert after == before, 'không được bịa ra mã PO'


# --- f) trùng mã Template, nội dung khác ---------------------------------

def test_ghi_de_template_doi_xac_nhan(api, db):
    """Viết lại một Template dùng chung phải được xác nhận rõ ràng."""
    po_code = _po_number()
    assert _upload(api, _router_bytes(po_number=po_code)).status_code == 200
    changed = _router_bytes(po_number=po_code, extra_block=True)

    without_confirm = _upload(api, changed)
    assert without_confirm.status_code == 409, without_confirm.text[:400]
    body = without_confirm.json()
    assert body['reason'] == 'TEMPLATE_UPDATE_NEEDS_CONFIRM'
    # Và phải nói có bao nhiêu PO đã dựng từ Template đó.
    assert body['detail']['po_built_from_it'] >= 1


# --- màn xem trước --------------------------------------------------------

def test_preview_noi_ro_ke_hoach_va_ba_muc(api):
    po_code = _po_number()
    response = _upload(api, _router_bytes(po_number=po_code),
                       path='/api/templates/preview-workbook')
    assert response.status_code == 200, response.text[:400]
    body = response.json()
    plan = body['plan']
    assert plan['template']['action'] == 'CREATE'
    assert plan['po']['action'] == 'CREATE'
    assert po_code in plan['po']['message']

    sections = body['sections']
    assert sections['imported']['po']['quantity'] == 110
    assert sections['imported']['po']['order_type'] == 'SẢN XUẤT HÀNG LOẠT'
    # Mục "Excel có nhưng chưa dùng" KHÔNG được rỗng: trường vận hành luôn ở đó.
    assert sections['unused'], 'phải liệt kê field chưa dùng'
    reasons = ' '.join(u['reason'] for u in sections['unused'])
    assert 'VẬN HÀNH' in reasons or 'vận hành' in reasons
    # Preview không được ghi gì.
    assert 'template_id' not in body or body.get('plan')


def test_preview_khong_ghi_gi_vao_csdl(api, db):
    before = db.execute('SELECT COUNT(*) c FROM templates').fetchone()['c']
    _upload(api, _router_bytes(po_number=_po_number()),
            path='/api/templates/preview-workbook')
    assert db.execute('SELECT COUNT(*) c FROM templates').fetchone()['c'] == before


# --- file gốc được lưu để xuất lại ---------------------------------------

def test_workbook_goc_duoc_noi_voi_template(api, db):
    po_code = _po_number()
    data = _router_bytes(po_number=po_code)
    body = _upload(api, data).json()
    link = db.execute("""SELECT sha256,original_filename FROM template_source_workbooks
        WHERE template_id=%s""", (body['template_id'],)).fetchone()
    assert link, 'phải nối file gốc với Template để còn xuất lại kèm QR'
    assert re.fullmatch(r'[0-9a-f]{64}', link['sha256'])

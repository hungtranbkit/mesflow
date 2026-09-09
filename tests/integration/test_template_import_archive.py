"""Every Template import keeps the workbook it came from -- including failures.

`templates.source_workbook` only ever held the uploaded file's NAME, and a
re-import overwrote it. So when TPL-6126 turned out to carry ten duplicate
Operation codes (2026-09-09) there was no way to open the sheet responsible,
and a rejected import left no trace at all.
"""
import io
import uuid

import pytest
from openpyxl import Workbook

from conftest import BASE_URL

pytestmark = pytest.mark.postgres


def _workbook(code, operations):
    """operations: list of (part_code, op_code, op_name)."""
    wb = Workbook()
    meta = wb.active
    meta.title = 'Template'
    meta.append(['template_code', 'template_name', 'product', 'version', 'active'])
    meta.append([code, f'Tên {code}', 'SP thử', '1.0', 1])
    parts = wb.create_sheet('Parts')
    parts.append(['part_code', 'part_name', 'sort_order'])
    for idx, part_code in enumerate(sorted({p for p, _c, _n in operations})):
        parts.append([part_code, f'Part {part_code}', idx])
    ops = wb.create_sheet('Operations')
    ops.append(['part_code', 'operation_code', 'operation_name', 'sort_order'])
    for idx, (part_code, op_code, op_name) in enumerate(operations):
        ops.append([part_code, op_code, op_name, idx])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _upload(api, filename, payload):
    return api.post(f'{BASE_URL}/api/templates/import-workbook',
                    files={'file': (filename, payload,
                                    'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')},
                    timeout=30)


def _cleanup(db, code):
    with db.cursor() as cur:
        cur.execute("""DELETE FROM template_import_events WHERE upper(template_code)=upper(%s)""", (code,))
        cur.execute("""DELETE FROM template_operations WHERE template_id IN
            (SELECT id FROM templates WHERE upper(code)=upper(%s))""", (code,))
        cur.execute("""DELETE FROM template_parts WHERE template_id IN
            (SELECT id FROM templates WHERE upper(code)=upper(%s))""", (code,))
        cur.execute('DELETE FROM templates WHERE upper(code)=upper(%s)', (code,))


def test_successful_import_is_archived_and_downloadable(db, api):
    code = f'TPL-ARC-{uuid.uuid4().hex[:6].upper()}'
    payload = _workbook(code, [('P1', 'P1-OP01', 'CẮT'), ('P1', 'P1-OP02', 'HÀN')])
    try:
        response = _upload(api, 'quy_trinh_goc.xlsx', payload)
        assert response.status_code == 200, response.text
        template_id = response.json()['template_id']

        history = api.get(f'{BASE_URL}/api/templates/{template_id}/import-history', timeout=15)
        assert history.status_code == 200, history.text
        rows = history.json()['items']
        assert len(rows) == 1
        assert rows[0]['original_filename'] == 'quy_trinh_goc.xlsx'
        assert rows[0]['outcome'] == 'CREATED'
        assert (rows[0]['part_count'], rows[0]['operation_count']) == (1, 2)

        # The archived bytes are the bytes that were uploaded.
        download = api.get(f"{BASE_URL}/api/templates/import-history/{rows[0]['id']}/download", timeout=20)
        assert download.status_code == 200
        assert download.content == payload
    finally:
        _cleanup(db, code)


def test_rejected_import_is_archived_with_the_reason(db, api):
    """The whole point: a bad sheet must still be openable afterwards."""
    code = f'TPL-BAD-{uuid.uuid4().hex[:6].upper()}'
    # Same Part, same Operation code twice -- the TPL-6126 shape.
    payload = _workbook(code, [('P1', 'P1-OP01', 'CẮT'), ('P1', 'P1-OP01', 'LÀM NGUỘI')])
    try:
        response = _upload(api, 'file_loi.xlsx', payload)

        history = api.get(f'{BASE_URL}/api/templates/import-history?template_code={code}', timeout=15)
        assert history.status_code == 200, history.text
        rows = history.json()['items']
        assert len(rows) == 1, 'a rejected import must still be recorded'
        entry = rows[0]
        assert entry['original_filename'] == 'file_loi.xlsx'
        if response.status_code >= 400:
            assert entry['outcome'] == 'FAILED'
            assert entry['error_message']
        else:
            # Imported but carrying duplicates: the archive says which ones,
            # so the bad sheet is identifiable either way.
            assert 'P1-OP01' in entry['duplicate_operation_codes']

        download = api.get(f"{BASE_URL}/api/templates/import-history/{entry['id']}/download", timeout=20)
        assert download.status_code == 200
        assert download.content == payload
    finally:
        _cleanup(db, code)


def test_reimport_keeps_every_attempt_and_stores_identical_bytes_once(db, api):
    code = f'TPL-RE-{uuid.uuid4().hex[:6].upper()}'
    payload = _workbook(code, [('P1', 'P1-OP01', 'CẮT')])
    try:
        first = _upload(api, 'lan_1.xlsx', payload)
        assert first.status_code == 200, first.text
        second = _upload(api, 'lan_2.xlsx', payload)
        assert second.status_code == 200, second.text
        template_id = second.json()['template_id']

        rows = api.get(f'{BASE_URL}/api/templates/{template_id}/import-history', timeout=15).json()['items']
        assert [r['original_filename'] for r in rows] == ['lan_2.xlsx', 'lan_1.xlsx'], 'newest first'
        assert rows[0]['outcome'] == 'REPLACED' and rows[1]['outcome'] == 'CREATED'
        # Identical content is stored once, referenced twice.
        assert rows[0]['sha256'] == rows[1]['sha256']
        stored = db.execute('SELECT COUNT(*) n FROM template_import_blobs WHERE sha256=%s',
                            (rows[0]['sha256'],)).fetchone()['n']
        assert stored == 1
    finally:
        _cleanup(db, code)

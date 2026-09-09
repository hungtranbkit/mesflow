"""Cụm Excel/Template -- bốn lỗi P1 tìm ra trong audit kiến trúc 2026-09-09.

Điểm chung của cả bốn: Excel là đường ghi duy nhất KHÔNG đi qua repository, nên
mọi bất biến mà phần còn lại của hệ thống giữ được đều phải tự viết lại ở đây --
và bốn chỗ này đã viết thiếu.

  P1-a  instantiate() nối dòng vật tư theo mã OP toàn cục. Từ khi mã OP chỉ
        unique trong phạm vi Part (2026-09-09), hai Part cùng có 'OP01' thì OP
        của Part B hút nguyên liệu từ OP01 của Part A.
  P1-b  export ghi cả OP phụ; import lại dựng chúng thành PRODUCTION không cha,
        lọt vào operation_count nên PO không bao giờ COMPLETED.
  P1-c  mode=replace DELETE không WHERE -- xoá cấu trúc Operation của MỌI PO
        trong cơ sở dữ liệu, không chỉ PO có trong file.
  P1-d  import lại Template = xoá rồi chèn, làm mất sạch những cột chỉ cấu hình
        được trên web (setup, dòng vật tư, bản vẽ) vì workbook không có chúng.
"""
import io
import uuid

import pytest
from openpyxl import Workbook

from conftest import BASE_URL

pytestmark = pytest.mark.postgres

HEADERS = ['operation_id', 'product', 'po', 'part', 'part_order', 'operation_name',
           'drawing', 'plan', 'done', 'defect', 'status', 'qr']


def _xlsx(sheets):
    """sheets: [(title, [row, ...]), ...] -> BytesIO sẵn sàng upload."""
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets:
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _upload(api, url, buf, *, filename='wb.xlsx', data=None):
    return api.post(url, files={'file': (filename, buf.getvalue(),
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')},
        data=data or {}, timeout=30)


# ---------------------------------------------------------------- P1-a


def _flow_template(db, suffix):
    """Hai Part, mỗi Part có OP01 -> OP02, OP02 hút nguyên liệu từ OP01.

    Cả hai Part dùng CÙNG mã OP -- hợp lệ từ khi mã OP unique theo Part. Đây
    chính là hình dạng làm lộ bug: bản đồ nguồn chỉ khoá theo mã.
    """
    with db.cursor() as cur:
        cur.execute("""INSERT INTO templates(code,name,product,version,active)
            VALUES(%s,'Tpl flow','SP','1.0',true) RETURNING id""", (f'TPL-FLOW-{suffix}',))
        template_id = cur.fetchone()['id']
        for idx, part_code in enumerate(('PA', 'PB')):
            cur.execute("""INSERT INTO template_parts(template_id,code,name,sort_order)
                VALUES(%s,%s,%s,%s) RETURNING id""", (template_id, part_code, f'Part {part_code}', idx))
            part_id = cur.fetchone()['id']
            cur.execute("""INSERT INTO template_operations(template_id,part_id,code,name,sort_order,
                    standard_seconds_per_unit) VALUES(%s,%s,'OP01','Cắt',0,10)""", (template_id, part_id))
            cur.execute("""INSERT INTO template_operations(template_id,part_id,code,name,sort_order,
                    standard_seconds_per_unit,input_flow_enabled,input_source_code,input_source_kind)
                VALUES(%s,%s,'OP02','Hàn',1,10,true,'OP01','GOOD')""", (template_id, part_id))
    return template_id


def _drop_template(db, template_id, po_code=None):
    with db.cursor() as cur:
        if po_code:
            cur.execute("""DELETE FROM operations WHERE production_order_id IN
                (SELECT id FROM production_orders WHERE code=%s)""", (po_code,))
            cur.execute("""DELETE FROM parts WHERE production_order_id IN
                (SELECT id FROM production_orders WHERE code=%s)""", (po_code,))
            cur.execute('DELETE FROM production_orders WHERE code=%s', (po_code,))
        cur.execute('DELETE FROM template_operations WHERE template_id=%s', (template_id,))
        cur.execute('DELETE FROM template_parts WHERE template_id=%s', (template_id,))
        cur.execute('DELETE FROM templates WHERE id=%s', (template_id,))


def test_input_flow_source_resolves_within_the_same_part(db, api):
    """OP02 của Part B phải trỏ về OP01 của Part B, không phải của Part A."""
    suffix = uuid.uuid4().hex[:8].upper()
    template_id = _flow_template(db, suffix)
    po_code = f'PO-FLOW-{suffix}'
    try:
        created = api.post(f'{BASE_URL}/api/templates/{template_id}/instantiate',
                           json={'code': po_code, 'planned_quantity': 10}, timeout=20)
        assert created.status_code in (200, 201), created.text
        rows = db.execute("""SELECT p.code part_code,o.code,o.id,o.input_source_operation_id src
            FROM operations o JOIN parts p ON p.id=o.part_id
            JOIN production_orders po ON po.id=o.production_order_id
            WHERE po.code=%s ORDER BY p.code,o.sort_order""", (po_code,)).fetchall()
        by_key = {(r['part_code'], r['code'].rsplit('-', 1)[-1]): r for r in rows}
        assert len(by_key) == 4, rows
        for part_code in ('PA', 'PB'):
            op1, op2 = by_key[(part_code, 'OP01')], by_key[(part_code, 'OP02')]
            assert op2['src'] == op1['id'], (
                f'{part_code}/OP02 hút nguyên liệu từ OP01 của Part khác: {rows}')
    finally:
        _drop_template(db, template_id, po_code)


# ---------------------------------------------------------------- P1-b


def _po_with_support_ops(db, suffix):
    """Một PO có 1 OP sản xuất, 1 SETUP con của nó, và 1 bàn SỬA HÀNG."""
    with db.cursor() as cur:
        cur.execute("""INSERT INTO production_orders(code,product,planned_quantity,status)
            VALUES(%s,'SP',100,'IN_PROGRESS') RETURNING id""", (f'PO-SUP-{suffix}',))
        po_id = cur.fetchone()['id']
        cur.execute("""INSERT INTO parts(production_order_id,code,name,sort_order,active)
            VALUES(%s,'PA','Thân',0,true) RETURNING id""", (po_id,))
        part_id = cur.fetchone()['id']
        cur.execute("""INSERT INTO operations(production_order_id,part_id,code,name,status,qr,sort_order)
            VALUES(%s,%s,%s,'Cắt phôi','PLANNED',%s,0) RETURNING id""",
            (po_id, part_id, f'SUP-{suffix}-OP1', f'WF|OP|SUP-{suffix}-OP1'))
        main_id = cur.fetchone()['id']
        cur.execute("""INSERT INTO operations(production_order_id,part_id,code,name,status,qr,sort_order,
                operation_type,parent_operation_id)
            VALUES(%s,%s,%s,'Setup Cắt phôi','PLANNED',%s,1,'SETUP',%s) RETURNING id""",
            (po_id, part_id, f'SUP-{suffix}-OP1-SU', f'WF|OP|SUP-{suffix}-OP1-SU', main_id))
        setup_id = cur.fetchone()['id']
        cur.execute("""INSERT INTO operations(production_order_id,part_id,code,name,status,qr,sort_order,
                operation_type) VALUES(%s,%s,%s,'SỬA HÀNG','PLANNED',%s,2,'REWORK') RETURNING id""",
            (po_id, part_id, f'SUP-{suffix}-RW', f'WF|OP|SUP-{suffix}-RW'))
        rework_id = cur.fetchone()['id']
    return {'po_id': po_id, 'po_code': f'PO-SUP-{suffix}', 'part_id': part_id,
            'main_id': main_id, 'setup_id': setup_id, 'rework_id': rework_id}


def _drop_po(db, po_id):
    with db.cursor() as cur:
        cur.execute('DELETE FROM operations WHERE production_order_id=%s', (po_id,))
        cur.execute('DELETE FROM parts WHERE production_order_id=%s', (po_id,))
        cur.execute('DELETE FROM production_orders WHERE id=%s', (po_id,))


def test_export_omits_support_operations(db, api):
    """SETUP và SỬA HÀNG không phải bước routing -- không xuất ra file."""
    from openpyxl import load_workbook

    suffix = uuid.uuid4().hex[:8].upper()
    graph = _po_with_support_ops(db, suffix)
    try:
        response = api.get(f'{BASE_URL}/api/operations/export.xlsx', timeout=60)
        assert response.status_code == 200, response.text[:300]
        ws = load_workbook(io.BytesIO(response.content), read_only=True)['Operations']
        codes = {str(row[0]) for row in ws.iter_rows(min_row=2, values_only=True) if row and row[0]}
        assert f'SUP-{suffix}-OP1' in codes, 'OP sản xuất phải có trong file'
        assert f'SUP-{suffix}-OP1-SU' not in codes, 'SETUP lọt vào file export'
        assert f'SUP-{suffix}-RW' not in codes, 'SỬA HÀNG lọt vào file export'
    finally:
        _drop_po(db, graph['po_id'])


def test_import_refuses_to_rewrite_a_support_operation(db, api):
    """File xuất TRƯỚC bản vá vẫn còn dòng OP phụ -- phải chặn, không ghi đè.

    Nếu để lọt, UPDATE sẽ giữ nguyên operation_type='SETUP' nhưng đổi Part/tên
    theo file, và OP đó bắt đầu đếm như một bước routing thật.
    """
    suffix = uuid.uuid4().hex[:8].upper()
    graph = _po_with_support_ops(db, suffix)
    try:
        buf = _xlsx([('Operations', [HEADERS, [
            f'SUP-{suffix}-OP1-SU', 'SP', graph['po_code'], 'Thân', 0,
            'Đổi thành bước thật', '', 100, 0, 0, 'PLANNED', '']])])
        response = _upload(api, f'{BASE_URL}/api/operations/import', buf, data={'mode': 'merge'})
        assert response.status_code >= 400, response.text
        assert 'OP phụ' in response.json()['message']
        row = db.execute("""SELECT name,COALESCE(operation_type,'PRODUCTION') operation_type,
            parent_operation_id FROM operations WHERE id=%s""", (graph['setup_id'],)).fetchone()
        assert row['operation_type'] == 'SETUP'
        assert row['parent_operation_id'] == graph['main_id']
        assert row['name'] == 'Setup Cắt phôi', 'OP phụ bị Excel ghi đè'
    finally:
        _drop_po(db, graph['po_id'])


# ---------------------------------------------------------------- P1-c


def test_replace_only_touches_the_pos_named_in_the_file(db, api):
    """Đây là bug nguy hiểm nhất trong cụm: Replace từng xoá cả hệ thống."""
    suffix = uuid.uuid4().hex[:8].upper()
    target = _po_with_support_ops(db, f'T{suffix}')
    bystander = _po_with_support_ops(db, f'B{suffix}')
    try:
        buf = _xlsx([('Operations', [HEADERS, [
            f'T{suffix}-NEW-OP1', 'SP', target['po_code'], 'Thân', 0,
            'Cắt phôi', '', 100, 0, 0, 'PLANNED', '']])])
        response = _upload(api, f'{BASE_URL}/api/operations/import', buf, data={'mode': 'replace'})
        assert response.status_code == 200, response.text

        survivors = db.execute('SELECT code FROM operations WHERE production_order_id=%s ORDER BY code',
                               (bystander['po_id'],)).fetchall()
        assert len(survivors) == 3, f'PO ngoài file bị Replace xoá mất: {survivors}'
        parts_left = db.execute('SELECT COUNT(*) c FROM parts WHERE production_order_id=%s',
                                (bystander['po_id'],)).fetchone()['c']
        assert parts_left == 1, 'Part của PO ngoài file bị xoá'

        replaced = db.execute('SELECT code FROM operations WHERE production_order_id=%s',
                              (target['po_id'],)).fetchall()
        assert [r['code'] for r in replaced] == [f'T{suffix}-NEW-OP1'], replaced
    finally:
        _drop_po(db, target['po_id'])
        _drop_po(db, bystander['po_id'])


def test_replace_with_no_po_code_in_file_is_rejected(db, api):
    """Không suy ra được phạm vi thì không được đoán -- càng không được xoá hết."""
    buf = _xlsx([('Operations', [HEADERS, [
        'NO-PO-OP1', 'SP', '', 'Thân', 0, 'Cắt phôi', '', 100, 0, 0, 'PLANNED', '']])])
    response = _upload(api, f'{BASE_URL}/api/operations/import', buf, data={'mode': 'replace'})
    assert response.status_code >= 400, response.text


# ---------------------------------------------------------------- P1-d


TEMPLATE_SHEETS = lambda code, ops: [
    ('Template', [['template_code', 'template_name', 'product', 'version', 'active'],
                  [code, 'Tpl giữ cấu hình', 'SP', '1.0', 1]]),
    ('Parts', [['part_code', 'part_name', 'sort_order'], ['PA', 'Thân', 0]]),
    ('Operations', [['part_code', 'operation_code', 'operation_name', 'equipment_code',
                     'cycle_time_value', 'cycle_time_unit', 'sort_order']] + ops),
]


def test_reimport_keeps_web_only_configuration(db, api):
    """Cấu hình chỉ có trên web phải sống sót qua một lần import lại.

    Kịch bản thật: kỹ sư nhập file, cấu hình hướng dẫn setup và dòng vật tư cho
    từng OP trên web, rồi đồng nghiệp upload lại đúng file đó sau khi sửa một
    chữ trong tên OP. Trước bản vá: mất sạch, kèm thông báo "Đã cập nhật".
    """
    code = f'TPL-KEEP-{uuid.uuid4().hex[:8].upper()}'
    template_id = None
    try:
        first = _upload(api, f'{BASE_URL}/api/templates/import-workbook',
            _xlsx(TEMPLATE_SHEETS(code, [
                ['PA', 'OP01', 'Cắt', 'EQ1', 10, 'second', 0],
                ['PA', 'OP02', 'Hàn', 'EQ2', 10, 'second', 1]])))
        assert first.status_code == 200, first.text
        template_id = first.json()['template_id']

        # Cấu hình trên web: setup cho OP01, dòng vật tư cho OP02, bản vẽ Part.
        with db.cursor() as cur:
            cur.execute("""UPDATE template_operations SET requires_setup=true,
                expected_setup_minutes=25,setup_note='Căn dao trước khi chạy'
                WHERE template_id=%s AND code='OP01'""", (template_id,))
            cur.execute("""UPDATE template_operations SET input_flow_enabled=true,
                input_source_code='OP01',input_source_kind='GOOD',defects_consume_input=false,
                repair_cycle_time_seconds_per_unit=45
                WHERE template_id=%s AND code='OP02'""", (template_id,))
            cur.execute("""UPDATE template_parts SET drawing_path='/dwg/than.pdf'
                WHERE template_id=%s AND code='PA'""", (template_id,))

        # Đồng nghiệp sửa tên một OP rồi upload lại cùng mã Template.
        second = _upload(api, f'{BASE_URL}/api/templates/import-workbook',
            _xlsx(TEMPLATE_SHEETS(code, [
                ['PA', 'OP01', 'Cắt phôi', 'EQ1', 10, 'second', 0],
                ['PA', 'OP02', 'Hàn', 'EQ2', 10, 'second', 1]])))
        assert second.status_code == 200, second.text
        assert second.json()['template_id'] == template_id, 'import lại phải cập nhật, không fork'

        rows = {r['code']: r for r in db.execute("""SELECT code,name,requires_setup,
            expected_setup_minutes,setup_note,input_flow_enabled,input_source_code,
            input_source_kind,defects_consume_input,repair_cycle_time_seconds_per_unit
            FROM template_operations WHERE template_id=%s""", (template_id,)).fetchall()}
        assert rows['OP01']['name'] == 'Cắt phôi', 'file vẫn phải ghi đè được cột nó sở hữu'
        assert rows['OP01']['requires_setup'] is True
        assert rows['OP01']['expected_setup_minutes'] == 25
        assert rows['OP01']['setup_note'] == 'Căn dao trước khi chạy'
        assert rows['OP02']['input_flow_enabled'] is True
        assert rows['OP02']['input_source_code'] == 'OP01'
        assert rows['OP02']['defects_consume_input'] is False
        assert float(rows['OP02']['repair_cycle_time_seconds_per_unit']) == 45.0
        drawing = db.execute('SELECT drawing_path FROM template_parts WHERE template_id=%s',
                             (template_id,)).fetchone()['drawing_path']
        assert drawing == '/dwg/than.pdf', 'bản vẽ Part cũng chỉ cấu hình được trên web'
    finally:
        if template_id:
            _drop_template(db, template_id)


def test_reimport_does_not_leak_configuration_onto_a_new_operation(db, api):
    """Khớp theo (mã Part, mã OP). OP mới trong file phải nhận mặc định sạch.

    Nếu bản đồ giữ-lại chỉ khoá theo mã OP thì OP mới trùng mã ở Part khác sẽ
    thừa hưởng cấu hình không phải của nó -- đúng họ bug với P1-a.
    """
    code = f'TPL-NEW-{uuid.uuid4().hex[:8].upper()}'
    template_id = None
    try:
        first = _upload(api, f'{BASE_URL}/api/templates/import-workbook',
            _xlsx(TEMPLATE_SHEETS(code, [['PA', 'OP01', 'Cắt', 'EQ1', 10, 'second', 0]])))
        assert first.status_code == 200, first.text
        template_id = first.json()['template_id']
        with db.cursor() as cur:
            cur.execute("""UPDATE template_operations SET requires_setup=true,
                expected_setup_minutes=25,setup_note='Của PA/OP01'
                WHERE template_id=%s AND code='OP01'""", (template_id,))

        # Thêm Part PB, cũng có OP01 -- hợp lệ, mã OP unique theo Part.
        sheets = TEMPLATE_SHEETS(code, [['PA', 'OP01', 'Cắt', 'EQ1', 10, 'second', 0],
                                        ['PB', 'OP01', 'Cắt', 'EQ1', 10, 'second', 0]])
        sheets[1] = ('Parts', [['part_code', 'part_name', 'sort_order'],
                               ['PA', 'Thân', 0], ['PB', 'Nắp', 1]])
        second = _upload(api, f'{BASE_URL}/api/templates/import-workbook', _xlsx(sheets))
        assert second.status_code == 200, second.text

        rows = {(r['part_code'], r['code']): r for r in db.execute("""SELECT p.code part_code,
            o.code,o.requires_setup,o.expected_setup_minutes,o.setup_note
            FROM template_operations o JOIN template_parts p ON p.id=o.part_id
            WHERE o.template_id=%s""", (template_id,)).fetchall()}
        assert rows[('PA', 'OP01')]['requires_setup'] is True
        assert rows[('PA', 'OP01')]['setup_note'] == 'Của PA/OP01'
        assert rows[('PB', 'OP01')]['requires_setup'] is False, 'cấu hình rò sang Part khác'
        assert rows[('PB', 'OP01')]['expected_setup_minutes'] is None
        assert rows[('PB', 'OP01')]['setup_note'] == ''
    finally:
        if template_id:
            _drop_template(db, template_id)


def test_reimport_drops_configuration_when_the_operation_code_changes(db, api):
    """Đổi mã OP = OP khác. Mất cấu hình cũ là đúng, và không được vỡ."""
    code = f'TPL-REN-{uuid.uuid4().hex[:8].upper()}'
    template_id = None
    try:
        first = _upload(api, f'{BASE_URL}/api/templates/import-workbook',
            _xlsx(TEMPLATE_SHEETS(code, [['PA', 'OP01', 'Cắt', 'EQ1', 10, 'second', 0]])))
        assert first.status_code == 200, first.text
        template_id = first.json()['template_id']
        with db.cursor() as cur:
            cur.execute("""UPDATE template_operations SET requires_setup=true,setup_note='Cũ'
                WHERE template_id=%s AND code='OP01'""", (template_id,))

        second = _upload(api, f'{BASE_URL}/api/templates/import-workbook',
            _xlsx(TEMPLATE_SHEETS(code, [['PA', 'OP99', 'Cắt', 'EQ1', 10, 'second', 0]])))
        assert second.status_code == 200, second.text
        rows = db.execute("""SELECT code,requires_setup,setup_note,input_source_kind
            FROM template_operations WHERE template_id=%s""", (template_id,)).fetchall()
        assert [r['code'] for r in rows] == ['OP99']
        assert rows[0]['requires_setup'] is False
        assert rows[0]['setup_note'] == ''
        # Các cột NOT NULL vẫn phải có giá trị hợp lệ khi chèn tường minh.
        assert rows[0]['input_source_kind'] == 'GOOD'
    finally:
        if template_id:
            _drop_template(db, template_id)

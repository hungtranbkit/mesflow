"""Operation codes are unique per Part, not per Template.

Design agreed 2026-09-09: an Operation code may repeat across different Parts
-- (Part code, Operation code) is the business key -- while Part codes stay
unique within a Template. operations.code and operations.qr remain globally
unique, so the generated code has to carry the Part whenever the Operation
code does not already name it (operation_code_suffix()).
"""
import io
import uuid

import pytest
from openpyxl import Workbook

from conftest import BASE_URL

pytestmark = pytest.mark.postgres


def _template(db, *, parts_ops):
    """parts_ops: {part_code: [op_code, ...]} inserted directly."""
    suffix = uuid.uuid4().hex[:8].upper()
    with db.cursor() as cur:
        cur.execute("""INSERT INTO templates(code,name,product,version,active)
            VALUES(%s,'Tpl scope','SP','1.0',true) RETURNING id""", (f'TPL-SCOPE-{suffix}',))
        template_id = cur.fetchone()['id']
        for p_idx, (part_code, ops) in enumerate(parts_ops.items()):
            cur.execute("""INSERT INTO template_parts(template_id,code,name,sort_order)
                VALUES(%s,%s,%s,%s) RETURNING id""", (template_id, part_code, f'Part {part_code}', p_idx))
            part_id = cur.fetchone()['id']
            for o_idx, op_code in enumerate(ops):
                cur.execute("""INSERT INTO template_operations(template_id,part_id,code,name,sort_order,
                        standard_seconds_per_unit,repair_cycle_time_seconds_per_unit)
                    VALUES(%s,%s,%s,%s,%s,10,0)""",
                    (template_id, part_id, op_code, f'OP {op_code}', o_idx))
    return template_id, suffix


def _create_po(api, template_id, code):
    return api.post(f'{BASE_URL}/api/templates/{template_id}/instantiate',
                    json={'code': code, 'planned_quantity': 10}, timeout=20)


def _cleanup(db, template_id, po_code=None):
    with db.cursor() as cur:
        if po_code:
            cur.execute("""DELETE FROM operations WHERE production_order_id IN
                (SELECT id FROM production_orders WHERE code=%s)""", (po_code,))
            cur.execute("""DELETE FROM parts WHERE production_order_id IN
                (SELECT id FROM production_orders WHERE code=%s)""", (po_code,))
            cur.execute('DELETE FROM production_orders WHERE code=%s', (po_code,))
        cur.execute('DELETE FROM template_operations WHERE template_id=%s', (template_id,))
        cur.execute('DELETE FROM template_parts WHERE template_id=%s', (template_id,))
        cur.execute('DELETE FROM templates WHERE id=%s', (template_id,))


def test_same_operation_code_in_two_parts_is_allowed(db, api):
    """The feature itself: OP01 under two Parts, cloned without collision."""
    template_id, _s = _template(db, parts_ops={'PA': ['OP01', 'OP02'], 'PB': ['OP01', 'OP02']})
    po_code = f'PO-SCOPE-{uuid.uuid4().hex[:6].upper()}'
    try:
        response = _create_po(api, template_id, po_code)
        assert response.status_code in (200, 201), response.text
        rows = db.execute("""SELECT p.code part_code,o.code,o.qr FROM operations o
            JOIN parts p ON p.id=o.part_id JOIN production_orders po ON po.id=o.production_order_id
            WHERE po.code=%s ORDER BY p.code,o.sort_order""", (po_code,)).fetchall()
        assert [(r['part_code'], r['code']) for r in rows] == [
            ('PA', f'{po_code}-PA-OP01'), ('PA', f'{po_code}-PA-OP02'),
            ('PB', f'{po_code}-PB-OP01'), ('PB', f'{po_code}-PB-OP02')]
        # operations.qr is globally unique too -- it must not collide either.
        assert len({r['qr'] for r in rows}) == 4
    finally:
        _cleanup(db, template_id, po_code)


def test_code_that_already_names_its_part_is_left_alone(db, api):
    """Existing convention keeps producing exactly the codes it does today."""
    template_id, _s = _template(db, parts_ops={'KM-3172005-08': ['KM-3172005-08-OP01',
                                                                'KM-3172005-08-OP02']})
    po_code = f'PO-KEEP-{uuid.uuid4().hex[:6].upper()}'
    try:
        assert _create_po(api, template_id, po_code).status_code in (200, 201)
        codes = [r['code'] for r in db.execute("""SELECT o.code FROM operations o
            JOIN production_orders po ON po.id=o.production_order_id
            WHERE po.code=%s ORDER BY o.sort_order""", (po_code,)).fetchall()]
        assert codes == [f'{po_code}-KM-3172005-08-OP01', f'{po_code}-KM-3172005-08-OP02']
    finally:
        _cleanup(db, template_id, po_code)


def test_duplicate_within_one_part_is_still_refused(db, api):
    """The TPL-6126 shape stays invalid -- one Part cannot have two OP02."""
    template_id, _s = _template(db, parts_ops={'PA': ['OP02', 'OP02']})
    po_code = f'PO-DUP1-{uuid.uuid4().hex[:6].upper()}'
    try:
        response = _create_po(api, template_id, po_code)
        assert response.status_code == 409, response.text
        message = response.json()['message']
        assert 'cùng một Part' in message
        assert db.execute('SELECT COUNT(*) n FROM production_orders WHERE code=%s',
                          (po_code,)).fetchone()['n'] == 0
    finally:
        _cleanup(db, template_id, po_code)


def test_mixed_conventions_inside_one_part_are_refused(db, api):
    """`OP01` and `PA-OP01` in the SAME Part generate one identical code.

    The pair (Part, Operation code) differs, so a naive per-pair check would
    pass this and then die on operations_code_key. Validation therefore checks
    the code that will actually be generated.
    """
    template_id, _s = _template(db, parts_ops={'PA': ['OP01', 'PA-OP01']})
    po_code = f'PO-MIX-{uuid.uuid4().hex[:6].upper()}'
    try:
        response = _create_po(api, template_id, po_code)
        assert response.status_code == 409, (
            f'both rows generate {po_code}-PA-OP01; got {response.status_code}: {response.text}')
        assert db.execute('SELECT COUNT(*) n FROM production_orders WHERE code=%s',
                          (po_code,)).fetchone()['n'] == 0
    finally:
        _cleanup(db, template_id, po_code)


def _workbook(code, rows):
    wb = Workbook()
    meta = wb.active; meta.title = 'Template'
    meta.append(['template_code', 'template_name', 'product', 'version', 'active'])
    meta.append([code, f'Tên {code}', 'SP', '1.0', 1])
    parts = wb.create_sheet('Parts'); parts.append(['part_code', 'part_name', 'sort_order'])
    for idx, part_code in enumerate(sorted({p for p, _c in rows})):
        parts.append([part_code, f'Part {part_code}', idx])
    ops = wb.create_sheet('Operations')
    ops.append(['part_code', 'operation_code', 'operation_name', 'sort_order'])
    for idx, (part_code, op_code) in enumerate(rows):
        ops.append([part_code, op_code, f'OP {op_code}', idx])
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def _upload(api, filename, payload):
    return api.post(f'{BASE_URL}/api/templates/import-workbook',
                    files={'file': (filename, payload,
                                    'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')},
                    timeout=30)


def _drop_template(db, code):
    with db.cursor() as cur:
        cur.execute('DELETE FROM template_import_events WHERE upper(template_code)=upper(%s)', (code,))
        cur.execute("""DELETE FROM template_operations WHERE template_id IN
            (SELECT id FROM templates WHERE upper(code)=upper(%s))""", (code,))
        cur.execute("""DELETE FROM template_parts WHERE template_id IN
            (SELECT id FROM templates WHERE upper(code)=upper(%s))""", (code,))
        cur.execute('DELETE FROM templates WHERE upper(code)=upper(%s)', (code,))


def test_excel_import_accepts_the_same_code_in_different_parts(db, api):
    code = f'TPL-XOK-{uuid.uuid4().hex[:6].upper()}'
    try:
        response = _upload(api, 'ok.xlsx', _workbook(code, [('PA', 'OP01'), ('PB', 'OP01')]))
        assert response.status_code == 200, response.text
        assert response.json()['operation_count'] == 2
    finally:
        _drop_template(db, code)


def test_excel_import_refuses_a_duplicate_inside_one_part(db, api):
    """Excel used to write straight through -- this is how TPL-6126 got in."""
    code = f'TPL-XBAD-{uuid.uuid4().hex[:6].upper()}'
    try:
        response = _upload(api, 'bad.xlsx', _workbook(code, [('PA', 'OP01'), ('PA', 'OP01')]))
        assert response.status_code >= 400, response.text
        message = response.json().get('message') or ''
        assert 'cùng một Part' in message
        assert 'OP01' in message
        # Points at the rows to edit, not just the code: the person reading
        # this has the workbook open.
        assert 'dòng' in message
        assert "sheet 'Operations'" in message, f'the sheet must be named: {message}'
        assert '2' in message and '3' in message, f'expected the two Operations rows: {message}'
        assert 'Part PA' in message
        assert db.execute('SELECT COUNT(*) n FROM templates WHERE upper(code)=upper(%s)',
                          (code,)).fetchone()['n'] == 0, 'nothing may be written'
        # ...and the rejected workbook is still archived for inspection.
        assert db.execute('SELECT COUNT(*) n FROM template_import_events WHERE upper(template_code)=upper(%s)',
                          (code,)).fetchone()['n'] == 1
    finally:
        _drop_template(db, code)


def _go_router_workbook(sheets):
    """The other real layout: one sheet per Part, `OPERATION # NN - name` rows.

    This is the shape the customer's own routing files use, and the one that
    produced TPL-6126 -- two blocks numbered the same inside one sheet.
    """
    wb = Workbook()
    wb.remove(wb.active)
    for title, ops in sheets.items():
        ws = wb.create_sheet(title[:31])
        ws.append(['PO NUMBER:', 'PO-GO-1'])
        ws.append(['QTY:', 10])
        ws.append(['MÃ BẢN VẼ', title])
        for seq, name in ops:
            ws.append([f'OPERATION # {seq:02d} - {name}'])
            ws.append(['', 'chi tiết'])
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def test_go_router_workbook_imports_duplicate_source_op_numbers_and_keeps_them_raw(db, api):
    """Trùng SỐ OP trong một sheet là dữ liệu thật, không phải lỗi đánh máy.

    Quyết định 2026-09-09 là TỪ CHỐI cả file và nêu sheet + dòng. Điều đó đã bị
    thay sau khi audit chính file của xưởng: 10 chỗ trùng số ấy là những công
    đoạn KHÁC NHAU dùng chung một số (OPERATION # 02 vừa là CHAMFER LỖ vừa là
    LÀM NGUỘI; OPERATION # 01 vừa là TIỆN BƯỚC 1 vừa là TIỆN BƯỚC 2). Từ chối
    file là vứt dữ liệu lộ trình thật; đổi số của khách là bịa ra một mã không
    có trên giấy.

    Hợp đồng hiện tại: nhập ĐỦ, giữ NGUYÊN số và tiêu đề gốc, mã nội bộ sinh
    duy nhất, danh tính canonical là operation.id, và mỗi chỗ trùng là một
    cảnh báo hiện trên màn xem trước -- không im lặng.
    """
    payload = _go_router_workbook({'KM-3172005-08': [(1, 'CẮT LASER'), (2, 'CHAMFER LỖ'),
                                                     (2, 'LÀM NGUỘI')]})
    response = _upload(api, 'go_router.xlsx', payload)
    assert response.status_code == 200, response.text[:400]
    template_id = response.json()['template_id']

    rows = db.execute("""SELECT o.code,o.source_op_no,o.source_title
        FROM template_operations o JOIN template_parts p ON p.id=o.part_id
        WHERE o.template_id=%s ORDER BY o.sort_order""", (template_id,)).fetchall()
    assert len(rows) == 3, 'không được mất Operation nào'
    # Số gốc giữ verbatim -- hai công đoạn cùng mang số 2.
    assert [r['source_op_no'] for r in rows] == [1, 2, 2]
    assert rows[1]['source_title'].endswith('CHAMFER LỖ')
    assert rows[2]['source_title'].endswith('LÀM NGUỘI')
    # Mã nội bộ thì phải duy nhất, nếu không hai dòng đụng nhau trong CSDL.
    codes = [r['code'] for r in rows]
    assert len(codes) == len(set(codes)), codes


"""Workbook Operation: danh tính đi bằng id, mã đi bằng mắt người.

Cột `operation_id` trong file KHÔNG phải operations.id -- nó là MÃ NGHIỆP VỤ,
và cái tên đó đã nằm trong file của khách từ lâu nên không đổi được. Danh tính
thật đi ở một cột mới, `operation_row_id`.

Điều phải giữ bằng mọi giá: file CŨ (không có cột mới) vẫn nhập được y như
trước. _parse_operations_sheet map theo tên cột, nên cột thiếu chỉ là ô rỗng --
nhưng "đáng lẽ phải thế" không phải bằng chứng, nên có bài test cho nó.
"""
from __future__ import annotations

import io
import os
import uuid

import pytest
from openpyxl import Workbook, load_workbook

from conftest import BASE_URL

pytestmark = pytest.mark.postgres

LEGACY_HEADERS = ['operation_id', 'product', 'po', 'part', 'part_order', 'operation_name',
                  'drawing', 'plan', 'done', 'defect', 'status', 'qr']


def _suffix() -> str:
    return uuid.uuid4().hex[:8].upper()


@pytest.fixture
def one_operation(db):
    sfx = _suffix()
    ids = {'suffix': sfx, 'code': f'XL-{sfx}', 'po_code': f'PO-XL-{sfx}'}
    with db.cursor() as cur:
        cur.execute("""INSERT INTO production_orders(code,product,planned_quantity,status)
            VALUES(%s,'SP',100,'IN_PROGRESS') RETURNING id""", (ids['po_code'],))
        ids['po_id'] = cur.fetchone()['id']
        cur.execute("INSERT INTO parts(production_order_id,code,name,sort_order) VALUES(%s,%s,%s,0) RETURNING id",
                    (ids['po_id'], f'PXL-{sfx}', f'Part {sfx}'))
        ids['part_id'] = cur.fetchone()['id']
        ids['part_name'] = f'Part {sfx}'
        cur.execute("""INSERT INTO operations(production_order_id,part_id,code,name,status,qr,sort_order)
            VALUES(%s,%s,%s,'Công đoạn','PLANNED',%s,0) RETURNING id""",
            (ids['po_id'], ids['part_id'], ids['code'], f"WF|OP|{ids['code']}"))
        ids['op_id'] = cur.fetchone()['id']
    yield ids
    with db.cursor() as cur:
        cur.execute("DELETE FROM operations WHERE production_order_id=%s", (ids['po_id'],))
        cur.execute("DELETE FROM parts WHERE production_order_id=%s", (ids['po_id'],))
        cur.execute("DELETE FROM production_orders WHERE id=%s", (ids['po_id'],))


def _upload(api, rows, headers):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Operations'
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    buf = io.BytesIO()
    workbook.save(buf)
    buf.seek(0)
    base = os.environ.get('MESFLOW_BASE_URL', BASE_URL).rstrip('/')
    return api.post(f'{base}/api/operations/import',
                    files={'file': ('ops.xlsx', buf,
                                    'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')},
                    data={'mode': 'merge'}, timeout=30)


def test_export_carries_the_row_id_without_dropping_any_legacy_column(api, one_operation):
    base = os.environ.get('MESFLOW_BASE_URL', BASE_URL).rstrip('/')
    response = api.get(f'{base}/api/operations/export.xlsx', timeout=60)
    assert response.status_code == 200, response.text[:200]
    sheet = load_workbook(io.BytesIO(response.content), data_only=True)['Operations']
    header = [c.value for c in next(sheet.iter_rows(max_row=1))]
    for legacy in LEGACY_HEADERS:
        assert legacy in header, f'xuất file đã mất cột cũ {legacy}: {header}'
    assert 'operation_row_id' in header, header
    code_at = header.index('operation_id')
    id_at = header.index('operation_row_id')
    mine = [r for r in sheet.iter_rows(min_row=2, values_only=True)
            if r[code_at] == one_operation['code']]
    assert mine, 'không thấy Operation trong file xuất'
    assert int(mine[0][id_at]) == one_operation['op_id']


def test_a_legacy_file_without_the_new_column_still_imports(api, db, one_operation):
    """Đây là bài quan trọng nhất của file: đừng phá file khách đang dùng."""
    response = _upload(api, [[one_operation['code'], 'SP', one_operation['po_code'],
                              one_operation['part_name'], 3, 'Tên mới', '', 100, 0, 0,
                              'PLANNED', f"WF|OP|{one_operation['code']}"]], LEGACY_HEADERS)
    assert response.status_code == 200, response.text[:300]
    with db.cursor() as cur:
        cur.execute("SELECT name,sort_order FROM operations WHERE id=%s", (one_operation['op_id'],))
        row = cur.fetchone()
    assert row['name'] == 'Tên mới' and row['sort_order'] == 3


def test_the_row_id_identifies_the_operation_even_when_the_code_column_is_stale(api, db, one_operation):
    """Mã trong file đã cũ đi so với hệ thống -> phải nói ra, không ghi đè bừa."""
    with db.cursor() as cur:
        cur.execute("UPDATE operations SET code=%s WHERE id=%s",
                    (f"{one_operation['code']}-NEW", one_operation['op_id']))
    response = _upload(api, [[one_operation['code'], 'SP', one_operation['po_code'],
                              one_operation['part_name'], 0, 'Tên', '', 100, 0, 0,
                              'PLANNED', f"WF|OP|{one_operation['code']}", one_operation['op_id']]],
                       LEGACY_HEADERS + ['operation_row_id'])
    assert response.status_code >= 400, response.text[:300]
    assert f"{one_operation['code']}-NEW" in response.text


def test_a_row_id_belonging_to_another_po_is_refused(api, db, one_operation):
    sfx = one_operation['suffix']
    with db.cursor() as cur:
        cur.execute("""INSERT INTO production_orders(code,product,planned_quantity,status)
            VALUES(%s,'SP',10,'IN_PROGRESS') RETURNING id""", (f'PO-XL2-{sfx}',))
        other_po = cur.fetchone()['id']
        cur.execute("INSERT INTO parts(production_order_id,code,name,sort_order) VALUES(%s,%s,%s,0)",
                    (other_po, f'PXL2-{sfx}', f'Part2 {sfx}'))
    try:
        response = _upload(api, [[one_operation['code'], 'SP', f'PO-XL2-{sfx}', f'Part2 {sfx}',
                                  0, 'Tên', '', 10, 0, 0, 'PLANNED',
                                  f"WF|OP|{one_operation['code']}", one_operation['op_id']]],
                           LEGACY_HEADERS + ['operation_row_id'])
        assert response.status_code >= 400, response.text[:300]
        assert one_operation['po_code'] in response.text
        with db.cursor() as cur:
            cur.execute("SELECT production_order_id FROM operations WHERE id=%s", (one_operation['op_id'],))
            assert cur.fetchone()['production_order_id'] == one_operation['po_id']
    finally:
        with db.cursor() as cur:
            cur.execute("DELETE FROM parts WHERE production_order_id=%s", (other_po,))
            cur.execute("DELETE FROM production_orders WHERE id=%s", (other_po,))


def test_a_row_id_that_no_longer_exists_is_refused(api, one_operation):
    response = _upload(api, [[one_operation['code'], 'SP', one_operation['po_code'],
                              one_operation['part_name'], 0, 'Tên', '', 100, 0, 0,
                              'PLANNED', f"WF|OP|{one_operation['code']}", 99999999]],
                       LEGACY_HEADERS + ['operation_row_id'])
    assert response.status_code >= 400, response.text[:300]


def test_a_non_numeric_row_id_is_refused_not_silently_ignored(api, one_operation):
    """Ô có nội dung mà không phải id thì KHÔNG được lặng lẽ rơi về tra theo mã."""
    response = _upload(api, [[one_operation['code'], 'SP', one_operation['po_code'],
                              one_operation['part_name'], 0, 'Tên', '', 100, 0, 0,
                              'PLANNED', f"WF|OP|{one_operation['code']}", 'khong-phai-so']],
                       LEGACY_HEADERS + ['operation_row_id'])
    assert response.status_code >= 400, response.text[:300]
    assert 'operation_row_id' in response.text


def test_a_row_id_round_trips_through_export_and_import(api, db, one_operation):
    """Xuất ra rồi nhập lại đúng file đó phải không đổi danh tính của dòng nào."""
    base = os.environ.get('MESFLOW_BASE_URL', BASE_URL).rstrip('/')
    exported = api.get(f'{base}/api/operations/export.xlsx', timeout=60)
    assert exported.status_code == 200
    sheet = load_workbook(io.BytesIO(exported.content), data_only=True)['Operations']
    header = [c.value for c in next(sheet.iter_rows(max_row=1))]
    code_at = header.index('operation_id')
    rows = [list(r) for r in sheet.iter_rows(min_row=2, values_only=True)
            if r[code_at] == one_operation['code']]
    assert rows

    response = _upload(api, rows, header)
    assert response.status_code == 200, response.text[:400]
    with db.cursor() as cur:
        cur.execute("""SELECT id,code,production_order_id,part_id FROM operations WHERE id=%s""",
                    (one_operation['op_id'],))
        row = cur.fetchone()
    assert row['code'] == one_operation['code']
    assert row['production_order_id'] == one_operation['po_id']
    assert row['part_id'] == one_operation['part_id']


def test_excel_cannot_create_an_ambiguous_legacy_label(api, db, one_operation):
    """Cột qr đi thẳng từ file vào bảng -- guard phải chặn ở đây nữa.

    operations_qr_key chỉ cấm hai tem giống hệt nhau. Ca làm tem cũ hoá mơ hồ
    là ca CHÉO CỘT: tem của dòng này trùng với MÃ của dòng khác.
    """
    sfx = one_operation['suffix']
    with db.cursor() as cur:
        cur.execute("""INSERT INTO operations(production_order_id,part_id,code,name,status,qr,sort_order)
            VALUES(%s,%s,%s,'Khác','PLANNED',%s,1) RETURNING id""",
            (one_operation['po_id'], one_operation['part_id'], f'OTHER-{sfx}', f'WF|OPID|Z-{sfx}'))
        other_id = cur.fetchone()['id']

    response = _upload(api, [[f'OTHER-{sfx}', 'SP', one_operation['po_code'],
                              one_operation['part_name'], 1, 'Khác', '', 100, 0, 0,
                              'PLANNED', f"WF|OP|{one_operation['code']}", other_id]],
                       LEGACY_HEADERS + ['operation_row_id'])
    assert response.status_code >= 400, f'Excel đã tạo được tem mơ hồ: {response.text[:300]}'
    with db.cursor() as cur:
        cur.execute("SELECT qr FROM operations WHERE id=%s", (other_id,))
        assert cur.fetchone()['qr'] == f'WF|OPID|Z-{sfx}', 'tem đã bị ghi đè'

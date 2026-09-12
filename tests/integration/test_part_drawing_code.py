"""MÃ BẢN VẼ phải sống được tới màn hình, không dừng ở parser.

Tờ GO ROUTER ghi ``MÃ BẢN VẼ`` ở đầu mỗi sheet, và cả xưởng gọi tên chi tiết
theo mã đó: thợ đọc nó trên bản vẽ giấy, kho dán nó lên thùng, QC ghi nó vào
phiếu. Trước đây importer đọc ô ấy chỉ để chế ra một mã nội bộ rồi thả rơi --
không bảng nào giữ, nên không màn hình nào hiện được, và người dùng không đối
chiếu được Part trên MESFlow với tờ giấy đang cầm.

Bốn điều được khoá ở đây, chạy trên API + PostgreSQL thật:

  a) nhập -> ``template_parts.drawing_code`` giữ NGUYÊN VĂN mã trên tờ giấy
     (khác ``code``, vốn đã lọc ký tự để làm mã nội bộ);
  b) tạo PO -> ``parts.drawing_code`` là BẢN CHỤP, không đọc xuyên sang
     Template;
  c) sửa Template SAU ĐÓ không đụng tới PO đã tạo -- đây là lý do có hai cột;
  d) lưu lại cây Template không được xoá mất mã (replace_tree xoá rồi ghi lại
     toàn bộ cây, nên quên mang theo một cột là mất sạch).

Tờ mức quy trình không khai mã bản vẽ thì giữ NULL, không bịa.
"""
import io
import os
import uuid

import pytest
from openpyxl import Workbook

BASE_URL = os.environ.get('MESFLOW_BASE_URL', 'http://mesflow-test-api:8080').rstrip('/')
pytestmark = pytest.mark.integration

BLOCK_HEIGHT = 12
#: Cố ý có khoảng trắng và dấu chấm như mã thật trên tờ NEWARK
#: ('KM- 967.232006L- 01'): mã nội bộ sẽ bị lọc, mã bản vẽ thì không.
DRAWING_CODE = 'KM- 967.232006L- 01'


def _router_bytes(po_number, *, drawing_code=DRAWING_CODE, process_sheet=False):
    """Một tờ Part có mã bản vẽ, tuỳ chọn thêm một tờ mức quy trình."""
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet('Chân ghế A')
    ws['A2'] = 'PO NUMBER:'; ws['C2'] = po_number
    ws['G2'] = 'LOẠI ĐƠN HÀNG'; ws['I2'] = 'SỐ LƯỢNG'
    ws['A3'] = 'QTY:'; ws['C3'] = 110
    ws['A4'] = 'TÊN BẢN VẼ:'; ws['C4'] = 'Chân ghế A'
    ws['G4'] = 'SẢN XUẤT HÀNG LOẠT'; ws['I4'] = 110
    ws['A5'] = 'MÃ BẢN VẼ:'; ws['C5'] = drawing_code
    row = 8
    ws.cell(row=row, column=1, value='OPERATION # 01- CẮT LASER')
    ws.cell(row=row + 1, column=1, value='Part Number ( Mã bản vẽ )')
    ws.cell(row=row + 2, column=12, value='Thời gian Setup ( phút )')
    ws.cell(row=row + 3, column=1, value='SETUP')
    ws.cell(row=row + 3, column=12, value=20)
    ws.cell(row=row + 5, column=12, value='Thời gian gia công / sản phẩm (s)')
    ws.cell(row=row + 6, column=1, value='Thời gian gia công')
    ws.cell(row=row + 6, column=12, value=100)
    if process_sheet:
        # Tờ mức quy trình: không TÊN, không MÃ BẢN VẼ, không block OPERATION.
        other = wb.create_sheet('SƠN TĨNH ĐIỆN')
        other['A2'] = 'PO NUMBER:'; other['C2'] = po_number
    buffer = io.BytesIO(); wb.save(buffer)
    return buffer.getvalue()


def _upload(api, data, path='/api/templates/import-workbook'):
    files = {'file': ('router.xlsx', data,
                      'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}
    return api.post(f'{BASE_URL}{path}', files=files, timeout=300)


def _po_number():
    return f'PO{uuid.uuid4().hex[:8].upper()}'


@pytest.fixture
def imported(api):
    po_code = _po_number()
    response = _upload(api, _router_bytes(po_code, process_sheet=True))
    assert response.status_code == 200, response.text[:400]
    body = response.json()
    return body['template_id'], body['production_order']['id'], po_code


# --- a) nhập -> Template giữ nguyên văn mã bản vẽ ------------------------

def test_nhap_giu_nguyen_van_ma_ban_ve_tren_template(imported, db):
    template_id, _po_id, _po_code = imported
    row = db.execute("""SELECT code,drawing_code FROM template_parts
        WHERE template_id=%s AND name='Chân ghế A'""", (template_id,)).fetchone()
    assert row['drawing_code'] == DRAWING_CODE, 'phải giữ nguyên văn, không lọc ký tự'
    assert row['code'] != DRAWING_CODE, 'mã nội bộ là thứ khác, đã lọc ký tự'


def test_to_muc_quy_trinh_khong_bi_bia_ma_ban_ve(imported, db):
    template_id, _po_id, _po_code = imported
    row = db.execute("""SELECT drawing_code FROM template_parts
        WHERE template_id=%s AND name='SƠN TĨNH ĐIỆN'""", (template_id,)).fetchone()
    assert row is not None, 'tờ quy trình vẫn phải thành một Part'
    assert row['drawing_code'] is None, 'không có bản vẽ thì để NULL, không bịa'


# --- b) tạo PO -> bản chụp ----------------------------------------------

def test_po_chup_lai_ma_ban_ve_luc_tao(imported, db):
    _template_id, po_id, _po_code = imported
    row = db.execute("""SELECT drawing_code FROM parts
        WHERE production_order_id=%s AND name='Chân ghế A'""", (int(po_id),)).fetchone()
    assert row['drawing_code'] == DRAWING_CODE


def test_api_parts_tra_ve_ma_ban_ve_cho_man_po(api, imported):
    """Màn PO dựng thẻ Part từ /api/parts -- API không trả thì giao diện không hiện."""
    _template_id, po_id, _po_code = imported
    listed = api.get(f'{BASE_URL}/api/parts?production_order_id={int(po_id)}&limit=1000',
                     timeout=60)
    assert listed.status_code == 200, listed.text[:400]
    part = next(p for p in listed.json()['items'] if p['name'] == 'Chân ghế A')
    assert part['drawing_code'] == DRAWING_CODE


# --- c) bản chụp KHÔNG hồi tố -------------------------------------------

def test_sua_template_sau_do_khong_viet_lai_tai_lieu_cua_po_dang_chay(api, imported, db):
    """Đây là lý do có HAI cột chứ không phải một.

    Template còn sửa được sau khi PO đã chạy. Nếu PO đọc xuyên sang Template
    thì một lần sửa Template sẽ lặng lẽ đổi mã bản vẽ trên tờ giấy của một PO
    đang sản xuất -- thợ cầm tờ in cũ, hệ thống nói mã khác.
    """
    template_id, po_id, _po_code = imported
    tree = api.get(f'{BASE_URL}/api/templates/{template_id}/tree', timeout=60).json()
    parts = [{**p, 'key': p.get('code') or p['id'],
              'drawing_code': ('KM-DOI-MOI' if p['name'] == 'Chân ghế A'
                               else p.get('drawing_code'))}
             for p in tree['parts']]
    key_of_part = {p['id']: p['key'] for p in parts}
    operations = [{**o, 'part_key': key_of_part[o['part_id']]} for o in tree['operations']]
    saved = api.put(f'{BASE_URL}/api/templates/{template_id}/tree',
                    json={'parts': parts, 'operations': operations,
                          'equipment': tree.get('equipment') or []}, timeout=60)
    assert saved.status_code == 200, saved.text[:400]

    assert db.execute("""SELECT drawing_code FROM template_parts
        WHERE template_id=%s AND name='Chân ghế A'""",
        (template_id,)).fetchone()['drawing_code'] == 'KM-DOI-MOI'
    assert db.execute("""SELECT drawing_code FROM parts
        WHERE production_order_id=%s AND name='Chân ghế A'""",
        (int(po_id),)).fetchone()['drawing_code'] == DRAWING_CODE, 'PO không được đổi theo'


# --- d) lưu lại cây Template không được đánh rơi mã ----------------------

def test_luu_lai_template_khong_xoa_mat_ma_ban_ve(api, imported, db):
    """replace_tree xoá rồi ghi lại TOÀN BỘ cây: quên một cột là mất sạch."""
    template_id, _po_id, _po_code = imported
    tree = api.get(f'{BASE_URL}/api/templates/{template_id}/tree', timeout=60).json()
    parts = [{**p, 'key': p.get('code') or p['id']} for p in tree['parts']]
    key_of_part = {p['id']: p['key'] for p in parts}
    operations = [{**o, 'part_key': key_of_part[o['part_id']]} for o in tree['operations']]
    # Sửa một thứ KHÁC hẳn, không đụng tới mã bản vẽ.
    for part in parts:
        if part['name'] == 'Chân ghế A':
            part['name'] = 'Chân ghế A (đã sửa tên)'
    saved = api.put(f'{BASE_URL}/api/templates/{template_id}/tree',
                    json={'parts': parts, 'operations': operations,
                          'equipment': tree.get('equipment') or []}, timeout=60)
    assert saved.status_code == 200, saved.text[:400]
    row = db.execute("""SELECT drawing_code FROM template_parts
        WHERE template_id=%s AND name='Chân ghế A (đã sửa tên)'""", (template_id,)).fetchone()
    assert row['drawing_code'] == DRAWING_CODE, 'sửa tên Part không được xoá mã bản vẽ'

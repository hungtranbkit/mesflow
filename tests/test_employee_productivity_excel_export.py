from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from mesflow.web.productivity_excel import (
    EXPORT_HEADERS,
    build_employee_productivity_xlsx,
    filter_and_sort_employee_rows,
)


def _report():
    return {
        "summary": {"from": "2026-09-01", "to": "2026-09-28"},
        "employees": [
            {"employee_id": 2, "employee_code": "NV02", "employee_name": "Bình", "department": "May", "team": "A",
             "completed_sessions": 4, "completed_valid_sessions": 3, "productivity_percent": 82.5,
             "good_qty": 120, "defect_qty": 2, "worked_seconds": 14400},
            {"employee_id": 1, "employee_code": "NV01", "employee_name": "An", "department": "May", "team": "B",
             "completed_sessions": 3, "completed_valid_sessions": 3, "productivity_percent": 75.0,
             "good_qty": 90, "defect_qty": 1, "worked_seconds": 10800},
            {"employee_id": 3, "employee_code": "NV03", "employee_name": "Cường", "department": "Kho", "team": "",
             "completed_sessions": 1, "completed_valid_sessions": 0, "productivity_percent": None,
             "good_qty": 10, "defect_qty": 0, "worked_seconds": 3600},
        ],
    }


def test_filter_and_sort_matches_visible_employee_list():
    rows = filter_and_sort_employee_rows(
        _report()["employees"], search="n", department="May", sort_key="employee_name", sort_dir="asc"
    )
    assert [row["employee_code"] for row in rows] == ["NV01", "NV02"]


def test_workbook_is_print_ready_and_has_signature_column():
    stream = build_employee_productivity_xlsx(
        _report(), search="an", department="May", sort_key="good_qty", sort_dir="desc"
    )
    wb = load_workbook(BytesIO(stream.getvalue()))
    ws = wb["Năng suất nhân viên"]
    assert [ws.cell(4, col).value for col in range(1, len(EXPORT_HEADERS) + 1)] == EXPORT_HEADERS
    assert "Bộ phận: May" in ws["A2"].value
    assert "Tìm nhân viên: an" in ws["A2"].value
    assert "Năng suất TB theo bộ lọc: 75.0%" in ws["A2"].value
    assert ws["B5"].value == "NV01"
    assert ws["H5"].value == 0.75
    assert ws["H5"].number_format == "0.0%"
    assert ws["K5"].number_format == "[h]:mm"
    assert ws["L5"].value is None
    assert ws.column_dimensions["L"].width >= 20
    assert ws.row_dimensions[5].height >= 30
    assert ws.freeze_panes == "A5"
    assert ws.page_setup.orientation == "landscape"
    assert ws.page_setup.fitToWidth == 1


def test_frontend_export_uses_current_filters_and_sort():
    js = Path("app/mesflow/web/static/pages/employee-productivity.js").read_text()
    assert 'id="epExport"' in js
    assert "/api/reports/employee-productivity/export.xlsx" in js
    for token in ["q.set('from'", "q.set('to'", "q.set('search'", "q.set('department'", "q.set('sort'", "q.set('dir'"]:
        assert token in js


def test_export_route_contract():
    source = Path("app/mesflow/web/analytics.py").read_text()
    assert "@bp.get('/reports/employee-productivity/export.xlsx')" in source
    assert "build_employee_productivity_xlsx" in source
    assert "int(employee) if employee else None" in source
    assert "if employe else None" not in source
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in source

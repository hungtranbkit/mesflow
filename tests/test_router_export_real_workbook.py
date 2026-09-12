"""Xuất tờ router phải là chính file gốc đã nhập, không phải file dựng lại.

Bài này chạy trên **file thật** (`tests/fixtures/router-newark-arm-chair.xlsx`,
44 sheet của xưởng) chứ không phải workbook dựng bằng openpyxl, vì thứ cần
chứng minh là *giữ nguyên biểu mẫu của khách*: merge, độ rộng cột, chiều cao
dòng, ảnh/logo, khung in, công thức. Một workbook dựng tay không có thứ nào
trong số đó nên không chứng minh được gì.

Ba lỗi thật đã bị bắt đúng bằng file này và không bị workbook dựng tay bắt:
mọi tem rơi xuống sheet phụ (mã OP thật mang tiền tố PO), dấu '-' trong ô thời
gian setup làm cả file bị từ chối, và số Operation đánh trùng trong cùng Part.
"""
import re
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from mesflow.web.excel_io import _parse_go_router_template
from mesflow.web.router_export import (
    EXTRA_SHEET_TITLE, QR_OP_LABEL, QR_SETUP_LABEL, _stamp_source_workbook)
from mesflow.db.repositories.master_data import operation_code_suffix

pytest.importorskip('PIL', reason='Pillow là dependency của tính năng xuất QR')

FIXTURE = Path(__file__).resolve().parent / 'fixtures/router-newark-arm-chair.xlsx'
PO = {'id': 1, 'code': 'PO-NEWARK', 'product': 'NEWARK ARM CHAIR',
      'planned_quantity': 110, 'source_template_id': 7}


@pytest.fixture(scope='module')
def source_bytes():
    assert FIXTURE.is_file(), f'thiếu fixture {FIXTURE}'
    return FIXTURE.read_bytes()


@pytest.fixture(scope='module')
def parsed(source_bytes):
    return _parse_go_router_template(load_workbook(BytesIO(source_bytes), data_only=True),
                                     FIXTURE.name)


@pytest.fixture(scope='module')
def rows(parsed):
    """Operation "thật" như CSDL sẽ có: mã mang tiền tố PO, SETUP khi phút > 0."""
    part_code_by_key = {p['key']: p['code'] for p in parsed['parts']}
    part_index = {p['key']: i for i, p in enumerate(parsed['parts'])}
    built = []
    next_id = 1000
    for op in parsed['operations']:
        part_code = part_code_by_key[op['part_key']]
        suffix = operation_code_suffix(part_code, op['code'])
        next_id += 1
        op_id = next_id
        setup_id = None
        if op['requires_setup']:
            next_id += 1
            setup_id = next_id
        built.append({
            'id': op_id, 'code': f"{PO['code']}-{suffix}", 'name': op['name'],
            'sort_order': op['sort_order'],
            'standard_seconds_per_unit': op['standard_seconds_per_unit'],
            'requires_setup': op['requires_setup'],
            'expected_setup_minutes': op['expected_setup_minutes'],
            'part_id': part_index[op['part_key']], 'part_code': part_code,
            'part_name': op['_excel_sheet'], 'part_sort': part_index[op['part_key']],
            'op_qr': f'WF|OPID|{op_id}',
            'setup_id': setup_id,
            'setup_code': f"{PO['code']}-{suffix}-SU" if setup_id else None,
            'setup_name': f"Setup {op['name']}" if setup_id else None,
            'setup_minutes': op['expected_setup_minutes'],
            'setup_qr': f'WF|OPID|{setup_id}' if setup_id else None,
        })
    return built


@pytest.fixture(scope='module')
def exported(source_bytes, rows):
    wb, placed, matched, unmatched = _stamp_source_workbook(source_bytes, PO, rows)
    buffer = BytesIO(); wb.save(buffer); buffer.seek(0)
    return load_workbook(buffer), placed, matched, unmatched


# --- file thật vào đúng như mong đợi --------------------------------------

def test_file_that_shop_uses_parses_completely(parsed):
    assert len(parsed['parts']) == 44
    assert len(parsed['operations']) == 112
    assert sum(1 for op in parsed['operations'] if op['requires_setup']) == 47
    # 10 chỗ đánh trùng số Operation trong cùng một Part -> tách được, có cảnh báo.
    assert len(parsed['warnings']) == 10
    codes = [op['code'] for op in parsed['operations']]
    assert len(codes) == len(set(codes)), 'không được còn mã trùng sau khi tách'


def test_setup_values_from_the_real_sheets(parsed):
    by_code = {op['code']: op for op in parsed['operations']}
    # Đúng những con số user mô tả trên file của họ.
    assert by_code['KM-967-232006L-01-OP01']['expected_setup_minutes'] == 20
    assert by_code['KM-967-232006L-01-OP02']['requires_setup'] is False
    assert by_code['KM-504539877-B4-2-OP01']['expected_setup_minutes'] == 120


# --- xuất = file gốc + tem, không phải file dựng lại ----------------------

def test_every_source_sheet_survives_the_export(source_bytes, exported):
    """Không mất sheet nào, không đổi tên sheet nào."""
    original = load_workbook(BytesIO(source_bytes))
    workbook, _, _, unmatched = exported
    assert unmatched == [], 'file gốc phải khớp được MỌI Operation'
    assert EXTRA_SHEET_TITLE not in workbook.sheetnames
    assert workbook.sheetnames == original.sheetnames


def test_layout_of_every_sheet_is_preserved(source_bytes, exported):
    """Merge, độ rộng cột, chiều cao dòng của vùng dữ liệu gốc không đổi."""
    original = load_workbook(BytesIO(source_bytes))
    workbook, _, _, _ = exported
    for name in original.sheetnames:
        before, after = original[name], workbook[name]
        assert {str(r) for r in before.merged_cells.ranges} == {
            str(r) for r in after.merged_cells.ranges}, f'{name}: merge bị đổi'
        # Cột của vùng dữ liệu (A..M) phải giữ nguyên bề rộng; làn QR nằm sau đó.
        for letter in 'ABCDEFGHIJKLM':
            assert (before.column_dimensions[letter].width
                    == after.column_dimensions[letter].width), f'{name}: cột {letter}'
        for index, dimension in before.row_dimensions.items():
            assert after.row_dimensions[index].height == dimension.height, (
                f'{name}: chiều cao dòng {index}')


def test_original_images_and_text_survive(source_bytes, exported):
    """Logo/ảnh của khách còn nguyên, và mọi ô có chữ giữ đúng giá trị cũ."""
    original = load_workbook(BytesIO(source_bytes))
    workbook, placed, _, _ = exported
    labels = {QR_OP_LABEL, QR_SETUP_LABEL}
    for name in original.sheetnames:
        before, after = original[name], workbook[name]
        assert len(after._images) >= len(before._images), f'{name}: mất ảnh gốc'
        for row in before.iter_rows():
            for cell in row:
                if cell.value in (None, ''):
                    continue
                assert after[cell.coordinate].value == cell.value, (
                    f'{name}!{cell.coordinate} bị ghi đè')
        # Chữ duy nhất được thêm vào là hai nhãn tem.
        added = {c.value for r in after.iter_rows() for c in r if c.value not in (None, '')}
        original_values = {c.value for r in before.iter_rows() for c in r
                           if c.value not in (None, '')}
        assert (added - original_values) <= labels


def test_qr_lane_never_covers_data_on_any_sheet(exported, parsed):
    """Trên MỌI sheet có tem, làn tem phải nằm sau ô có chữ xa nhất."""
    workbook, _, _, _ = exported
    labels = {QR_OP_LABEL, QR_SETUP_LABEL}
    sheets_with_operations = {op['_excel_sheet'] for op in parsed['operations']}
    assert len(sheets_with_operations) == 43
    for name in sheets_with_operations:
        ws = workbook[name]
        rightmost_data = max(
            (c.column for row in ws.iter_rows() for c in row
             if c.value not in (None, '') and c.value not in labels), default=0)
        label_columns = [c.column for row in ws.iter_rows() for c in row
                         if c.value in labels]
        assert label_columns, f'{name}: không có tem nào'
        assert min(label_columns) > rightmost_data, (
            f'{name}: tem nằm chồng lên vùng dữ liệu')


def test_sheet_without_any_operation_block_is_left_untouched(source_bytes, exported, parsed):
    """Sheet không có block nào ('SƠN TĨNH ĐIỆN') phải được giữ y nguyên.

    Nó vẫn là một tờ trong biểu mẫu của khách, chỉ là không có công đoạn nào để
    gắn tem. Sửa nó -- kể cả chỉ đặt lại khổ in -- là đụng vào thứ không cần.
    """
    original = load_workbook(BytesIO(source_bytes))
    workbook, placed, _, _ = exported
    untouched = set(original.sheetnames) - {op['_excel_sheet'] for op in parsed['operations']}
    assert untouched == {'SƠN TĨNH ĐIỆN'}
    for name in untouched:
        assert not any(item['sheet'] == name for item in placed)
        assert len(workbook[name]._images) == len(original[name]._images)
        assert (workbook[name].page_setup.fitToWidth
                == original[name].page_setup.fitToWidth)


def test_labels_land_on_the_block_of_their_own_operation(parsed, rows, exported):
    """Tem của một Operation phải nằm trên ĐÚNG sheet chứa block của nó.

    Đếm đủ số tem không chứng minh được gì: lỗi thật là cả 159 tem rơi xuống
    một sheet phụ mà tổng số vẫn đúng.
    """
    workbook, placed, matched, _ = exported
    assert matched == len(rows)
    sheet_of_operation = {}
    part_code_by_key = {p['key']: p['code'] for p in parsed['parts']}
    for op in parsed['operations']:
        suffix = operation_code_suffix(part_code_by_key[op['part_key']], op['code'])
        sheet_of_operation[f"{PO['code']}-{suffix}"] = op['_excel_sheet']
    by_id = {row['id']: row for row in rows}
    setup_parent = {row['setup_id']: row for row in rows if row['setup_id']}
    for item in placed:
        row = by_id.get(item['operation_id']) or setup_parent[item['operation_id']]
        assert item['sheet'] == sheet_of_operation[row['code']], (
            f"tem của {row['code']} nằm nhầm sheet {item['sheet']}")


def test_setup_label_only_where_a_setup_operation_exists(rows, exported):
    _, placed, _, _ = exported
    setup_ids = {row['setup_id'] for row in rows if row['setup_id']}
    assert len(setup_ids) == 47
    assert {i['operation_id'] for i in placed if i['kind'] == 'SETUP'} == setup_ids
    assert {i['operation_id'] for i in placed if i['kind'] == 'OP'} == {
        row['id'] for row in rows}
    assert len(placed) == 112 + 47


def test_every_payload_is_the_canonical_operation_id(rows, exported):
    _, placed, _, _ = exported
    payloads = [item['payload'] for item in placed]
    assert len(payloads) == len(set(payloads)), 'hai tem cùng payload'
    for item in placed:
        assert re.fullmatch(r'WF\|OPID\|\d+', item['payload']), item['payload']
        assert item['payload'] == f"WF|OPID|{item['operation_id']}"


def test_two_sheets_sharing_an_operation_name_get_different_labels(parsed, rows, exported):
    """'CẮT LASER' xuất hiện ở nhiều Part -- tuyệt đối không dùng chung id."""
    _, placed, _, _ = exported
    by_name = {}
    for row in rows:
        by_name.setdefault(row['name'], []).append(row['id'])
    repeated = {name: ids for name, ids in by_name.items() if len(ids) > 1}
    assert repeated, 'file thật phải có công đoạn trùng tên giữa các Part'
    payload_of = {i['operation_id']: i['payload'] for i in placed}
    for name, ids in repeated.items():
        payloads = {payload_of[i] for i in ids}
        assert len(payloads) == len(ids), f'{name}: tem bị dùng chung'


def test_print_setup_fits_the_qr_lane_onto_the_page(exported, parsed):
    """Sheet nào có thêm tem thì khổ in phải ôm được làn tem đó."""
    workbook, _, _, _ = exported
    for name in {op['_excel_sheet'] for op in parsed['operations']}:
        ws = workbook[name]
        assert ws.sheet_properties.pageSetUpPr.fitToPage is True, name
        assert ws.page_setup.fitToWidth == 1, name

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from zoneinfo import ZoneInfo

EXPORT_HEADERS = [
    "STT", "Mã NV", "Nhân viên", "Bộ phận", "Nhóm",
    "Phiên hoàn tất", "Phiên hợp lệ", "Năng suất TB (%)",
    "Sản lượng đạt", "Lỗi", "Tổng thời gian", "NV xác nhận",
]

SORT_FIELDS = {
    "employee_name": "employee_name",
    "employee": "employee_name",
    "session": "completed_sessions",
    "completed_sessions": "completed_sessions",
    "productivity": "productivity_percent",
    "productivity_percent": "productivity_percent",
    "good_qty": "good_qty",
    "output": "good_qty",
    "worked_seconds": "worked_seconds",
    "time": "worked_seconds",
}


def filter_and_sort_employee_rows(
    employees: Iterable[dict[str, Any]],
    *,
    search: str = "",
    department: str = "",
    sort_key: str = "productivity_percent",
    sort_dir: str = "desc",
) -> list[dict[str, Any]]:
    query = (search or "").strip().casefold()
    department = (department or "").strip()
    rows: list[dict[str, Any]] = []
    for raw in employees:
        row = dict(raw)
        if department and str(row.get("department") or "") != department:
            continue
        if query:
            haystack = f"{row.get('employee_name') or ''} {row.get('employee_code') or ''}".casefold()
            if query not in haystack:
                continue
        rows.append(row)

    field = SORT_FIELDS.get(sort_key, "productivity_percent")
    reverse = str(sort_dir).lower() != "asc"
    if field == "employee_name":
        rows.sort(
            key=lambda row: (
                str(row.get("employee_name") or "").casefold(),
                str(row.get("employee_code") or "").casefold(),
            ),
            reverse=reverse,
        )
        return rows

    present = [row for row in rows if row.get(field) is not None]
    missing = [row for row in rows if row.get(field) is None]
    present.sort(
        key=lambda row: (
            float(row.get(field) or 0),
            str(row.get("employee_code") or "").casefold(),
        ),
        reverse=reverse,
    )
    return present + missing


def _duration_days(seconds: Any) -> float:
    try:
        return max(float(seconds or 0), 0.0) / 86400.0
    except (TypeError, ValueError):
        return 0.0


def build_employee_productivity_xlsx(
    report: dict[str, Any],
    *,
    search: str = "",
    department: str = "",
    sort_key: str = "productivity_percent",
    sort_dir: str = "desc",
) -> BytesIO:
    rows = filter_and_sort_employee_rows(
        report.get("employees") or [],
        search=search,
        department=department,
        sort_key=sort_key,
        sort_dir=sort_dir,
    )
    wb = Workbook()
    ws = wb.active
    ws.title = "Năng suất nhân viên"
    _write_summary_sheet(ws, report, rows, search=search, department=department)
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream


def _write_summary_sheet(ws, report: dict[str, Any], rows: list[dict[str, Any]], *, search: str, department: str) -> None:
    """The "Tổng hợp" sheet. Shared verbatim by the summary export and the
    detail workbook's "Tong hop NV" sheet, so both show the same numbers."""
    scores = [float(row["productivity_percent"]) for row in rows if row.get("productivity_percent") is not None]
    filtered_avg = (sum(scores) / len(scores)) if scores else None
    summary = report.get("summary") or {}
    ws.sheet_view.showGridLines = False
    last_col = len(EXPORT_HEADERS)
    last_letter = get_column_letter(last_col)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws["A1"] = "BÁO CÁO NĂNG SUẤT NHÂN VIÊN"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    filters = [
        f"Từ ngày: {summary.get('from') or '—'}",
        f"Đến ngày: {summary.get('to') or '—'}",
        f"Bộ phận: {department or 'Tất cả'}",
        f"Tìm nhân viên: {search.strip() or 'Tất cả'}",
        f"Năng suất TB theo bộ lọc: {filtered_avg:.1f}%" if filtered_avg is not None else "Năng suất TB theo bộ lọc: —",
    ]
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    ws["A2"] = " | ".join(filters)
    ws["A2"].font = Font(size=10)
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center")

    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=last_col)
    ws["A3"] = 'In bảng này và để nhân viên ký/xác nhận tại cột "NV xác nhận".'
    ws["A3"].font = Font(italic=True, size=10)
    ws["A3"].alignment = Alignment(horizontal="left", vertical="center")

    header_row = 4
    thin = Side(style="thin", color="B8C2CC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for col, header in enumerate(EXPORT_HEADERS, 1):
        cell = ws.cell(header_row, col, header)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[header_row].height = 28

    for index, item in enumerate(rows, 1):
        row_idx = header_row + index
        pct = item.get("productivity_percent")
        values = [
            index,
            item.get("employee_code") or "",
            item.get("employee_name") or "",
            item.get("department") or "",
            item.get("team") or "",
            int(item.get("completed_sessions") or 0),
            int(item.get("completed_valid_sessions") or 0),
            (float(pct) / 100.0) if pct is not None else None,
            int(item.get("good_qty") or 0),
            int(item.get("defect_qty") or 0),
            _duration_days(item.get("worked_seconds")),
            "",
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row_idx, col, value)
            cell.border = border
            cell.alignment = Alignment(
                horizontal="left" if col in (2, 3, 4, 5, 12) else "center",
                vertical="center",
                wrap_text=col in (3, 4, 5, 12),
            )
        ws.cell(row_idx, 8).number_format = "0.0%"
        ws.cell(row_idx, 11).number_format = "[h]:mm"
        ws.row_dimensions[row_idx].height = 30

    last_row = max(header_row, header_row + len(rows))
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A{header_row}:{last_letter}{last_row}"
    ws.print_area = f"A1:{last_letter}{last_row}"
    ws.print_title_rows = f"{header_row}:{header_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.45
    ws.page_margins.bottom = 0.45
    ws.oddFooter.center.text = "Trang &[Page]/&[Pages]"

    widths = [6, 13, 24, 18, 16, 13, 13, 15, 15, 11, 16, 22]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width


# --- "Chi tiết từng nhân viên" ------------------------------------------------
# One row = one session behind the summary numbers
# (ReportRepository.employee_productivity_sessions, same scope and the same
# per-session score). "% năng suất" is that session's completion_percent =
# thời gian định mức / thời gian thực tế; the employee's summary % is the
# average of the non-empty values in their rows -- no new formula.
DETAIL_HEADERS = [
    "STT", "Mã NV", "Nhân viên", "Bộ phận", "Ngày", "PO", "Part", "Mã OP", "Operation",
    "Bắt đầu", "Kết thúc", "Thời gian thực tế", "Sản lượng đạt", "Lỗi",
    "Định mức (giây/SP)", "Thời gian định mức", "% năng suất", "Trạng thái / ghi chú",
    "NV xác nhận", "Mã phiên",
]
DETAIL_WIDTHS = [6, 12, 24, 16, 11, 14, 16, 16, 24, 16, 16, 13, 11, 8, 12, 13, 12, 34, 18, 10]
DETAIL_SHEET = "Chi tiet"
DETAIL_SUMMARY_SHEET = "Tong hop NV"


def _local_naive(value: Any, tz: ZoneInfo) -> Any:
    """Excel has no time zones: show the factory's local wall-clock time."""
    if isinstance(value, datetime):
        if value.tzinfo is not None:
            value = value.astimezone(tz)
        return value.replace(tzinfo=None)
    return value


def _part_label(session: dict[str, Any]) -> str:
    code = str(session.get("part_code") or "").strip()
    name = str(session.get("part_name") or "").strip()
    return f"{code} · {name}" if code and name and name != code else (code or name)


def session_status_note(session: dict[str, Any]) -> str:
    """Why a session does / does not carry a %, plus the review flags the
    session already has. Only existing data, no new rule."""
    parts: list[str] = []
    if session.get("completion_percent") is not None:
        parts.append("Tính năng suất")
    elif session.get("is_repair"):
        parts.append("Ca sửa hàng · không có định mức")
    elif float(session.get("standard_seconds_per_unit") or 0) <= 0:
        parts.append("Thiếu định mức · không tính năng suất")
    elif int(session.get("good_qty") or 0) + int(session.get("defect_qty") or 0) <= 0:
        # Standard exists but nothing was reported -> expected time is 0,
        # so the existing score is empty (same rule, clearer reason).
        parts.append("Không có sản lượng · không tính năng suất")
    else:
        parts.append("Không đủ dữ liệu thời gian · không tính năng suất")
    if session.get("closed_by_system"):
        parts.append("Hệ thống tự đóng")
    if session.get("quantity_confirmed") is False:
        parts.append("Chưa xác nhận sản lượng")
    for key in ("close_reason", "note"):
        text = str(session.get(key) or "").strip()
        if text:
            parts.append(text)
    return " · ".join(parts)


def detail_rows_for_employees(
    employees: list[dict[str, Any]], sessions: Iterable[dict[str, Any]]
) -> list[dict[str, Any]]:
    """Sessions of the (already filtered + sorted) employees, grouped in the
    same employee order as the summary, each employee's sessions by start."""
    order = {row.get("employee_id"): index for index, row in enumerate(employees)}
    picked = [dict(s) for s in sessions if s.get("employee_id") in order]
    picked.sort(key=lambda s: (order[s.get("employee_id")], str(s.get("started_at") or ""), s.get("session_id") or 0))
    return picked


def build_employee_productivity_detail_xlsx(
    report: dict[str, Any],
    detail: dict[str, Any],
    *,
    search: str = "",
    department: str = "",
    sort_key: str = "productivity_percent",
    sort_dir: str = "desc",
    timezone_name: str = "Asia/Ho_Chi_Minh",
) -> BytesIO:
    employees = filter_and_sort_employee_rows(
        report.get("employees") or [],
        search=search,
        department=department,
        sort_key=sort_key,
        sort_dir=sort_dir,
    )
    rows = detail_rows_for_employees(employees, detail.get("sessions") or [])
    tz = ZoneInfo(timezone_name)
    summary = report.get("summary") or {}

    wb = Workbook()
    ws = wb.active
    ws.title = DETAIL_SHEET
    ws.sheet_view.showGridLines = False
    last_col = len(DETAIL_HEADERS)
    last_letter = get_column_letter(last_col)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws["A1"] = "BÁO CÁO NĂNG SUẤT CHI TIẾT TỪNG NHÂN VIÊN"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 26
    filters = [
        f"Từ ngày: {summary.get('from') or '—'}",
        f"Đến ngày: {summary.get('to') or '—'}",
        f"Bộ phận: {department or 'Tất cả'}",
        f"Tìm nhân viên: {search.strip() or 'Tất cả'}",
        f"{len(employees)} nhân viên · {len(rows)} phiên làm việc đã kết thúc",
    ]
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    ws["A2"] = " | ".join(filters)
    ws["A2"].font = Font(size=10)
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=last_col)
    note = ('Mỗi dòng là một phiên làm việc đã kết thúc. % năng suất = thời gian định mức / thời gian thực tế; '
            'năng suất của nhân viên (sheet "Tong hop NV") là trung bình các % có giá trị.')
    if detail.get("truncated"):
        note += " CẢNH BÁO: vượt giới hạn số dòng, danh sách bị cắt -- thu hẹp khoảng ngày."
    ws["A3"] = note
    ws["A3"].font = Font(italic=True, size=10, color="C43232" if detail.get("truncated") else None)

    header_row = 4
    thin = Side(style="thin", color="B8C2CC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for col, header in enumerate(DETAIL_HEADERS, 1):
        cell = ws.cell(header_row, col, header)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[header_row].height = 32

    text_cols = {2, 3, 4, 6, 7, 8, 9, 18, 19}
    for index, item in enumerate(rows, 1):
        row_idx = header_row + index
        pct = item.get("completion_percent")
        values = [
            index,
            item.get("employee_code") or "",
            item.get("employee_name") or "",
            item.get("department") or "",
            item.get("work_date"),
            item.get("po_code") or "",
            _part_label(item),
            item.get("operation_code") or "",
            item.get("operation_name") or "",
            _local_naive(item.get("started_at"), tz),
            _local_naive(item.get("ended_at"), tz),
            _duration_days(item.get("actual_seconds")),
            int(item.get("good_qty") or 0),
            int(item.get("defect_qty") or 0),
            float(item.get("standard_seconds_per_unit") or 0) or None,
            _duration_days(item.get("expected_seconds")) if item.get("expected_seconds") else None,
            (float(pct) / 100.0) if pct is not None else None,
            session_status_note(item),
            "",
            item.get("session_id"),
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row_idx, col, value)
            cell.border = border
            cell.alignment = Alignment(
                horizontal="left" if col in text_cols else ("right" if col >= 12 else "center"),
                vertical="center",
                wrap_text=col in (3, 9, 18),
            )
        ws.cell(row_idx, 5).number_format = "dd/mm/yyyy"
        ws.cell(row_idx, 10).number_format = "dd/mm/yyyy hh:mm"
        ws.cell(row_idx, 11).number_format = "dd/mm/yyyy hh:mm"
        ws.cell(row_idx, 12).number_format = "[h]:mm:ss"
        ws.cell(row_idx, 15).number_format = "0.##"
        ws.cell(row_idx, 16).number_format = "[h]:mm:ss"
        ws.cell(row_idx, 17).number_format = "0.0%"

    last_row = max(header_row, header_row + len(rows))
    ws.freeze_panes = "D5"  # header rows + Mã NV / Nhân viên stay visible
    ws.auto_filter.ref = f"A{header_row}:{last_letter}{last_row}"
    ws.print_title_rows = f"{header_row}:{header_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.center.text = "Trang &[Page]/&[Pages]"
    for col, width in enumerate(DETAIL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col)].width = width

    _write_summary_sheet(wb.create_sheet(DETAIL_SUMMARY_SHEET), report, employees, search=search, department=department)

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream

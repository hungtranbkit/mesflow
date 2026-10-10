"""XLSX renderer: XlsxWriter (preferred), openpyxl fallback.

Safety: every text cell is written with write_string(), and the workbook is
opened with strings_to_formulas/strings_to_urls OFF, so a value such as
"=HYPERLINK(...)" typed into an employee name or note is stored as plain
text and never evaluated by Excel (formula injection). The openpyxl fallback
gets the same guarantee by forcing data_type 's'.

Layout per section/sheet: title, report metadata (range, timezone, export
time/user, filters), header row (frozen, auto-filter, repeated on print),
data rows, bold totals row. A4 landscape, fit to width, page footer.
"""
from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from typing import Any
from zoneinfo import ZoneInfo

from mesflow.reporting.datasets import DATE, DATETIME, DURATION, INT, NUMBER, PERCENT, Column, ReportDocument

SHEET_NAME_MAX = 31
_SHEET_FORBIDDEN = set('[]:*?/\\')
_RESERVED = {'history'}
HEADER_ROW = 4  # 0-based row index of the column header (row 5 in Excel)


def xlsxwriter_available() -> bool:
    try:
        import xlsxwriter  # noqa: F401
        return True
    except Exception:
        return False


def sheet_names(titles: list[str]) -> list[str]:
    """Excel-valid, unique (case-insensitive), <= 31 chars."""
    used = set(_RESERVED)
    out = []
    for i, title in enumerate(titles, 1):
        clean = ''.join(' ' if ch in _SHEET_FORBIDDEN or ord(ch) < 32 else ch for ch in str(title or ''))
        base = ' '.join(clean.split()).strip("'").strip()[:SHEET_NAME_MAX].rstrip() or f'Sheet {i}'
        cand, n = base, 1
        while cand.casefold() in used:
            n += 1
            suffix = f' ({n})'
            cand = base[:SHEET_NAME_MAX - len(suffix)].rstrip() + suffix
        used.add(cand.casefold())
        out.append(cand)
    return out


def cell_value(col: Column, value: Any, tz: ZoneInfo) -> Any:
    """Normalize a raw dataset value to what goes into the cell."""
    if value is None or value == '':
        return None
    kind = col.kind
    if kind == PERCENT:
        return float(value) / 100.0
    if kind == DURATION:
        return max(float(value), 0.0) / 86400.0  # Excel days -> [h]:mm:ss
    if kind == INT:
        return int(value)
    if kind == NUMBER:
        return float(value)
    if kind == DATETIME and isinstance(value, datetime):
        return (value.astimezone(tz) if value.tzinfo else value).replace(tzinfo=None)
    if kind == DATE and isinstance(value, (date, datetime)):
        return value if not isinstance(value, datetime) else value.date()
    return str(value)


NUMBER_FORMATS = {INT: '#,##0', NUMBER: '#,##0.##', PERCENT: '0.0%', DURATION: '[h]:mm:ss',
                  DATETIME: 'dd/mm/yyyy hh:mm', DATE: 'dd/mm/yyyy'}


def _meta_lines(doc: ReportDocument) -> tuple[str, str]:
    m = doc.metadata
    line1 = (f"Từ ngày: {m.get('date_from') or '—'} | Đến ngày: {m.get('date_to') or '—'} | "
             f"Múi giờ: {m.get('timezone')} | Xuất lúc: {m.get('generated_at')} bởi {m.get('generated_by')}")
    filters = ' | '.join(f'{k}: {v}' for k, v in m.get('filters') or []) or 'Bộ lọc: không'
    if doc.warnings:
        filters += ' | CẢNH BÁO: ' + ' '.join(doc.warnings)
    return line1, filters


def render_xlsx(doc: ReportDocument, *, timezone_name: str, force_fallback: bool = False) -> tuple[bytes, str]:
    """Returns (bytes, engine_name)."""
    if not force_fallback and xlsxwriter_available():
        return _render_xlsxwriter(doc, ZoneInfo(timezone_name)), 'xlsxwriter'
    return _render_openpyxl(doc, ZoneInfo(timezone_name)), 'openpyxl'


def _render_xlsxwriter(doc: ReportDocument, tz: ZoneInfo) -> bytes:
    import xlsxwriter

    out = BytesIO()
    wb = xlsxwriter.Workbook(out, {'in_memory': True, 'strings_to_formulas': False, 'strings_to_urls': False,
                                   'strings_to_numbers': False, 'remove_timezone': True})
    wb.set_properties({'title': doc.title, 'subject': doc.metadata.get('dataset_title', ''),
                       'author': 'MESFlow Report Engine', 'comments': f"generated_by={doc.metadata.get('generated_by')}"})
    f_title = wb.add_format({'bold': True, 'font_size': 15})
    f_meta = wb.add_format({'font_size': 10, 'italic': True})
    f_head = wb.add_format({'bold': True, 'bg_color': '#D9EAF7', 'border': 1, 'border_color': '#B8C2CC',
                            'text_wrap': True, 'align': 'center', 'valign': 'vcenter'})
    base = {'border': 1, 'border_color': '#B8C2CC', 'valign': 'top'}
    f_text = wb.add_format(dict(base, text_wrap=True))
    f_kind = {k: wb.add_format(dict(base, num_format=fmt)) for k, fmt in NUMBER_FORMATS.items()}
    f_total_text = wb.add_format(dict(base, bold=True, bg_color='#F2F5F8'))
    f_total = {k: wb.add_format(dict(base, bold=True, bg_color='#F2F5F8', num_format=fmt))
               for k, fmt in NUMBER_FORMATS.items()}
    line1, line2 = _meta_lines(doc)
    for name, section in zip(sheet_names([s.sheet_name for s in doc.sections]), doc.sections):
        ws = wb.add_worksheet(name)
        last = len(section.columns) - 1
        ws.hide_gridlines(2)
        ws.merge_range(0, 0, 0, last, '', f_title)
        ws.write_string(0, 0, f'{doc.title} · {section.title}', f_title)
        ws.merge_range(1, 0, 1, last, '', f_meta)
        ws.write_string(1, 0, line1, f_meta)
        ws.merge_range(2, 0, 2, last, '', f_meta)
        facts = ' | '.join(f'{k}: {v}' for k, v in section.facts)
        ws.write_string(2, 0, (facts + ' | ' if facts else '') + line2, f_meta)
        for c, col in enumerate(section.columns):
            ws.write_string(HEADER_ROW, c, col.label, f_head)
            ws.set_column(c, c, col.width)
        ws.set_row(HEADER_ROW, 30)
        r = HEADER_ROW
        for r_off, row in enumerate(section.rows, 1):
            r = HEADER_ROW + r_off
            for c, col in enumerate(section.columns):
                _xw_write(ws, r, c, cell_value(col, row.get(col.key), tz), col, f_text, f_kind)
        last_data = HEADER_ROW + len(section.rows)
        if section.totals is not None:
            r = last_data + 1
            for c, col in enumerate(section.columns):
                _xw_write(ws, r, c, cell_value(col, section.totals.get(col.key), tz), col, f_total_text, f_total)
        ws.freeze_panes(HEADER_ROW + 1, 0)
        ws.autofilter(HEADER_ROW, 0, max(last_data, HEADER_ROW), last)
        ws.repeat_rows(HEADER_ROW)
        ws.set_landscape()
        ws.set_paper(9)  # A4
        ws.fit_to_pages(1, 0)
        ws.set_margins(left=0.25, right=0.25, top=0.45, bottom=0.5)
        ws.set_footer('&L&8MESFlow · ' + doc.metadata.get('generated_at', '').replace('&', '&&') + '&R&8Trang &P/&N')
        ws.print_area(0, 0, max(r, HEADER_ROW), last)
    wb.close()
    return out.getvalue()


def _xw_write(ws, r: int, c: int, value: Any, col: Column, f_text, f_kind) -> None:
    if value is None:
        ws.write_blank(r, c, None, f_kind.get(col.kind, f_text))
    elif col.kind in (DATETIME, DATE) and isinstance(value, (date, datetime)):
        ws.write_datetime(r, c, value if isinstance(value, datetime) else datetime(value.year, value.month, value.day),
                          f_kind[col.kind])
    elif isinstance(value, (int, float)) and not isinstance(value, bool):
        ws.write_number(r, c, value, f_kind.get(col.kind, f_text))
    else:
        ws.write_string(r, c, str(value), f_text)  # never write(): no formula/url coercion


def _render_openpyxl(doc: ReportDocument, tz: ZoneInfo) -> bytes:
    """Fallback when XlsxWriter is not installed. Same sheets, cells and
    number formats; plainer styling."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    line1, line2 = _meta_lines(doc)
    for name, section in zip(sheet_names([s.sheet_name for s in doc.sections]), doc.sections):
        ws = wb.create_sheet(name)
        _op_text(ws.cell(1, 1), f'{doc.title} · {section.title}').font = Font(bold=True, size=15)
        _op_text(ws.cell(2, 1), line1)
        facts = ' | '.join(f'{k}: {v}' for k, v in section.facts)
        _op_text(ws.cell(3, 1), (facts + ' | ' if facts else '') + line2)
        hr = HEADER_ROW + 1
        for c, col in enumerate(section.columns, 1):
            cell = _op_text(ws.cell(hr, c), col.label)
            cell.font = Font(bold=True)
            cell.fill = PatternFill('solid', fgColor='D9EAF7')
            ws.column_dimensions[get_column_letter(c)].width = col.width
        rows = list(section.rows) + ([section.totals] if section.totals is not None else [])
        for r_off, row in enumerate(rows, 1):
            for c, col in enumerate(section.columns, 1):
                value = cell_value(col, row.get(col.key), tz)
                cell = ws.cell(hr + r_off, c)
                if isinstance(value, str):
                    _op_text(cell, value)
                else:
                    cell.value = value
                    if col.kind in NUMBER_FORMATS and value is not None:
                        cell.number_format = NUMBER_FORMATS[col.kind]
                if section.totals is not None and r_off == len(rows):
                    cell.font = Font(bold=True)
        last_letter = get_column_letter(len(section.columns))
        ws.freeze_panes = f'A{hr + 1}'
        ws.auto_filter.ref = f'A{hr}:{last_letter}{hr + len(section.rows)}'
        ws.print_title_rows = f'{hr}:{hr}'
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.oddFooter.right.text = 'Trang &[Page]/&[Pages]'
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def _op_text(cell, text: str):
    cell.value = text
    cell.data_type = 's'  # never a formula, even when text starts with '='
    return cell
